# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The two tender workbooks: the bill sent out to bidders and the price comparison that comes back.

The **bill for bidders** is the package's own lines (the part of the bill the
package was raised over, not the whole bill) with quantities and descriptions
and nothing of the contractor's own pricing: no unit rate, no total, no
markup, no metadata. Each line has a yellow unit price cell to fill and a line
total that calculates itself. The sheet is protected without a password so a
quantity is not changed by accident, and the price and note cells stay open.
The greyed reference column carries the position id so a priced sheet can be
matched back line by line.

The **price comparison** is the comparison endpoint's answer laid out as a
sheet: one row per bill line, the contractor's own estimate, then a unit price
and a line total per bidder. It does not recompute anything the endpoint
answers. It adds three readings of those numbers, the same ones the screen
shows: the lowest unit price per line, a missing price written as a word
rather than as 0, and the outliers ``classifyCell`` marks in
``frontend/src/features/tendering/analysis.ts`` (more than
:data:`OUTLIER_THRESHOLD` away from the median of the line's prices).

Both renderers are pure: they take plain values and return bytes, so the route
can run them in a worker thread.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

from app.core.csv_safety import neutralise_formula
from app.modules.tendering.xlsx_translations import xt

__all__ = [
    "LOWEST_FILL",
    "MISSING_FILL",
    "OUTLIER_THRESHOLD",
    "BidderLine",
    "bidder_lines",
    "classify_rate",
    "render_bidder_xlsx",
    "render_comparison_xlsx",
    "unit_code",
]

# Kept equal to the default ``threshold`` of ``classifyCell`` in
# frontend/src/features/tendering/analysis.ts so the sheet flags what the
# screen flags.
OUTLIER_THRESHOLD = Decimal("0.15")

LOWEST_FILL = "C6EFCE"
MISSING_FILL = "FCE4D6"
INPUT_FILL = "FFF2CC"
_SECTION_FILL = "E7E6E6"
_HEADER_FILL = "D9D9D9"
_REFERENCE_FILL = "F2F2F2"
_LOWEST_FONT = "006100"
_HIGH_FONT = "C00000"
_LOW_FONT = "9C5700"
_MISSING_FONT = "974706"

_MONEY = "#,##0.00"
_QTY = "#,##0.###"
_HIGH_FORMAT = '#,##0.00" ▲"'
_LOW_FORMAT = '#,##0.00" ▼"'

_SECTION_UNITS = frozenset({"", "section"})


def _is_section(unit: object) -> bool:
    # The BOQ Excel export reads a row the same way: a header carries no unit.
    return str(unit or "").strip().lower() in _SECTION_UNITS


# Locale trade spellings, the ones ``LOCALE_UNIT_CODES`` in
# frontend/src/shared/lib/unitLabels.ts applies in the BOQ grid, for the
# languages this catalogue writes: a German bill says "psch" and "Stk", not the
# stored "lsum" and "pcs".
_LOCALE_UNIT_CODES: dict[str, dict[str, str]] = {
    "de": {"lsum": "psch", "ls": "psch", "lump_sum": "psch", "pcs": "Stk", "ea": "Stk"},
    # Kept equal to the editor's table (tests/unit/test_tender_unit_codes_match_the_editor.py),
    # so it is ready once the sheet catalogue speaks Croatian.
    "hr": {"lm": "m'", "lsum": "pauš.", "ls": "pauš.", "lump_sum": "pauš.", "pcs": "kom", "ea": "kom"},
}


def unit_code(unit: str, locale: str) -> str:
    """The unit as the BOQ editor shows it: ``m2`` as ``m²``, ``lsum`` as ``psch`` in German.

    Display only, like the editor's ``localizedUnitCode``; the stored token is
    unchanged. Unknown tokens pass through.

    Args:
        unit: The stored unit token.
        locale: The sheet language.

    Returns:
        The unit to print.
    """
    from app.core.unit_conversion import display_unit_for

    token = (unit or "").strip()
    if not token:
        return ""
    trade = _LOCALE_UNIT_CODES.get(locale, {}).get(token.lower())
    return trade or display_unit_for(token)


def _dec(value: object) -> Decimal:
    try:
        d = Decimal(str(value).strip())
    except Exception:  # noqa: BLE001 - any unreadable figure is zero, like the service's parser
        return Decimal("0")
    return d if d.is_finite() else Decimal("0")


def classify_rate(rate: Decimal, line_rates: Sequence[Decimal]) -> str | None:
    """Flag a unit price as far from the median of the line, as the comparison screen does.

    Port of ``classifyCell`` in ``frontend/src/features/tendering/analysis.ts``.
    Only prices above zero count, and a line needs two of them to have a median
    worth comparing with.

    Args:
        rate: The bidder's unit price on the line.
        line_rates: Every comparable bidder's price on the same line.

    Returns:
        ``"high"``, ``"low"`` or ``None``.
    """
    priced = sorted(r for r in line_rates if r > 0)
    if len(priced) < 2 or rate <= 0:
        return None
    mid = len(priced) // 2
    median = priced[mid] if len(priced) % 2 else (priced[mid - 1] + priced[mid]) / 2
    if median <= 0:
        return None
    if rate >= median * (1 + OUTLIER_THRESHOLD):
        return "high"
    if rate <= median * (1 - OUTLIER_THRESHOLD):
        return "low"
    return None


def _set_properties(wb: Any, title: str) -> None:
    wb.properties.creator = "OpenConstructionERP"
    wb.properties.lastModifiedBy = "OpenConstructionERP"
    wb.properties.title = title


def _save(wb: Any) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── Price comparison ────────────────────────────────────────────────────────


def render_comparison_xlsx(
    comparison: Any,
    *,
    project_name: str,
    reporting_currency: str,
    locale: str,
    exported_at: datetime | None = None,
) -> bytes:
    """Write the price comparison workbook from the comparison endpoint's answer.

    Args:
        comparison: A ``BidComparisonResponse``.
        project_name: The project the package belongs to.
        reporting_currency: The currency bids are ranked in (the project's).
            A bid quoted in another currency is shown but never marked lowest
            or as an outlier, since its figures are not comparable. ``""`` when
            unknown, in which case every bid is compared.
        locale: A member of the workbook catalogue.
        exported_at: The moment printed in the header, now when omitted.

    Returns:
        The xlsx bytes.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from app.core.xlsx_branding import apply_company_header
    from app.core.xlsx_text import store_strings_as_text

    reporting = (reporting_currency or "").strip().upper()
    bids = list(comparison.bid_totals or [])
    bid_ids = [str(b.get("bid_id")) for b in bids]

    def _ccy(bid: dict) -> str:
        return str(bid.get("currency") or "").strip().upper()

    comparable = {str(b.get("bid_id")) for b in bids if not reporting or not _ccy(b) or _ccy(b) == reporting}

    wb = Workbook()
    ws = wb.active
    ws.title = xt(locale, "comparison_sheet")[:31]

    bold = Font(bold=True)
    header_fill = PatternFill(start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type="solid")
    section_fill = PatternFill(start_color=_SECTION_FILL, end_color=_SECTION_FILL, fill_type="solid")
    lowest_fill = PatternFill(start_color=LOWEST_FILL, end_color=LOWEST_FILL, fill_type="solid")
    missing_fill = PatternFill(start_color=MISSING_FILL, end_color=MISSING_FILL, fill_type="solid")
    top_border = Border(top=Side(style="medium"))
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    missing_marker = xt(locale, "not_priced")

    # Two header rows: the group (estimate, each bidder) over its two columns,
    # then what each column holds.
    fixed = [xt(locale, "item"), xt(locale, "description"), xt(locale, "unit"), xt(locale, "quantity")]
    est_col = len(fixed) + 1
    first_bid_col = est_col + 2
    last_col = first_bid_col + 2 * len(bids) - 1 if bids else est_col + 1

    def _bid_col(index: int) -> int:
        return first_bid_col + 2 * index

    groups: list[tuple[int, str]] = [(est_col, xt(locale, "own_estimate"))]
    for i, bid in enumerate(bids):
        label = str(bid.get("company_name") or "")
        if _ccy(bid):
            label = f"{label} ({_ccy(bid)})"
        groups.append((_bid_col(i), label))
    for col, label in groups:
        cell = ws.cell(row=1, column=col, value=neutralise_formula(label))
        cell.font = bold
        cell.alignment = centre
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
    for col, label in enumerate(fixed, start=1):
        ws.cell(row=2, column=col, value=label)
    for col, _label in groups:
        ws.cell(row=2, column=col, value=xt(locale, "unit_price"))
        ws.cell(row=2, column=col + 1, value=xt(locale, "line_total"))
    for row in (1, 2):
        for col in range(1, last_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = header_fill
            cell.font = bold
            if row == 2:
                cell.alignment = Alignment(wrap_text=True, vertical="center")

    current = 3
    line_count = 0
    priced_count = dict.fromkeys(bid_ids, 0)
    priced_sum = dict.fromkeys(bid_ids, Decimal("0"))
    estimate_line_sum = Decimal("0")

    for row in comparison.rows:
        ws.cell(row=current, column=1, value=neutralise_formula(row.ordinal or ""))
        ws.cell(row=current, column=2, value=neutralise_formula(row.description or ""))
        if _is_section(row.unit):
            for col in range(1, last_col + 1):
                ws.cell(row=current, column=col).fill = section_fill
                ws.cell(row=current, column=col).font = bold
            current += 1
            continue

        line_count += 1
        ws.cell(row=current, column=3, value=neutralise_formula(unit_code(row.unit or "", locale)))
        qty = ws.cell(row=current, column=4, value=_dec(row.budget_quantity))
        qty.number_format = _QTY
        est_rate = ws.cell(row=current, column=est_col, value=_dec(row.budget_rate))
        est_rate.number_format = _MONEY
        est_total = ws.cell(row=current, column=est_col + 1, value=_dec(row.budget_total))
        est_total.number_format = _MONEY
        estimate_line_sum += _dec(row.budget_total)

        entries = {str(e.get("bid_id")): e for e in (row.bids or [])}
        comparable_rates = [
            _dec(e.get("unit_rate")) for bid_id, e in entries.items() if bid_id in comparable and e.get("priced", True)
        ]
        positive = [r for r in comparable_rates if r > 0]
        lowest = min(positive) if len(positive) >= 2 else None

        for i, bid_id in enumerate(bid_ids):
            col = _bid_col(i)
            entry = entries.get(bid_id)
            # ``priced`` is absent only on an answer older than the flag; there a
            # line the bid holds is the best reading available.
            if entry is None or not entry.get("priced", True):
                marker = ws.cell(row=current, column=col, value=missing_marker)
                marker.font = Font(italic=True, color=_MISSING_FONT)
                marker.fill = missing_fill
                marker.alignment = Alignment(horizontal="right")
                ws.cell(row=current, column=col + 1).fill = missing_fill
                continue
            rate = _dec(entry.get("unit_rate"))
            total = _dec(entry.get("total"))
            priced_count[bid_id] += 1
            priced_sum[bid_id] += total
            rate_cell = ws.cell(row=current, column=col, value=rate)
            rate_cell.number_format = _MONEY
            total_cell = ws.cell(row=current, column=col + 1, value=total)
            total_cell.number_format = _MONEY
            if bid_id not in comparable:
                continue
            flag = classify_rate(rate, comparable_rates)
            if flag == "high":
                rate_cell.number_format = _HIGH_FORMAT
                rate_cell.font = Font(color=_HIGH_FONT)
            elif flag == "low":
                rate_cell.number_format = _LOW_FORMAT
                rate_cell.font = Font(color=_LOW_FONT)
            if lowest is not None and rate == lowest:
                rate_cell.fill = lowest_fill
                rate_cell.font = Font(bold=True, color=rate_cell.font.color.rgb if flag else _LOWEST_FONT)
        current += 1

    # ── Footer: what the lines add up to, what each bidder submitted ────────
    footer_start = current
    totals_by_bid = {str(b.get("bid_id")): b for b in bids}

    def _label(row: int, key: str) -> None:
        cell = ws.cell(row=row, column=2, value=xt(locale, key))
        cell.font = bold

    _label(current, "sum_priced_lines")
    ws.cell(row=current, column=est_col + 1, value=estimate_line_sum).number_format = _MONEY
    for i, bid_id in enumerate(bid_ids):
        ws.cell(row=current, column=_bid_col(i) + 1, value=priced_sum[bid_id]).number_format = _MONEY
    current += 1

    _label(current, "bid_total")
    est_cell = ws.cell(row=current, column=est_col + 1, value=_dec(comparison.budget_total))
    est_cell.number_format = _MONEY
    est_cell.font = bold
    # A bid that left lines unpriced is cheaper by what it left out, so only
    # complete bids compete for the lowest total.
    complete = {b for b in bid_ids if priced_count[b] == line_count}
    comparable_totals = [
        _dec(totals_by_bid[b].get("total"))
        for b in bid_ids
        if b in comparable and b in complete and _dec(totals_by_bid[b].get("total")) > 0
    ]
    lowest_total = min(comparable_totals) if len(comparable_totals) >= 2 else None
    for i, bid_id in enumerate(bid_ids):
        total = _dec(totals_by_bid[bid_id].get("total"))
        cell = ws.cell(row=current, column=_bid_col(i) + 1, value=total)
        cell.number_format = _MONEY
        cell.font = bold
        if lowest_total is not None and bid_id in comparable and bid_id in complete and total == lowest_total:
            cell.fill = lowest_fill
            cell.font = Font(bold=True, color=_LOWEST_FONT)
    bid_total_row = current
    current += 1

    _label(current, "lines_priced")
    for i, bid_id in enumerate(bid_ids):
        cell = ws.cell(
            row=current,
            column=_bid_col(i) + 1,
            value=xt(locale, "lines_priced_value", priced=priced_count[bid_id], total=line_count),
        )
        cell.alignment = Alignment(horizontal="right")
        if bid_id not in complete:
            cell.fill = missing_fill
            cell.font = Font(bold=True, color=_MISSING_FONT)
    current += 1

    _label(current, "deviation")
    for i, bid_id in enumerate(bid_ids):
        bt = totals_by_bid[bid_id]
        col = _bid_col(i) + 1
        if bt.get("deviation_known"):
            cell = ws.cell(row=current, column=col, value=_dec(bt.get("deviation_pct")) / 100)
            cell.number_format = "+0.0%;-0.0%;0.0%"
        else:
            cell = ws.cell(row=current, column=col, value=xt(locale, "not_comparable"))
            cell.alignment = Alignment(horizontal="right")
    current += 1

    for col in range(1, last_col + 1):
        ws.cell(row=footer_start, column=col).border = top_border
    for col in range(1, last_col + 1):
        ws.cell(row=bid_total_row, column=col).border = Border(top=Side(style="thin"))

    # ── Legend ──────────────────────────────────────────────────────────────
    current += 1
    ws.cell(row=current, column=2, value=xt(locale, "legend")).font = bold
    current += 1
    pct = int(OUTLIER_THRESHOLD * 100)
    legend = [
        (xt(locale, "legend_lowest"), lowest_fill, None),
        (xt(locale, "legend_high", pct=pct), None, _HIGH_FONT),
        (xt(locale, "legend_low", pct=pct), None, _LOW_FONT),
        (xt(locale, "legend_missing", marker=missing_marker), missing_fill, _MISSING_FONT),
        (xt(locale, "legend_incomplete"), missing_fill, _MISSING_FONT),
        (xt(locale, "legend_totals"), None, None),
    ]
    for bid in bids:
        if str(bid.get("bid_id")) not in comparable:
            legend.append(
                (
                    f"{bid.get('company_name') or ''}: "
                    + xt(locale, "other_currency", currency=_ccy(bid), reporting=reporting),
                    None,
                    None,
                )
            )
    for text, fill, colour in legend:
        cell = ws.cell(row=current, column=2, value=neutralise_formula(text))
        if fill is not None:
            cell.fill = fill
        if colour is not None:
            cell.font = Font(color=colour)
        current += 1

    # ── Layout ──────────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 12
    for col in range(est_col, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15
    for row_cells in ws.iter_rows(min_row=3, max_row=footer_start - 1, min_col=2, max_col=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=3, column=est_col)

    when = (exported_at or datetime.now(tz=UTC)).astimezone(UTC).date().isoformat()
    details = [
        f"{xt(locale, 'project')}: {project_name}" if project_name else None,
        f"{xt(locale, 'package')}: {comparison.package_name}",
        "  |  ".join(
            [*([f"{xt(locale, 'currency')}: {reporting}"] if reporting else []), f"{xt(locale, 'date')}: {when}"]
        ),
    ]
    store_strings_as_text(ws)
    apply_company_header(ws, title=xt(locale, "comparison_title"), details=details)
    _set_properties(wb, f"{xt(locale, 'comparison_title')} - {comparison.package_name}")
    return _save(wb)


# ── Bill for bidders ────────────────────────────────────────────────────────


@dataclass(frozen=True)
class BidderLine:
    """One row of the bill sent to bidders. Carries nothing priced."""

    is_section: bool
    ordinal: str
    description: str
    unit: str
    quantity: Decimal
    position_id: str


def bidder_lines(all_positions: Sequence[Any], in_scope: Sequence[Any]) -> list[BidderLine]:
    """The rows of the bill for bidders: the package's lines under their section headers.

    Args:
        all_positions: Every exportable position of the bill, headers
            included, in bill order.
        in_scope: The positions the package was raised over
            (``_positions_in_scope``), which may or may not hold headers.

    Returns:
        The package's priceable lines in bill order, each preceded by the
        headers above it. A header with no line of the package under it is
        left out, so a bidder is not shown sections nobody asked them to price.
    """
    by_id = {str(p.id): p for p in all_positions}
    lines = {str(p.id) for p in in_scope if not _is_section(getattr(p, "unit", ""))}
    headers: set[str] = set()
    for pid in lines:
        parent = getattr(by_id.get(pid), "parent_id", None)
        seen: set[str] = set()
        while parent is not None and str(parent) not in seen:
            seen.add(str(parent))
            headers.add(str(parent))
            parent = getattr(by_id.get(str(parent)), "parent_id", None)
    out: list[BidderLine] = []
    for p in all_positions:
        pid = str(p.id)
        is_header = _is_section(getattr(p, "unit", ""))
        if (is_header and pid in headers) or (not is_header and pid in lines):
            out.append(
                BidderLine(
                    is_section=is_header,
                    ordinal=str(getattr(p, "ordinal", "") or ""),
                    description=str(getattr(p, "description", "") or ""),
                    unit="" if is_header else str(getattr(p, "unit", "") or ""),
                    quantity=Decimal("0") if is_header else _dec(getattr(p, "quantity", 0)),
                    position_id=pid,
                )
            )
    return out


def render_bidder_xlsx(
    lines: Sequence[BidderLine],
    *,
    project_name: str,
    package_name: str,
    currency: str,
    deadline: str,
    locale: str,
    exported_at: datetime | None = None,
) -> bytes:
    """Write the bill a bidder prices: quantities in, unit prices to fill, no contractor rates.

    Args:
        lines: From :func:`bidder_lines`.
        project_name: The project the package belongs to.
        package_name: The tender package.
        currency: The currency to price in, ``""`` when the project has none.
        deadline: The return date as text, ``""`` when the package has none.
        locale: A member of the workbook catalogue.
        exported_at: The moment printed in the header, now when omitted.

    Returns:
        The xlsx bytes.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    from app.core.xlsx_branding import apply_company_header
    from app.core.xlsx_text import store_strings_as_text

    wb = Workbook()
    ws = wb.active
    ws.title = xt(locale, "bidder_sheet")[:31]

    bold = Font(bold=True)
    header_fill = PatternFill(start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type="solid")
    section_fill = PatternFill(start_color=_SECTION_FILL, end_color=_SECTION_FILL, fill_type="solid")
    input_fill = PatternFill(start_color=INPUT_FILL, end_color=INPUT_FILL, fill_type="solid")
    reference_fill = PatternFill(start_color=_REFERENCE_FILL, end_color=_REFERENCE_FILL, fill_type="solid")
    open_cell = Protection(locked=False)

    # Who is pricing, then how, then the table.
    ws.cell(row=1, column=1, value=xt(locale, "bidder")).font = bold
    name = ws.cell(row=1, column=2)
    name.fill = input_fill
    name.protection = open_cell
    instructions = ws.cell(row=2, column=1, value=xt(locale, "instructions"))
    instructions.font = Font(italic=True)
    instructions.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    ws.row_dimensions[2].height = 45

    price_header = xt(locale, "unit_price") + (f" ({currency})" if currency else "")
    total_header = xt(locale, "line_total") + (f" ({currency})" if currency else "")
    headers = [
        xt(locale, "item"),
        xt(locale, "description"),
        xt(locale, "unit"),
        xt(locale, "quantity"),
        price_header,
        total_header,
        xt(locale, "bidder_note"),
        xt(locale, "reference"),
    ]
    header_row = 4
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    ws.add_data_validation(validation)

    current = header_row + 1
    first_line = current
    line_rows: list[int] = []
    for line in lines:
        ws.cell(row=current, column=1, value=neutralise_formula(line.ordinal))
        desc = ws.cell(row=current, column=2, value=neutralise_formula(line.description))
        desc.alignment = Alignment(wrap_text=True, vertical="top")
        ref = ws.cell(row=current, column=8, value=line.position_id)
        ref.fill = reference_fill
        ref.font = Font(color="808080", size=8)
        if line.is_section:
            for col in range(1, 8):
                ws.cell(row=current, column=col).fill = section_fill
                ws.cell(row=current, column=col).font = bold
            current += 1
            continue
        ws.cell(row=current, column=3, value=neutralise_formula(unit_code(line.unit, locale)))
        ws.cell(row=current, column=4, value=line.quantity).number_format = _QTY
        price = ws.cell(row=current, column=5)
        price.fill = input_fill
        price.number_format = _MONEY
        price.protection = open_cell
        validation.add(price.coordinate)
        ws.cell(row=current, column=6).number_format = _MONEY
        line_rows.append(current)
        note = ws.cell(row=current, column=7)
        note.protection = open_cell
        note.alignment = Alignment(wrap_text=True, vertical="top")
        current += 1
    last_line = current - 1

    ws.cell(row=current, column=2, value=xt(locale, "total")).font = bold
    grand = ws.cell(row=current, column=6)
    grand.number_format = _MONEY
    grand.font = bold
    for col in range(1, 9):
        ws.cell(row=current, column=col).border = Border(top=Side(style="medium"))

    for letter, width in (("A", 12), ("B", 56), ("C", 8), ("D", 12), ("E", 16), ("F", 16), ("G", 28), ("H", 12)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)

    # Quantities and texts are the tender; the bidder only adds prices and
    # notes. No password: this guards against a slip, not against the bidder.
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.formatCells = False

    when = (exported_at or datetime.now(tz=UTC)).astimezone(UTC).date().isoformat()
    third = [f"{xt(locale, 'currency')}: {currency}"] if currency else []
    if deadline:
        third.append(f"{xt(locale, 'return_by')}: {deadline}")
    third.append(f"{xt(locale, 'date')}: {when}")
    details = [
        f"{xt(locale, 'project')}: {project_name}" if project_name else None,
        f"{xt(locale, 'package')}: {package_name}",
        "  |  ".join(third),
    ]
    # The bill's text goes in as text first; only then the sheet's own
    # formulas, which store_strings_as_text would otherwise retype.
    store_strings_as_text(ws)
    for row in line_rows:
        ws.cell(row=row, column=6, value=f'=IF(E{row}="","",D{row}*E{row})')
    grand.value = f"=SUM(F{first_line}:F{max(last_line, first_line)})"
    apply_company_header(ws, title=xt(locale, "bidder_title"), details=details)
    _set_properties(wb, f"{xt(locale, 'bidder_title')} - {package_name}")
    return _save(wb)
