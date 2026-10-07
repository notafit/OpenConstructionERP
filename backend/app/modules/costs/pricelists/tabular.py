# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Price lists published as CSV, XLSX or JSON tables, read through one profile.

Each region lays its table out its own way. Read from the published files:

* Puglia 2026 (CSV, comma): ``codice_pug2026``, ``tipologia``, ``capitolo``,
  ``voce``, ``articolo``, ``um``, ``TOTALE_AT/RU/PR``, ``INCIDENZA_AT/RU/PR``,
  ``SPESE_GENERALI``, ``UTILE_DI_IMPRESA``, ``TOTALE_GENERALE``; float noise
  such as "8.6300000000000008".
* Campania 2024 (CSV, semicolon, multi-line quoted text) and Toscana 2025
  (CSV, pipe), one layout: ``Codice``, ``Tipologia``, ``Capitolo``, ``Voce``,
  ``Articolo``, ``Unita' di misura``, ``Prezzo``, ``Manodopera %`` (the
  labour amount in euro, with values such as "15.403.443" whose decimal
  separator is lost), ``Incidenza Manodopera sul totale`` (the percentage).
* Piemonte 2023 (CSV, semicolon, quoted): ``CODICE_TOTALE``, ``SEZIONE``,
  ``CAPITOLO``, ``ARTICOLO``, ``SUBARTICOLO``, ``SIMBOLO``, ``PREZZO_LORDO``,
  ``INCIDENZA_MANODOPERA``, ``SOMMA_MANODOPERA``.
* Umbria 2025 (XLSX one sheet per chapter, and the same rows as cp1252 JSON):
  ``Numero d'ordine``, ``Descrizione dell'articolo``, ``u.m.``, ``prezzo €``,
  ``costo minimo manodopera €`` (labour in euro, not percent); headings are
  rows with no price, a ``.0`` code is the parent of the codes under it.
* Lazio 2023 (CSV with no header row at all; part A in cp1252, the other parts
  in UTF-8): part letter, code with trailing dot and non-breaking spaces,
  description, unit, "€ 1 826,34"; part E puts the code in the first column.

So the reader maps columns by header aliases (the "profile"), and a file with
no header row is read by position around its price column. Where rows carry
no chapter columns, the chapter path comes from the code: a row without a
price is a heading, and the headings whose codes prefix a voce's code are its
chapters and parent descriptions.
"""

from __future__ import annotations

import codecs
import csv
import io
import itertools
import json
import os
import re
import shutil
import tempfile
import unicodedata
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from app.modules.costs.pricelists.base import (
    REGION_NAMES,
    BrokenNumber,
    NumberConvention,
    PriceListRow,
    PriceListSource,
    amount_or_flag,
    clean_text,
    normalise_unit,
    parse_amount,
    region_from_code,
)
from app.modules.costs.pricelists.containers import ContainerRefused, Member, guard_archive

FORMAT_ID = "tabular"

# JSON is read whole; a list published as JSON is a few MB.
_MAX_JSON_BYTES = 16 * 1024 * 1024
_HEADER_SCAN_ROWS = 30
_MAX_COLUMNS = 80
csv.field_size_limit(16 * 1024 * 1024)


@dataclass
class TabularProfile:
    """Header aliases per role, normalised (lower case, no accents, no punctuation).

    The first alias a file carries wins for each role, so a list keeps working
    when a region adds a column that also matches a later alias.
    """

    aliases: dict[str, tuple[str, ...]] = field(
        default_factory=lambda: {
            "code": (
                "codice totale",
                "codice",
                "codice articolo",
                "numero d ordine",
                "n ordine",
                "tariffa",
                "codice voce",
                "cod",
            ),
            "description": (
                "descrizione dell articolo",
                "descrizione estesa",
                "descrizione",
                "declaratoria",
                "descrizione voce",
            ),
            "group_text": ("voce",),
            "item_text": ("subarticolo", "articolo", "descrizione breve", "sottovoce"),
            "unit": ("unita di misura", "u m", "um", "udm", "simbolo", "unita misura", "unita"),
            "rate": ("prezzo", "prezzo lordo", "totale generale", "prezzo unitario", "prezzo euro", "importo", "euro"),
            "labour_pct": (
                "incidenza manodopera",
                "incidenza ru",
                "incidenza manodopera sul totale",
                "% manodopera",
            ),
            # "Manodopera %" holds the labour amount in euro despite its name:
            # on the Toscana 2025 list (same layout as Campania) the column
            # reads 5.50138 where the XML edition gives incidenzamanodopera
            # valore="5.50138" percentuale="40.36" for the same voce.
            "labour_amount": (
                "costo minimo manodopera",
                "somma manodopera",
                "totale ru",
                "costo manodopera",
                "manodopera %",
            ),
            "safety_pct": (
                "di cui oneri di sicurezza afferenti l impresa %",
                "incidenza oneri di sicurezza sul totale",
                "incidenza sicurezza",
                "oneri sicurezza %",
            ),
            "safety_amount": ("di cui oneri di sicurezza afferenti l impresa", "oneri di sicurezza", "oneri sicurezza"),
            "material_amount": ("totale pr",),
            "equipment_amount": ("totale at",),
            "edition": ("edizione",),
        }
    )
    #: Chapter columns, outermost first; each may come with a ``codice <name>`` column.
    chapter_columns: tuple[str, ...] = (
        "tipologia",
        "tipologia famiglia",
        "sezione",
        "settore",
        "capitolo",
        "categoria",
    )
    #: A header that names the region's own prefixed code (``codice_pug2026``).
    prefixed_code_header: re.Pattern[str] = field(
        default_factory=lambda: re.compile(r"^codice (?:" + "|".join(c.lower() for c in REGION_NAMES) + r")\d{2,4}$")
    )


DEFAULT_PROFILE = TabularProfile()

# Layouts recognised by their headers, for the region a file does not state in
# its codes. Advisory: the preview shows the guess and the user confirms it.
_LAYOUT_REGIONS: tuple[tuple[frozenset[str], str], ...] = (
    (frozenset({"codice totale", "prezzo lordo"}), "PIE"),
    (frozenset({"numero d ordine", "costo minimo manodopera"}), "UMB"),
    (frozenset({"di cui oneri di sicurezza afferenti l impresa"}), "CAM"),
    (frozenset({"incidenza ru", "totale generale"}), "PUG"),
)

_MONEY_CELL_RE = re.compile(r"^\s*€\s*[\d\s .]*\d(?:,\d+)?\s*$")

# The columns that hold amounts, whose cells say which decimal separator a
# file writes. Codes ("1.1.20") and descriptions are not asked.
_AMOUNT_ROLES: tuple[str, ...] = (
    "rate",
    "labour_pct",
    "labour_amount",
    "safety_pct",
    "safety_amount",
    "material_amount",
    "equipment_amount",
)

# Rows in a row without an amount written as text, after which a table is
# taken to write its amounts as real numbers and is not read further for them.
_TEXTLESS_ROWS = 2000

_NUMBERED_HEADING_RE = re.compile(r"^([A-Z]{0,2}\d+(?:\.\d+)*)\.?\s+(\S.*)$")


def normalise_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower()
    text = text.replace("€", " ").replace("%", " % ")
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    text = re.sub(r"\beur(?:o)?\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def sniff(head: bytes, name: str) -> bool:
    return os.path.splitext(name.lower())[1] in {".csv", ".txt", ".xlsx", ".json"}


# ── Raw tables ───────────────────────────────────────────────────────────────


# The wide Unicode forms by their byte-order mark, longest first: the UTF-32
# little-endian mark begins with the UTF-16 one. The codec named strips the mark.
_WIDE_MARKS: tuple[tuple[bytes, str, str], ...] = (
    (codecs.BOM_UTF32_LE, "utf-32", "UTF-32"),
    (codecs.BOM_UTF32_BE, "utf-32", "UTF-32"),
    (codecs.BOM_UTF16_LE, "utf-16", "UTF-16"),
    (codecs.BOM_UTF16_BE, "utf-16", "UTF-16"),
)


def _unmarked_utf16(head: bytes) -> str | None:
    """``utf-16-le`` or ``utf-16-be`` when ``head`` has the NUL every other byte that UTF-16 text has."""
    pairs = len(head) // 2
    if pairs < 8:
        return None
    odd = sum(1 for i in range(1, pairs * 2, 2) if head[i] == 0)
    even = sum(1 for i in range(0, pairs * 2, 2) if head[i] == 0)
    if odd >= pairs * 0.6 and even <= pairs * 0.05:
        return "utf-16-le"
    if even >= pairs * 0.6 and odd <= pairs * 0.05:
        return "utf-16-be"
    return None


def _decodes(member: Member, encoding: str) -> bool:
    decoder = codecs.getincrementaldecoder(encoding)()
    with member.open() as stream:
        try:
            while chunk := stream.read(1024 * 1024):
                decoder.decode(chunk, final=False)
            decoder.decode(b"", final=True)
        except UnicodeDecodeError:
            return False
    return True


def _detect_encoding(member: Member) -> str:
    """The encoding to read a text table in.

    UTF-16 or UTF-32 by its byte-order mark, the way Excel saves "Unicode
    text", or UTF-16 without a mark by its NUL every other byte; then UTF-8
    with or without a mark if the whole member decodes as UTF-8, else
    Windows-1252. Checked over the whole file, not its head: a Lazio part can
    be plain ASCII for its first hundred kilobytes and cp1252 after.

    Raises:
        ContainerRefused: ``text_encoding_unreadable`` when the file is in a
            wide Unicode form that does not decode (``encoding`` names it), or
            holds NUL bytes no text encoding we read explains (``encoding``
            is None): Windows-1252 would read such a file as nonsense.
    """
    head = member.head(64 * 1024)
    for mark, codec, label in _WIDE_MARKS:
        if head.startswith(mark):
            if _decodes(member, codec):
                return codec
            raise ContainerRefused("text_encoding_unreadable", encoding=label, member=member.name)
    wide = _unmarked_utf16(head)
    if wide is not None:
        if _decodes(member, wide):
            return wide
        raise ContainerRefused("text_encoding_unreadable", encoding="UTF-16", member=member.name)
    if head.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if _decodes(member, "utf-8"):
        return "utf-8"
    if b"\x00" in head:
        raise ContainerRefused("text_encoding_unreadable", encoding=None, member=member.name)
    return "cp1252"


def _sniff_delimiter(sample: str) -> str:
    lines = [line for line in sample.splitlines()[:40] if line.strip()]
    counts = {d: sum(line.count(d) for line in lines) for d in (";", ",", "\t", "|")}
    return max(counts, key=lambda d: counts[d]) if any(counts.values()) else ";"


def _csv_rows(member: Member) -> Iterator[list[Any]]:
    encoding = _detect_encoding(member)
    with member.open() as raw:
        text = io.TextIOWrapper(raw, encoding=encoding, errors="replace", newline="")
        # The sample may end mid-line; the rest of that line goes with it.
        sample = text.read(64 * 1024) + text.readline()
        reader = csv.reader(itertools.chain(io.StringIO(sample), text), delimiter=_sniff_delimiter(sample))
        for row in reader:
            yield row[:_MAX_COLUMNS]


def _xlsx_tables(member: Member) -> Iterator[tuple[str, Iterator[list[Any]]]]:
    from openpyxl import load_workbook

    with member.open() as raw, tempfile.TemporaryFile() as spool:
        # openpyxl needs a seekable file; a ZIP member is not cheaply seekable.
        shutil.copyfileobj(raw, spool, 1024 * 1024)
        spool.seek(0)
        # A workbook found inside a ZIP passed the outer guard as one member;
        # its own sheets are checked here before openpyxl inflates them.
        try:
            guard_archive(zipfile.ZipFile(spool))
        except zipfile.BadZipFile as exc:
            raise ContainerRefused("zip_unreadable", member=member.name) from exc
        spool.seek(0)
        workbook = load_workbook(spool, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                yield sheet.title, (list(row[:_MAX_COLUMNS]) for row in sheet.iter_rows(values_only=True))
        finally:
            workbook.close()


def _json_tables(member: Member) -> Iterator[tuple[str, Iterator[list[Any]]]]:
    if member.size > _MAX_JSON_BYTES:
        raise ValueError("json_too_large")
    with member.open() as raw:
        data = raw.read(_MAX_JSON_BYTES + 1)
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            payload = json.loads(data.decode(encoding))
            break
        except (UnicodeDecodeError, json.JSONDecodeError):
            continue
    else:
        raise ValueError("json_unreadable")
    tables = payload.items() if isinstance(payload, dict) else [("", payload)]
    for title, records in tables:
        if not isinstance(records, list) or not records or not isinstance(records[0], dict):
            continue
        header = list(records[0].keys())

        def _rows(records: list[Any] = records, header: list[str] = header) -> Iterator[list[Any]]:
            yield list(header)
            for rec in records:
                if isinstance(rec, dict):
                    yield [rec.get(h) for h in header]

        yield str(title), _rows()


def iter_tables(member: Member) -> Iterator[tuple[str, Iterator[list[Any]]]]:
    """``(table title, row iterator)`` for every table in the member."""
    ext = member.extension
    if ext == ".xlsx":
        yield from _xlsx_tables(member)
    elif ext == ".json":
        yield from _json_tables(member)
    else:
        yield member.name, _csv_rows(member)


# ── Mapping ──────────────────────────────────────────────────────────────────


def _cell(row: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _text(row: list[Any], idx: int | None) -> str:
    value = _cell(row, idx)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean_text(value)


def _map_header(cells: list[Any], profile: TabularProfile) -> dict[str, int] | None:
    """Column index per role, or ``None`` when the row is not a header row.

    "articolo" means two things: the item text in Puglia and Campania, where
    the group text is "voce", and the group text in Piemonte, where the item
    text is "subarticolo". The presence of a "subarticolo" column decides.
    """
    names = [normalise_header(c) for c in cells]
    aliases = dict(profile.aliases)
    if "subarticolo" in names:
        aliases["group_text"] = ("articolo",)
        aliases["item_text"] = tuple(a for a in aliases["item_text"] if a != "articolo")
    mapping: dict[str, int] = {}
    for idx, name in enumerate(names):
        if profile.prefixed_code_header.match(name):
            mapping["code"] = idx
            break
    for role, role_aliases in aliases.items():
        if role in mapping:
            continue
        for alias in role_aliases:
            if alias in names and names.index(alias) not in mapping.values():
                mapping[role] = names.index(alias)
                break
    found = set(mapping)
    if ("rate" in found or "unit" in found) and ({"code", "description", "item_text"} & found) and len(found) >= 3:
        return mapping
    return None


@dataclass
class _Hierarchy:
    """Headings by code, for lists whose rows carry no chapter columns.

    A heading contains a code when the code continues it after a dot
    ("1.01" contains "1.01.2.a"), or, for codes written without dots, when it
    simply starts with it (Lazio part E: "E01001" contains "E01001a").
    """

    stack: list[tuple[str, str]] = field(default_factory=list)

    @staticmethod
    def contains(parent: str, code: str) -> bool:
        if code.startswith(parent + "."):
            return True
        return "." not in parent and "." not in code and code != parent and code.startswith(parent)

    def ancestors(self, code: str) -> list[tuple[str, str]]:
        while self.stack and not self.contains(self.stack[-1][0], code):
            self.stack.pop()
        return list(self.stack)

    def push(self, code: str, text: str) -> None:
        if code.endswith(".0") and code.count(".") > 0:
            code = code[:-2]  # Umbria: "1.1.20.0" heads the "1.1.20.x" rows
        self.ancestors(code)
        self.stack.append((code, text))


def _clean_code(value: Any) -> str:
    text = clean_text(value).replace(" ", "")
    return text.rstrip(".")


def _row_from_hierarchy(code: str, own: str, hierarchy: _Hierarchy) -> tuple[list[tuple[str, str]], str]:
    ancestors = hierarchy.ancestors(code)

    # The top two levels are chapters, deeper headings are the voce's parent
    # text. A code without dots has no levels to count, so only its outermost
    # heading is a chapter.
    def is_chapter(index: int, heading: str) -> bool:
        return len(heading.split(".")) <= 2 if "." in heading else index == 0

    chapters = [(c, t[:160]) for i, (c, t) in enumerate(ancestors) if is_chapter(i, c)]
    articles = [t for i, (c, t) in enumerate(ancestors) if not is_chapter(i, c)]
    description = " - ".join([*articles, own]) if articles else own
    return chapters, description


def read(member: Member, source: PriceListSource, profile: TabularProfile = DEFAULT_PROFILE) -> Iterator[PriceListRow]:
    """Yield every priced row of a tabular price list."""
    stem = os.path.splitext(member.name)[0]
    _source_from_name(stem, source)
    convention = settle_convention(member, profile)
    for title, rows in iter_tables(member):
        year = re.search(r"\b(20\d{2})\b", title or "")
        if year and not source.edition:
            source.edition = year.group(1)
        yield from _read_table(rows, source, profile, member.name, convention)


def settle_convention(member: Member, profile: TabularProfile = DEFAULT_PROFILE) -> NumberConvention:
    """The decimal separator the member's amount cells write, read before any row is yielded.

    A separate pass over the file, so a "1.250" on the first row is read with
    the evidence of the last one, and the preview and the import (two passes
    each) read every row alike. It stops once enough cells agree; a file that
    disagrees with itself further down is a file that settles nothing anyway.
    Real number cells (a workbook's, a JSON number) prove nothing and need no
    convention, so a table that has written no amount as text for
    :data:`_TEXTLESS_ROWS` rows in a row is not read to its end for one.
    """
    convention = NumberConvention()
    for _title, rows in iter_tables(member):
        head, mapping, _names = _find_header(rows, profile)
        columns = [mapping[role] for role in _AMOUNT_ROLES if role in mapping] if mapping is not None else None
        textless = 0
        for row in _replay(head, rows):
            if columns is None:
                cells = ["" if c is None else str(c) for c in row]
                amounts: list[Any] = [next((c for c in cells if _MONEY_CELL_RE.match(c)), None)]
            else:
                amounts = [_cell(row, idx) for idx in columns]
            texts = [a for a in amounts if isinstance(a, str) and a.strip()]
            textless = 0 if texts else textless + 1
            for text in texts:
                convention.observe(text)
            if convention.settled or textless >= _TEXTLESS_ROWS:
                break
        if convention.settled:
            return convention
    return convention


def _find_header(
    rows: Iterator[list[Any]], profile: TabularProfile
) -> tuple[list[list[Any]], dict[str, int] | None, list[str]]:
    """The rows read while looking for the header, its role mapping and its normalised names."""
    head: list[list[Any]] = []
    for row in rows:
        head.append(row)
        mapping = _map_header(row, profile)
        if mapping is not None:
            return head, mapping, [normalise_header(c) for c in row]
        if len(head) >= _HEADER_SCAN_ROWS:
            break
    return head, None, []


def _source_from_name(stem: str, source: PriceListSource) -> None:
    lowered = normalise_header(stem)
    if not source.edition:
        year = re.search(r"\b(20\d{2})\b", lowered)
        if year:
            source.edition = year.group(1)
    if not source.region_code:
        for code, name in REGION_NAMES.items():
            if normalise_header(name) in lowered.split() or normalise_header(name) == lowered:
                source.region_code = code
                source.detected_from = "filename"
                break


def _read_table(
    rows: Iterator[list[Any]],
    source: PriceListSource,
    profile: TabularProfile,
    file_name: str,
    convention: NumberConvention,
) -> Iterator[PriceListRow]:
    head, mapping, header_names = _find_header(rows, profile)
    if mapping is None:
        yield from _read_headerless(_replay(head, rows), source, file_name, convention)
        return
    source.profile = source.profile or "headed"
    if not source.region_code:
        names = set(header_names)
        for marker, region in _LAYOUT_REGIONS:
            if marker <= names:
                source.region_code = region
                source.detected_from = "layout"
                break
    chapter_cols = [
        (header_names.index(f"codice {c}") if f"codice {c}" in header_names else None, header_names.index(c))
        for c in profile.chapter_columns
        if c in header_names
    ]
    hierarchy = _Hierarchy()
    for row in rows:
        if not any(v not in (None, "") for v in row):
            continue
        yield from _headed_row(row, mapping, chapter_cols, hierarchy, source, file_name, convention)


def _replay(head: list[list[Any]], rest: Iterator[list[Any]]) -> Iterator[list[Any]]:
    yield from head
    yield from rest


def _headed_row(
    row: list[Any],
    mapping: dict[str, int],
    chapter_cols: list[tuple[int | None, int]],
    hierarchy: _Hierarchy,
    source: PriceListSource,
    file_name: str,
    convention: NumberConvention,
) -> Iterator[PriceListRow]:
    flags: list[str] = []
    code = _clean_code(_cell(row, mapping.get("code")))
    description = _text(row, mapping.get("description"))
    group = _text(row, mapping.get("group_text"))
    item = _text(row, mapping.get("item_text"))
    raw_rate = _cell(row, mapping.get("rate"))
    source_unit = _text(row, mapping.get("unit"))
    if source.edition is None and mapping.get("edition") is not None:
        edition = _text(row, mapping.get("edition"))
        source.edition = edition or None
    if raw_rate in (None, "") and not source_unit:
        if code:
            hierarchy.push(code, description or group or item)
        return
    rate = amount_or_flag(raw_rate, "rate", flags, convention)
    own = description or item
    if group and own and own not in group and not own.startswith(group):
        long_text = f"{group} - {own}"
    else:
        # Campania repeats the voce text at the head of the articolo text.
        long_text = own if own.startswith(group) else group or own
    if chapter_cols:
        chapters = [
            (_text(row, code_idx) if code_idx is not None else "", _text(row, name_idx)[:160])
            for code_idx, name_idx in chapter_cols
            if _text(row, name_idx)
        ]
    else:
        chapters, long_text = _row_from_hierarchy(code, long_text, hierarchy)
    if code and source.detected_from in (None, "layout", "filename"):
        # A region written in the codes outranks one guessed from the layout or the file name.
        found = region_from_code(code)
        if found:
            source.region_code, year = found
            source.edition = source.edition or year
            source.detected_from = "code_prefix"
    if code and source.edition is None:
        found = region_from_code(code)
        if found:
            source.edition = found[1]
    out = PriceListRow(
        code=code,
        description=long_text,
        short_description=item or own or long_text[:160],
        unit=normalise_unit(source_unit),
        source_unit=source_unit,
        rate=rate,
        chapters=chapters,
        labour_share_pct=amount_or_flag(_cell(row, mapping.get("labour_pct")), "labour_share", flags, convention),
        labour_amount=amount_or_flag(_cell(row, mapping.get("labour_amount")), "labour_amount", flags, convention),
        safety_share_pct=amount_or_flag(_cell(row, mapping.get("safety_pct")), "safety_share", flags, convention),
        safety_amount=amount_or_flag(_cell(row, mapping.get("safety_amount")), "safety_amount", flags, convention),
        flags=flags,
    )
    out.extra["source_file"] = file_name
    _resource_split(out, row, mapping, convention)
    yield out


def _resource_split(out: PriceListRow, row: list[Any], mapping: dict[str, int], convention: NumberConvention) -> None:
    """Puglia states labour, plant and material totals; they become components."""
    if out.rate is None or not {"labour_amount", "material_amount", "equipment_amount"} <= set(mapping):
        return
    parts: list[tuple[str, str, Decimal]] = []
    for role, kind, name in (
        ("labour_amount", "labor", "Manodopera"),
        ("equipment_amount", "equipment", "Attrezzature e noli"),
        ("material_amount", "material", "Materiali"),
    ):
        try:
            value = parse_amount(_cell(row, mapping[role]), convention=convention)
        except BrokenNumber:
            return
        if value:
            parts.append((kind, name, value))
    total = sum((v for _k, _n, v in parts), Decimal(0))
    balance = out.rate - total
    if not parts or balance < 0:
        return
    if balance > 0:
        parts.append(("other", "Spese generali e utile d'impresa", balance))
    out.components = [
        {
            "code": "",
            "name": name,
            "unit": out.unit,
            "quantity": 1.0,
            "unit_rate": float(value),
            "cost": float(value),
            "type": kind,
        }
        for kind, name, value in parts
    ]


def _read_headerless(
    rows: Iterator[list[Any]], source: PriceListSource, file_name: str, convention: NumberConvention
) -> Iterator[PriceListRow]:
    """A table with no header row: columns found around the price cell (Lazio)."""
    source.profile = source.profile or "headerless"
    hierarchy = _Hierarchy()
    for row in rows:
        cells = ["" if c is None else str(c) for c in row]
        price_idx = next((i for i, c in enumerate(cells) if _MONEY_CELL_RE.match(c)), None)
        if price_idx is None:
            code, text = _headerless_heading(cells)
            if code and text:
                hierarchy.push(code, clean_text(text))
            continue
        flags: list[str] = []
        code_idx = price_idx - 3
        if code_idx < 0:
            continue
        part = clean_text(cells[code_idx - 1]) if code_idx >= 1 else ""
        code = _clean_code(cells[code_idx])
        own = clean_text(cells[price_idx - 2])
        source_unit = clean_text(cells[price_idx - 1])
        chapters, description = _row_from_hierarchy(code, own, hierarchy)
        full_code = f"{part}{code}" if len(part) == 1 and part.isalpha() else code
        out = PriceListRow(
            code=full_code,
            description=description,
            short_description=own,
            unit=normalise_unit(source_unit),
            source_unit=source_unit,
            rate=amount_or_flag(cells[price_idx], "rate", flags, convention),
            chapters=chapters,
            flags=flags,
        )
        out.extra["source_file"] = file_name
        if len(part) == 1 and part.isalpha():
            out.extra["part"] = part
            out.chapters.insert(0, (part, f"Parte {part}"))
        yield out


def _headerless_heading(cells: list[str]) -> tuple[str, str]:
    filled = [(i, c) for i, c in enumerate(cells) if c.strip()]
    if len(filled) == 1:
        # "E01. IMPIANTI IDRO-SANITARI E GAS DOMESTICO" in a single cell.
        match = _NUMBERED_HEADING_RE.match(filled[0][1].strip())
        return (match.group(1), match.group(2)) if match else ("", "")
    if len(filled) < 2:
        return "", ""
    first_idx, first = filled[0]
    if len(first.strip()) == 1 and first.strip().isalpha() and len(filled) >= 3:
        return _clean_code(filled[1][1]), filled[2][1]
    if len(first.strip()) == 1 and first.strip().isalpha():
        return "", ""
    return _clean_code(first), filled[1][1]


__all__ = [
    "DEFAULT_PROFILE",
    "FORMAT_ID",
    "TabularProfile",
    "iter_tables",
    "normalise_header",
    "read",
    "settle_convention",
    "sniff",
]
