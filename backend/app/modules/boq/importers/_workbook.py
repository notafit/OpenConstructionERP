"""Open an .xlsx or an Excel 97-2003 .xls workbook behind one reading surface.

The spreadsheet readers (the generic header mapper in ``excel.py`` and the
Hungarian profiles in ``hungary_workbook.py``) were written against the part of
openpyxl's read-only workbook they use: ``sheetnames``, ``workbook[name]``,
``active``, ``close()``, and on a sheet ``title``, ``sheet_state`` and
``iter_rows(min_row=, max_row=, values_only=True)``. openpyxl reads OOXML only,
and a great many bills still travel as .xls because the estimating programs
that write them export that format. :func:`open_workbook` returns openpyxl's
own workbook for an .xlsx and an adapter over xlrd for an .xls, so both go
through the same readers, the same profiles and the same validation, and
neither reader has to know which one it holds.

The adapter hands back cell values the way openpyxl would for the same cell,
because the readers were tuned on openpyxl's values:

* an empty cell is ``None``, not xlrd's ``""``;
* a whole number is an ``int``. The .xls format stores every number as a
  float, and a chapter code or a tetelrend segment read as ``11.0`` would no
  longer be the code ``11``;
* a date is a ``datetime``, decoded with the workbook's own date system;
* an error cell is its display text (``#DIV/0!``), which is what openpyxl
  returns with ``data_only=True``.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

from app.core.file_signature import detect as detect_signature
from app.modules.boq.importers._base import ImporterParseError

# The directory entry an encrypted OOXML package is stored under, in the
# UTF-16 the compound file directory spells it in. A password-protected
# .xlsx is not a zip but an OLE2 container holding this stream, so it looks
# like an .xls until someone tries to read a workbook out of it.
_ENCRYPTED_PACKAGE = "EncryptedPackage".encode("utf-16-le")

_SHEET_STATES = {0: "visible", 1: "hidden", 2: "veryHidden"}


def is_legacy_workbook(head: bytes) -> bool:
    """True when the bytes start an OLE2 compound file, the .xls container."""
    return detect_signature(head) == "ole"


def open_workbook(content: bytes) -> Any:
    """Open workbook bytes for reading, whichever of the two formats they are in.

    Raises:
        ImporterParseError: for a compound file that holds no readable
            workbook: a password-protected workbook, or another Office file
            (a .doc, a .msg) given a spreadsheet's name.
        Exception: whatever openpyxl raises on a broken .xlsx, unchanged, so
            the callers keep reporting those as they always did.
    """
    if is_legacy_workbook(content[:8]):
        return _open_legacy(content)

    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content), read_only=True, data_only=True)


def _open_legacy(content: bytes) -> LegacyWorkbook:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content, on_demand=False)
    except xlrd.XLRDError as exc:
        text = str(exc)
        # Asked only once xlrd has found no workbook: the bytes of a valid .xls
        # can spell anything in a cell, the stream name included.
        if "encrypted" in text.lower() or _ENCRYPTED_PACKAGE in content:
            raise ImporterParseError(
                "This workbook is protected with a password. Open it, remove the password and save it again.",
                code="workbook_password_protected",
            ) from exc
        raise ImporterParseError(
            "This file is an Office document but holds no Excel workbook that can be read: " + text,
            code="workbook_not_found",
        ) from exc
    except Exception as exc:  # noqa: BLE001 - xlrd raises plain errors on truncated files
        raise ImporterParseError(
            f"Could not read the Excel 97-2003 workbook: {exc}", code="workbook_unreadable"
        ) from exc
    return LegacyWorkbook(book)


class LegacySheet:
    """One .xls worksheet, read the way openpyxl's read-only worksheet reads."""

    def __init__(self, sheet: Any, datemode: int) -> None:
        self._sheet = sheet
        self._datemode = datemode
        self.title: str = sheet.name
        self.sheet_state: str = _SHEET_STATES.get(sheet.visibility, "hidden")

    def iter_rows(
        self,
        min_row: int | None = None,
        max_row: int | None = None,
        values_only: bool = False,
    ) -> Iterator[tuple[Any, ...]]:
        """Yield each row as a tuple of values, rows numbered from 1 like openpyxl."""
        if not values_only:
            raise NotImplementedError("Only values_only=True is supported for .xls sheets")
        first = max((min_row or 1) - 1, 0)
        last = self._sheet.nrows if max_row is None else min(max_row, self._sheet.nrows)
        for index in range(first, last):
            yield tuple(self._value(cell) for cell in self._sheet.row(index))

    def _value(self, cell: Any) -> Any:
        import xlrd

        kind = cell.ctype
        if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return None
        if kind == xlrd.XL_CELL_TEXT:
            return cell.value if cell.value != "" else None
        if kind == xlrd.XL_CELL_NUMBER:
            value = float(cell.value)
            return int(value) if value.is_integer() else value
        if kind == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate.xldate_as_datetime(cell.value, self._datemode)
            except (xlrd.xldate.XLDateError, ValueError, OverflowError):
                return float(cell.value)
        if kind == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if kind == xlrd.XL_CELL_ERROR:
            return xlrd.error_text_from_code.get(cell.value, "#N/A")
        return cell.value


class LegacyWorkbook:
    """An .xls workbook with the part of openpyxl's workbook the readers use."""

    def __init__(self, book: Any) -> None:
        self._book = book
        self._sheets = [LegacySheet(sheet, book.datemode) for sheet in book.sheets()]

    @property
    def sheetnames(self) -> list[str]:
        return [sheet.title for sheet in self._sheets]

    def __getitem__(self, name: str) -> LegacySheet:
        for sheet in self._sheets:
            if sheet.title == name:
                return sheet
        raise KeyError(f"Worksheet {name} does not exist.")

    @property
    def active(self) -> LegacySheet | None:
        """The sheet the workbook opens on, or the first one."""
        if not self._sheets:
            return None
        for sheet in self._sheets:
            if getattr(sheet._sheet, "sheet_visible", 0):
                return sheet
        return self._sheets[0]

    def close(self) -> None:
        self._book.release_resources()


def is_legacy(workbook: Any) -> bool:
    """True for a workbook :func:`open_workbook` opened from .xls bytes."""
    return isinstance(workbook, LegacyWorkbook)


__all__ = [
    "LegacySheet",
    "LegacyWorkbook",
    "is_legacy",
    "is_legacy_workbook",
    "open_workbook",
]
