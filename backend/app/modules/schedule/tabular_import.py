# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Read a schedule from an Excel workbook or a CSV file (pure parsing core).

A planner's schedule often lives in a spreadsheet: one row per activity with an
id, a name, dates or a duration and a predecessor list such as ``"A10FS+2d;
A20SS"``. :func:`preview` reads such a file and returns everything an import
dialog needs before anything is written:

* the column mapping it detected, each column with a confidence
  (1.0 exact header, 0.8 synonym, 0.5 fuzzy, 0 none) and every collision;
* sample rows, the parsed activities and the dependency network;
* every problem as an :class:`Issue` with a stable ``code`` the frontend
  translates, the 1-based sheet row and the 0-based column index;
* the SHA-256 of the bytes, so a later commit can prove it re-reads the same file.

The parsed result is an ``oce-schedule-interchange`` document
(:mod:`app.modules.schedule.schedule_interchange`), so committing it is a call
to ``ScheduleInterchangeService.import_schedule``. Each activity additionally
carries its source row in ``metadata["import_row"]`` and ``client_visible``,
which is always ``False`` here: the rows the sheet marks for the client are only
listed in ``client_visible_suggested``, and the commit applies the list a person
confirmed.

The module is pure and synchronous: no database, no ORM, no network. Callers
run it in a worker thread. Durations are working days counted inclusively, the
way ``schedule.service.compute_duration`` counts them; the caller passes the
project's working week and holidays (Monday to Friday when omitted).
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import zipfile
from collections import deque
from collections.abc import Iterable, Iterator, Mapping
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any, Literal

from app.core.csv_dialect import sep_directive, sniff_delimiter
from app.core.file_signature import detect as detect_signature
from app.core.sheet_header import find_header_row
from app.modules.boq.importers._encoding import decode_text_bytes, parse_numeric_cell
from app.modules.schedule.schedule_interchange import FORMAT, FORMAT_VERSION
from app.modules.schedule.tabular_headers import (
    CONFIDENCE_EXACT,
    CONFIDENCE_NONE,
    FIELDS,
    flag_value,
    known_header,
    match_header,
    month_number,
    relationship_type,
    unit_kind,
)
from app.modules.schedule.xer_encoding import sniff_code_page

# ── Limits ────────────────────────────────────────────────────────────────────

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 60
MAX_CELL_CHARS = 2000
#: Longest duration or lag read, in working days (about 40 years).
MAX_SPAN_DAYS = 10000
#: Column widths of the rows an import writes (``schedule.models.Activity`` and
#: ``schemas.ActivityResource``); a longer value is refused in the preview.
_FIELD_LIMITS: dict[str, int] = {"name": 255, "id": 50, "wbs": 50, "resource": 255}
_SCHEDULE_NAME_LIMIT = 255
#: Issues reported per preview, and per issue code; the rest are counted by ``issues_truncated``.
MAX_ISSUES = 1000
MAX_ISSUES_PER_CODE = 50
#: Activity rows returned as ``sample_rows``.
SAMPLE_ROWS = 10
#: Hours in one working day, for durations and lags given in hours.
HOURS_PER_DAY = 8
#: A run of this many empty rows ends the sheet (a formatted but empty tail).
_BLANK_RUN_END = 500
#: Cycles reported before the loop search stops.
_MAX_CYCLES = 20

DateOrder = Literal["dmy", "mdy"]
Severity = Literal["error", "warning", "info"]

_DEFAULT_WEEK: frozenset[int] = frozenset({0, 1, 2, 3, 4})
_SPREADSHEET_EXTENSIONS = (".xlsx", ".xlsm", ".xls")


# ── Result types ──────────────────────────────────────────────────────────────


@dataclass
class Issue:
    """One problem found in the file, addressed by sheet row and column index."""

    code: str
    severity: Severity
    message: str
    row: int | None = None
    column: int | None = None
    params: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "severity": self.severity,
            "row": self.row,
            "column": self.column,
            "message": self.message,
            "params": dict(self.params),
        }


@dataclass
class ColumnMatch:
    """How one sheet column was read."""

    index: int
    header: str
    field: str | None
    confidence: float
    tier: str  # "exact" | "synonym" | "fuzzy" | "override" | "none"

    def to_dict(self) -> dict[str, Any]:
        return {
            "index": self.index,
            "header": self.header,
            "field": self.field,
            "confidence": self.confidence,
            "tier": self.tier,
        }


@dataclass
class TabularPreview:
    """Everything :func:`preview` learned about one uploaded file."""

    sha256: str
    filename: str
    file_format: str | None = None  # "xlsx" | "xls" | "csv"
    encoding: str | None = None
    delimiter: str | None = None
    sheet: str | None = None
    header_row: int | None = None
    columns: list[ColumnMatch] = field(default_factory=list)
    mapping: dict[str, int] = field(default_factory=dict)
    date_order: DateOrder | None = None
    date_order_source: str | None = None  # "explicit" | "values" | "suggested" | None
    outline_source: str | None = None  # "outline_level" | "wbs" | "indent" | "leading_spaces" | None
    row_count: int = 0
    sample_rows: list[dict[str, Any]] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    document: dict[str, Any] | None = None
    #: Refs of the activities the sheet marks for the client. A suggestion only:
    #: every activity in ``document`` stays hidden until a commit names it.
    client_visible_suggested: list[str] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        """True when an error blocks the import."""
        return any(issue.severity == "error" for issue in self.issues)

    @property
    def activity_count(self) -> int:
        return len(self.document["activities"]) if self.document else 0

    @property
    def relationship_count(self) -> int:
        return len(self.document["relationships"]) if self.document else 0

    def issue_codes(self) -> list[str]:
        return [issue.code for issue in self.issues]

    def to_dict(self) -> dict[str, Any]:
        return {
            "sha256": self.sha256,
            "filename": self.filename,
            "file_format": self.file_format,
            "encoding": self.encoding,
            "delimiter": self.delimiter,
            "sheet": self.sheet,
            "header_row": self.header_row,
            "columns": [column.to_dict() for column in self.columns],
            "mapping": dict(self.mapping),
            "date_order": self.date_order,
            "date_order_source": self.date_order_source,
            "outline_source": self.outline_source,
            "row_count": self.row_count,
            "activity_count": self.activity_count,
            "relationship_count": self.relationship_count,
            "has_errors": self.has_errors,
            "sample_rows": [dict(row) for row in self.sample_rows],
            "issues": [issue.to_dict() for issue in self.issues],
            "client_visible_suggested": list(self.client_visible_suggested),
            "document": self.document,
        }


class _Issues:
    """Issue collector with a cap per code, so one noisy problem cannot hide the others."""

    def __init__(self) -> None:
        self.items: list[Issue] = []
        self.per_code: dict[str, int] = {}
        self.dropped: dict[str, int] = {}

    def add(
        self,
        code: str,
        severity: Severity,
        message: str,
        *,
        row: int | None = None,
        column: int | None = None,
        **params: Any,
    ) -> None:
        seen = self.per_code.get(code, 0)
        self.per_code[code] = seen + 1
        if seen >= MAX_ISSUES_PER_CODE or len(self.items) >= MAX_ISSUES:
            self.dropped[code] = self.dropped.get(code, 0) + 1
            return
        self.items.append(Issue(code, severity, message, row, column, params))

    def finish(self) -> list[Issue]:
        if self.dropped:
            total = sum(self.dropped.values())
            self.items.append(
                Issue(
                    "issues_truncated",
                    "info",
                    f"{total} more issues were not listed",
                    params={"count": total, "by_code": dict(self.dropped)},
                )
            )
        return self.items


# ── Reading the file into rows ────────────────────────────────────────────────


@dataclass
class _Sheet:
    """Rows of one sheet as plain values, with the cell indent where the format has one."""

    name: str | None
    rows: list[tuple[Any, ...]]
    indents: list[tuple[float, ...]] | None
    truncated: bool = False


def _reader_errors() -> tuple[type[BaseException], ...]:
    """Exceptions the workbook readers raise on a file they cannot read."""
    from openpyxl.utils.exceptions import InvalidFileException

    from app.modules.boq.importers._base import ImporterParseError

    legacy: tuple[type[BaseException], ...] = ()
    try:
        from xlrd import XLRDError
        from xlrd.compdoc import CompDocError

        legacy = (XLRDError, CompDocError)
    except ImportError:
        # Without xlrd an .xls cannot be opened at all; that ImportError is
        # reported as an unreadable file below.
        pass

    return (
        *legacy,
        ImporterParseError,
        InvalidFileException,
        ImportError,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        OSError,
        EOFError,
    )


def _trim(row: Iterable[Any]) -> tuple[Any, ...]:
    """The row without its trailing empty cells."""
    values = list(row)
    while values and (values[-1] is None or (isinstance(values[-1], str) and not values[-1].strip())):
        values.pop()
    return tuple(values)


def _row_cap() -> int:
    from app.core.sheet_header import HEADER_SEARCH_ROWS

    return MAX_ROWS + HEADER_SEARCH_ROWS + 1


def _collect(rows: Iterator[tuple[tuple[Any, ...], tuple[float, ...] | None]]) -> tuple[list, list, bool]:
    """Read rows until the row cap or a long blank run; returns ``(values, indents, truncated)``."""
    values: list[tuple[Any, ...]] = []
    indents: list[tuple[float, ...]] = []
    filled = 0
    blank_run = 0
    cap = _row_cap()
    for row_values, row_indents in rows:
        trimmed = _trim(row_values)
        if not trimmed:
            blank_run += 1
            if blank_run >= _BLANK_RUN_END:
                break
        else:
            blank_run = 0
            filled += 1
            if filled > cap:
                return values, indents, True
        values.append(trimmed)
        indents.append(row_indents or ())
    while values and not values[-1]:
        values.pop()
        indents.pop()
    return values, indents, False


def _xlsx_rows(worksheet: Any) -> Iterator[tuple[tuple[Any, ...], tuple[float, ...]]]:
    for cells in worksheet.iter_rows(max_col=MAX_COLUMNS + 1):
        row_values = tuple(getattr(cell, "value", None) for cell in cells)
        row_indents = []
        for cell in cells:
            alignment = getattr(cell, "alignment", None)
            row_indents.append(float(getattr(alignment, "indent", 0) or 0) if alignment is not None else 0.0)
        yield row_values, tuple(row_indents)


def _xls_rows(worksheet: Any) -> Iterator[tuple[tuple[Any, ...], None]]:
    for row_values in worksheet.iter_rows(values_only=True):
        yield tuple(row_values)[: MAX_COLUMNS + 1], None


def _read_workbook(data: bytes, file_format: str) -> list[_Sheet]:
    """Every visible sheet of the workbook, read up to the row cap."""
    from app.modules.boq.importers._workbook import open_workbook

    workbook = open_workbook(data)
    sheets: list[_Sheet] = []
    try:
        for name in workbook.sheetnames:
            worksheet = workbook[name]
            if getattr(worksheet, "sheet_state", "visible") != "visible":
                continue
            source = _xlsx_rows(worksheet) if file_format == "xlsx" else _xls_rows(worksheet)
            values, indents, truncated = _collect(source)
            sheets.append(_Sheet(name, values, indents if file_format == "xlsx" else None, truncated))
    finally:
        close = getattr(workbook, "close", None)
        if close is not None:
            close()
    return sheets


#: Single-byte pages a CSV saved by Excel outside Western Europe comes in. Cyrillic
#: (Russian, Ukrainian, Bulgarian, Kazakh), Central European (Polish, Czech) and Turkish.
_LEGACY_PAGES = ("cp1251", "cp1250", "cp1254")
_HIGH_RUN = re.compile(rb"[\x80-\xff]+")
#: Rows searched for header words when telling code pages apart.
_ENCODING_HEADER_ROWS = 15
#: Characters of each reading weighed; the rest of a large file adds time, not evidence.
_ENCODING_SAMPLE_CHARS = 65536


def _header_hits(text: str, delimiter: str) -> int:
    """The most cells of one early row that read as a known header."""
    best = 0
    lines = text.splitlines()[:_ENCODING_HEADER_ROWS]
    for row in csv.reader(lines, delimiter=delimiter):
        best = max(best, sum(1 for cell in row[:MAX_COLUMNS] if known_header(cell)))
    return best


def _symbols_inside_words(text: str) -> int:
    """Non-ASCII symbols with a letter on both sides, which real words do not contain.

    A symbol next to a digit or a space is left alone: ``m³``, ``5 €`` and
    ``2–3`` are ordinary text in any page.
    """
    return sum(
        1
        for index in range(1, len(text) - 1)
        if ord(text[index]) > 0x7F
        and not text[index].isalpha()
        and text[index - 1].isalpha()
        and text[index + 1].isalpha()
    )


def _decode_csv(data: bytes) -> tuple[str, str]:
    """Decode CSV bytes, telling a Cyrillic, Central European or Turkish code page from cp1252.

    UTF-8 and a byte-order mark are taken as they come. Bytes that only a
    single-byte page reads are decoded with each candidate, and the reading
    kept is the one that (1) spells the most header words, then (2) produces the
    fewest non-letter characters above ASCII (Polish ``ą`` read as cp1252 is a
    superscript ``¹``, the Russian letter che a division sign), then (3) agrees with the
    letter-frequency vote of :func:`~app.modules.schedule.xer_encoding.sniff_code_page`,
    or failing that is Cyrillic when the high bytes come in whole words rather
    than one accented letter at a time, and otherwise stays cp1252. Turkish
    cp1254 differs from cp1252 only in letters, so it is chosen on header
    words alone.
    """
    text, encoding = decode_text_bytes(data)
    if encoding not in ("cp1252", "latin-1"):
        return text, encoding
    sniffed = sniff_code_page(data[:_ENCODING_SAMPLE_CHARS])
    runs = _HIGH_RUN.findall(data[:_ENCODING_SAMPLE_CHARS])
    high = sum(len(run) for run in runs)
    in_words = sum(len(run) for run in runs if len(run) >= 3) / high if high else 0.0
    delimiter = sniff_delimiter(text)

    def prior(page: str) -> int:
        if page == sniffed:
            return 4
        if page == "cp1251" and sniffed is None and in_words >= 0.5:
            return 3
        if page == encoding:
            return 2
        # One accented letter at a time is how Central European text looks, too.
        return 1 if page == "cp1250" and in_words < 0.5 else 0

    best: tuple[tuple[int, int, int], str, str] | None = None
    for page in dict.fromkeys((encoding, *([sniffed] if sniffed else []), *_LEGACY_PAGES)):
        try:
            candidate = text if page == encoding else data.decode(page)
        except (UnicodeDecodeError, LookupError):
            continue
        head = candidate[:_ENCODING_SAMPLE_CHARS]
        score = (_header_hits(head, delimiter), -_symbols_inside_words(head), prior(page))
        if best is None or score > best[0]:
            best = (score, candidate, page)
    return (best[1], best[2]) if best else (text, encoding)


def _read_csv(data: bytes) -> tuple[_Sheet, str, str]:
    """The CSV as one sheet, with the encoding and the delimiter it was read with."""
    text, encoding = _decode_csv(data)
    _named, blanked = sep_directive(text)
    delimiter = sniff_delimiter(text)
    reader = csv.reader(io.StringIO(blanked, newline=""), delimiter=delimiter)
    values, indents, truncated = _collect((tuple(row), None) for row in reader)
    return _Sheet(None, values, None, truncated), encoding, delimiter


# ── Cell helpers ──────────────────────────────────────────────────────────────

_FORMULA_ESCAPED = re.compile(r"^'[=+\-@\t\r]")


def _text(value: Any) -> str:
    """A cell as display text: whole floats without ``.0``, dates as ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value)
    # The schedule CSV export guards formula-looking cells with a leading
    # apostrophe (csv_safety.neutralise_formula); read them back without it.
    if _FORMULA_ESCAPED.match(text):
        text = text[1:]
    return text


def _round_half_up(value: float) -> int:
    return int(math.floor(abs(value) + 0.5)) * (1 if value >= 0 else -1)


# ── Working-day arithmetic ────────────────────────────────────────────────────


@dataclass(frozen=True)
class _Week:
    weekdays: frozenset[int]
    holidays: frozenset[date]

    def working(self, day: date) -> bool:
        return day.weekday() in self.weekdays and day not in self.holidays

    def count(self, start: date, end: date) -> int:
        """Working days from ``start`` to ``end`` inclusive, 0 when ``end`` precedes ``start``."""
        if end < start:
            return 0
        span = (end - start).days + 1
        if span > 20000:
            span = 20000
        return sum(1 for offset in range(span) if self.working(start + timedelta(days=offset)))

    def finish_from(self, start: date, days: int) -> date:
        """The finish date of a ``days``-long activity starting ``start`` (``count`` inverts it)."""
        if days <= 0:
            return start
        current = start
        guard = 0
        while not self.working(current) and guard < 400:
            current += timedelta(days=1)
            guard += 1
        remaining = days - 1
        while remaining > 0 and guard < 40000:
            current += timedelta(days=1)
            guard += 1
            if self.working(current):
                remaining -= 1
        return current

    def start_from(self, finish: date, days: int) -> date:
        """The start date of a ``days``-long activity finishing ``finish``."""
        if days <= 0:
            return finish
        current = finish
        guard = 0
        while not self.working(current) and guard < 400:
            current -= timedelta(days=1)
            guard += 1
        remaining = days - 1
        while remaining > 0 and guard < 40000:
            current -= timedelta(days=1)
            guard += 1
            if self.working(current):
                remaining -= 1
        return current


# ── Dates ─────────────────────────────────────────────────────────────────────

_ISO_DATE = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[T\s].*)?$")
_NUMERIC_DATE = re.compile(r"^(\d{1,2})([./-])(\d{1,2})\2(\d{4}|\d{2})(?:[T\s,].*)?$")
# Month names in any language of ``tabular_headers.MONTH_NAMES``: "14-Mar-26",
# "1. März 2026", "1er mars 2026", "12 de marzo de 2026", "4 maja 2026 r.",
# "March 1st, 2026", and a Russian genitive with the year mark glued on.
_DAY_MONTH_NAME = re.compile(
    r"^(\d{1,2})(?:st|nd|rd|th|er|º|°)?(?:\.|\s+de)?[\s.-]+([^\W\d_]{3,12})\.?(?:\s+de)?[\s.,-]+'?(\d{4}|\d{2})"
    r"(?:\s*г\.?)?(?:\s.*)?$"
)
_MONTH_NAME_DAY = re.compile(r"^([^\W\d_]{3,12})\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+'?(\d{4}|\d{2})(?:\s.*)?$")
# A leading weekday ("Mon", "Montag,", "lundi", a Russian two-letter one), dropped only after the
# month patterns failed on the whole cell.
_WEEKDAY_PREFIX = re.compile(r"^[^\W\d_]{2,12}\.?,?\s+(?=\d)")
#: Excel serial day numbers that plausibly are dates (1954 to 2119).
_SERIAL_RANGE = (20000, 80000)
_EXCEL_EPOCH = date(1899, 12, 30)


def _year(text: str) -> int:
    year = int(text)
    if len(text) == 2:
        year += 2000 if year < 70 else 1900
    return year


@dataclass
class _DateCell:
    """A date cell before the day/month order is known."""

    row: int
    column: int
    raw: str
    value: date | None = None  # set when the cell was unambiguous on its own
    first: int = 0  # the two leading numbers of an a/b/YYYY date
    second: int = 0
    year: int = 0
    invalid: bool = False
    serial: bool = False


def _read_date(value: Any, row: int, column: int) -> _DateCell | None:
    """Classify a date cell; ``None`` when it is empty."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return _DateCell(row, column, _text(value), value=value.date())
    if isinstance(value, date):
        return _DateCell(row, column, _text(value), value=value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if _SERIAL_RANGE[0] <= value <= _SERIAL_RANGE[1]:
            return _DateCell(row, column, _text(value), value=_EXCEL_EPOCH + timedelta(days=int(value)), serial=True)
        return _DateCell(row, column, _text(value), invalid=True)
    raw = _text(value).strip()
    text = raw
    for _ in range(2):
        iso = _ISO_DATE.match(text)
        if iso:
            try:
                return _DateCell(row, column, raw, value=date(int(iso[1]), int(iso[2]), int(iso[3])))
            except ValueError:
                return _DateCell(row, column, raw, invalid=True)
        numeric = _NUMERIC_DATE.match(text)
        if numeric:
            return _DateCell(row, column, raw, first=int(numeric[1]), second=int(numeric[3]), year=_year(numeric[4]))
        named = _DAY_MONTH_NAME.match(text)
        month = month_number(named[2]) if named else None
        if named and month:
            return _named_date(raw, row, column, int(named[1]), month, named[3])
        named = _MONTH_NAME_DAY.match(text)
        month = month_number(named[1]) if named else None
        if named and month:
            return _named_date(raw, row, column, int(named[2]), month, named[3])
        stripped = _WEEKDAY_PREFIX.sub("", text, count=1)
        if stripped == text:
            break
        text = stripped
    return _DateCell(row, column, raw, invalid=True)


def _named_date(raw: str, row: int, column: int, day: int, month: int, year: str) -> _DateCell:
    try:
        return _DateCell(row, column, raw, value=date(_year(year), month, day))
    except ValueError:
        return _DateCell(row, column, raw, invalid=True)


def _resolve(cell: _DateCell, order: DateOrder) -> date | None:
    if cell.invalid:
        return None
    if cell.value is not None:
        return cell.value
    day, month = (cell.first, cell.second) if order == "dmy" else (cell.second, cell.first)
    try:
        return date(cell.year, month, day)
    except ValueError:
        return None


# ── Durations and lags ────────────────────────────────────────────────────────

_AMOUNT = re.compile(r"^([+-]?\d+(?:[.,]\d+)?)\s*(.*?)\s*\??$")


@dataclass
class _Amount:
    """A duration or lag read into working days."""

    days: int | None = None
    error: str | None = None  # issue code
    warnings: list[tuple[str, dict[str, Any]]] = field(default_factory=list)


def _to_days(number: float, unit: str | None, week_length: int, *, allow_negative: bool) -> _Amount:
    """Convert ``number`` in ``unit`` (``None`` means days) to whole working days."""
    kind = "day" if not unit else unit_kind(unit)
    if kind is None or kind == "unsupported":
        return _Amount(error="unit_unsupported")
    if kind == "elapsed":
        return _Amount(error="elapsed_unsupported")
    if number < 0 and not allow_negative:
        return _Amount(error="negative")
    out = _Amount()
    days = float(number)
    if kind == "week":
        days *= week_length
    elif kind == "hour":
        days /= HOURS_PER_DAY
        out.warnings.append(("hours_converted", {"hours": number, "days": days}))
    if abs(days) > MAX_SPAN_DAYS:
        return _Amount(error="invalid")
    if not days.is_integer():
        rounded = _round_half_up(days)
        if rounded == 0 and days != 0:
            rounded = 1 if days > 0 else -1
        out.warnings.append(("rounded", {"exact_days": days, "days": rounded}))
        out.days = rounded
    else:
        out.days = int(days)
    return out


def _read_duration(value: Any, week_length: int, unit_hint: str | None) -> _Amount | None:
    """A duration cell in working days; ``None`` when empty."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return _Amount(error="invalid")
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            return _Amount(error="invalid")
        return _to_days(float(value), unit_hint, week_length, allow_negative=False)
    text = _text(value).strip()
    match = _AMOUNT.match(text)
    if not match:
        return _Amount(error="invalid")
    number, problem = parse_numeric_cell(match[1])
    if problem is not None or number is None:
        return _Amount(error="invalid")
    unit = match[2].strip() or unit_hint
    return _to_days(number, unit, week_length, allow_negative=False)


_UNIT_HINT = re.compile(r"[(\[]([^)\]]+)[)\]]")


def _header_unit(header: str) -> str | None:
    """The unit a header names in brackets, ``"Dauer [Std]"`` -> ``"Std"``, when it is one."""
    for found in _UNIT_HINT.findall(header):
        if unit_kind(found) in ("day", "week", "hour", "elapsed", "unsupported"):
            return found
    return None


# ── Predecessors ──────────────────────────────────────────────────────────────

_TOKEN_SPLIT = re.compile(r"[;,\n]")
_LAG = re.compile(r"^(?P<head>.*?)\s*(?P<sign>[+-])\s*(?P<num>\d+(?:[.,]\d+)?)\s*(?P<unit>[^\d\s+-][^\d+-]*)?$")


@dataclass
class _Link:
    predecessor: str
    link_type: str
    lag_days: int
    warnings: list[tuple[str, dict[str, Any]]] = field(default_factory=list)


@dataclass
class _LinkError:
    code: str
    token: str
    params: dict[str, Any] = field(default_factory=dict)


class _IdIndex:
    """Activity refs, looked up exactly and then case-insensitively."""

    def __init__(self, refs: Iterable[str]) -> None:
        self.exact = set(refs)
        folded: dict[str, str | None] = {}
        for ref in self.exact:
            key = ref.casefold()
            folded[key] = None if key in folded and folded[key] != ref else ref
        self.folded = folded

    def find(self, text: str) -> str | None:
        text = text.strip()
        if not text:
            return None
        if text in self.exact:
            return text
        if text.endswith(".0") and text[:-2] in self.exact:
            return text[:-2]
        return self.folded.get(text.casefold())


def _parse_link(token: str, ids: _IdIndex, week_length: int) -> _Link | _LinkError:
    """One predecessor token: ``<id>[type][+/-lag[unit]]``.

    The whole token is tried as an id first, so an id that happens to end in
    letters ("A10FS") is never split; only then are a lag and a type suffix
    stripped.
    """
    head = token.strip()
    resolved = _resolve_link_head(head, ids)
    if resolved is not None:
        return _Link(resolved[0], resolved[1], 0)

    lag_match = _LAG.match(head)
    if not (lag_match and lag_match["head"].strip()):
        return _LinkError("predecessor_unknown", token, {"id": head})
    number, problem = parse_numeric_cell(lag_match["num"])
    if problem is not None or number is None:
        return _LinkError("predecessor_invalid", token)
    signed = -number if lag_match["sign"] == "-" else number
    lag = _to_days(signed, (lag_match["unit"] or "").strip() or None, week_length, allow_negative=True)
    resolved = _resolve_link_head(lag_match["head"].strip(), ids)
    if resolved is None:
        return _LinkError("predecessor_unknown", token, {"id": lag_match["head"].strip()})
    if lag.error == "elapsed_unsupported":
        return _LinkError("lag_elapsed_unsupported", token)
    if lag.error is not None:
        return _LinkError("predecessor_invalid", token, {"reason": "lag_unit"})
    return _Link(resolved[0], resolved[1], lag.days or 0, lag.warnings)


def _resolve_link_head(head: str, ids: _IdIndex) -> tuple[str, str] | None:
    """``(id, type)`` for ``<id>`` or ``<id><type>``; the bare id is tried first."""
    found = ids.find(head)
    if found is not None:
        return found, "FS"
    if len(head) > 2:
        mapped = relationship_type(head[-2:])
        if mapped is not None:
            found = ids.find(head[:-2].rstrip(" -_/"))
            if found is not None:
                return found, mapped
    return None


# ── Outline ───────────────────────────────────────────────────────────────────


def _levels_from_offsets(offsets: list[float | None]) -> list[int | None]:
    """Indent amounts (spaces or cell indent) to 1-based levels."""
    present = sorted({value for value in offsets if value is not None})
    if len(present) < 2:
        return [1 if value is not None else None for value in offsets]
    base = present[0]
    steps = [b - a for a, b in zip(present, present[1:], strict=False) if b - a > 0]
    unit = min(steps) if steps else 1.0
    return [None if value is None else int(round((value - base) / unit)) + 1 for value in offsets]


# ── The parse ─────────────────────────────────────────────────────────────────


@dataclass
class _Row:
    """One data row, its cells keyed by field."""

    number: int  # 1-based sheet row
    ordinal: int  # 1-based activity number
    cells: dict[str, Any]
    indent: float | None = None
    values: tuple[Any, ...] = ()  # every column as read, for the sample


def _parse_mapping_override(raw: Mapping[Any, Any] | None, width: int, issues: _Issues) -> dict[int, str] | None:
    """Validate an explicit ``{column index: field or ""}`` override."""
    if not raw:
        return None
    overrides: dict[int, str] = {}
    for key, value in raw.items():
        try:
            index = int(key)
        except (TypeError, ValueError):
            index = -1
        target = value or ""
        if index < 0 or index >= width or not isinstance(target, str) or (target and target not in FIELDS):
            issues.add(
                "column_mapping_invalid",
                "error",
                f"column_mapping entry {key!r}: {value!r} is not a column of this file and a known field",
                column=index if index >= 0 else None,
                key=str(key),
                value=str(value),
            )
            continue
        overrides[index] = target
    targets = [target for target in overrides.values() if target]
    for target in sorted(set(targets)):
        if targets.count(target) > 1:
            issues.add(
                "column_mapping_invalid",
                "error",
                f"column_mapping maps more than one column to {target}",
                field=target,
            )
    return overrides


def _map_columns(
    header: tuple[Any, ...],
    overrides: dict[int, str] | None,
    issues: _Issues,
    header_row: int,
) -> tuple[list[ColumnMatch], dict[str, int]]:
    columns: list[ColumnMatch] = []
    for index, value in enumerate(header):
        found = match_header(value)
        columns.append(ColumnMatch(index, _text(value).strip(), found.field, found.confidence, found.tier))
    if overrides:
        for index, target in overrides.items():
            if not target:
                columns[index].field, columns[index].confidence, columns[index].tier = None, CONFIDENCE_NONE, "none"
                continue
            for column in columns:
                if column.field == target and column.index != index:
                    column.field, column.confidence, column.tier = None, CONFIDENCE_NONE, "none"
            columns[index].field, columns[index].confidence, columns[index].tier = target, CONFIDENCE_EXACT, "override"

    mapping: dict[str, int] = {}
    for target in FIELDS:
        claimants = [column for column in columns if column.field == target]
        if not claimants:
            continue
        winner = max(claimants, key=lambda column: (column.confidence, -column.index))
        mapping[target] = winner.index
        losers = [column for column in claimants if column is not winner]
        if losers:
            issues.add(
                "header_collision",
                "warning",
                f"{len(claimants)} columns read as {target}; using column {winner.index} ({winner.header!r})",
                row=header_row,
                column=winner.index,
                field=target,
                used=winner.index,
                used_header=winner.header,
                ignored=[column.index for column in losers],
                ignored_headers=[column.header for column in losers],
            )
            for column in losers:
                column.field, column.confidence, column.tier = None, CONFIDENCE_NONE, "none"
    return columns, mapping


def _pick_sheet(sheets: list[_Sheet], issues: _Issues) -> _Sheet | None:
    """The first visible sheet whose header row names an activity column."""
    if not sheets:
        return None
    for sheet in sheets:
        found = find_header_row(iter(sheet.rows), known_header)
        fields = {known_header(_text(value)) for value in (found.values or ()) if value is not None}
        if "name" in fields or len(fields - {None}) >= 2:
            if len(sheets) > 1:
                issues.add(
                    "sheet_selected",
                    "info",
                    f"Read sheet {sheet.name!r}",
                    sheet=sheet.name,
                    sheets=[candidate.name for candidate in sheets],
                )
            return sheet
    return sheets[0]


def preview(
    data: bytes,
    filename: str,
    *,
    column_mapping: Mapping[Any, Any] | None = None,
    date_order: DateOrder | None = None,
    work_weekdays: Iterable[int] | None = None,
    holidays: Iterable[date] | None = None,
) -> TabularPreview:
    """Read a schedule spreadsheet and report what an import would create.

    Args:
        data: The uploaded bytes (.xlsx, .xls or a delimited text file).
        filename: The uploaded name; its extension only breaks ties.
        column_mapping: ``{column index: field}`` overrides of the detected
            mapping, an empty field unmaps the column. Indices are those of
            ``TabularPreview.columns``.
        date_order: ``"dmy"`` or ``"mdy"`` for numeric dates such as
            ``03/04/2026`` the values alone cannot settle.
        work_weekdays: Working weekdays (0 Monday .. 6 Sunday) of the target
            project; Monday to Friday when omitted.
        holidays: Non-working dates of the target project.

    Returns:
        A :class:`TabularPreview`. Its ``document`` is ``None`` only when the
        file could not be read at all; otherwise it is built even when errors
        were found, so the dialog can show what was read.
    """
    result = TabularPreview(sha256=hashlib.sha256(data).hexdigest(), filename=filename)
    issues = _Issues()
    try:
        _run(result, issues, data, filename, column_mapping, date_order, work_weekdays, holidays)
    finally:
        result.issues = issues.finish()
    return result


def _run(
    result: TabularPreview,
    issues: _Issues,
    data: bytes,
    filename: str,
    column_mapping: Mapping[Any, Any] | None,
    date_order: DateOrder | None,
    work_weekdays: Iterable[int] | None,
    holidays: Iterable[date] | None,
) -> None:
    sheet = _open(result, issues, data, filename)
    if sheet is None:
        return
    if sheet.truncated:
        issues.add(
            "too_many_rows",
            "error",
            f"The file has more than {MAX_ROWS} rows; split it into smaller files",
            limit=MAX_ROWS,
        )
        return

    found = find_header_row(iter(sheet.rows), known_header)
    header = found.values or ()
    result.header_row = found.number
    if not header:
        issues.add("no_header", "error", "The file has no rows", row=1)
        return
    if len(header) > MAX_COLUMNS:
        issues.add(
            "too_many_columns",
            "error",
            f"The header has more than {MAX_COLUMNS} columns",
            row=found.number,
            limit=MAX_COLUMNS,
        )
        return

    overrides = _parse_mapping_override(column_mapping, len(header), issues)
    columns, mapping = _map_columns(header, overrides, issues, found.number)
    result.columns = columns

    if "name" not in mapping and "notes" in mapping:
        mapping["name"] = mapping.pop("notes")
        column = columns[mapping["name"]]
        column.field = "name"
        issues.add(
            "name_from_description",
            "info",
            f"No name column; activity names are read from {column.header!r}",
            row=found.number,
            column=column.index,
            header=column.header,
        )
    result.mapping = dict(mapping)
    if "name" not in mapping:
        issues.add(
            "name_column_missing",
            "error",
            "No column holds the activity name; map one by hand",
            row=found.number,
            headers=[column.header for column in columns],
        )
        return

    rows = _data_rows(sheet, found.number, header, mapping, issues)
    result.row_count = len(rows)
    if len(rows) > MAX_ROWS:
        issues.add(
            "too_many_rows",
            "error",
            f"The file has {len(rows)} activity rows, more than {MAX_ROWS}",
            limit=MAX_ROWS,
            count=len(rows),
        )
        return
    if not rows:
        issues.add("no_rows", "error", "The file has a header but no activity rows", row=found.number)

    week = _Week(
        frozenset(work_weekdays) if work_weekdays is not None else _DEFAULT_WEEK,
        frozenset(holidays or ()),
    )
    if not week.weekdays:
        week = _Week(_DEFAULT_WEEK, week.holidays)
    builder = _Builder(result, issues, rows, header, mapping, week, date_order)
    builder.build()
    result.sample_rows = [
        {
            "row": row.number,
            "values": [_text(value)[:200] for value in row.values],
            "cells": {key: _text(value)[:200] for key, value in row.cells.items()},
        }
        for row in rows[:SAMPLE_ROWS]
    ]


def _open(result: TabularPreview, issues: _Issues, data: bytes, filename: str) -> _Sheet | None:
    """Size, type and bomb checks, then the rows of the sheet to read."""
    if not data:
        issues.add("file_empty", "error", "The file is empty")
        return None
    if len(data) > MAX_FILE_BYTES:
        issues.add(
            "file_too_large",
            "error",
            f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB",
            limit_mb=MAX_FILE_BYTES // (1024 * 1024),
            size=len(data),
        )
        return None

    signature = detect_signature(data[:16])
    lowered = filename.lower()
    if signature == "zip":
        file_format = "xlsx"
    elif signature == "ole":
        file_format = "xls"
    elif signature is None or signature == "xml":
        if lowered.endswith(_SPREADSHEET_EXTENSIONS):
            issues.add(
                "file_type_mismatch",
                "error",
                "The file is named as a workbook but its content is not one",
                filename=filename,
            )
            return None
        if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            issues.add("unsupported_file_type", "error", "The file is not a spreadsheet or a CSV")
            return None
        file_format = "csv"
    else:
        issues.add(
            "unsupported_file_type",
            "error",
            f"The file is a {signature}, not a spreadsheet or a CSV",
            detected=signature,
        )
        return None
    result.file_format = file_format

    if file_format == "csv":
        try:
            sheet, encoding, delimiter = _read_csv(data)
        except (csv.Error, UnicodeDecodeError) as exc:
            issues.add("file_unreadable", "error", f"The CSV could not be read: {exc}")
            return None
        result.encoding, result.delimiter = encoding, delimiter
        return sheet

    if file_format == "xlsx":
        from fastapi import HTTPException

        from app.core.upload_guards import reject_if_xlsx_bomb

        try:
            reject_if_xlsx_bomb(data, max_uncompressed=MAX_UNCOMPRESSED_BYTES)
        except HTTPException:
            issues.add(
                "file_too_large_uncompressed",
                "error",
                f"The workbook unpacks to more than {MAX_UNCOMPRESSED_BYTES // (1024 * 1024)} MB",
                limit_mb=MAX_UNCOMPRESSED_BYTES // (1024 * 1024),
            )
            return None
    try:
        sheets = _read_workbook(data, file_format)
    except _reader_errors() as exc:
        issues.add("file_unreadable", "error", f"The workbook could not be read: {exc}")
        return None
    sheet = _pick_sheet(sheets, issues)
    if sheet is None:
        issues.add("no_header", "error", "The workbook has no visible sheet")
        return None
    result.sheet = sheet.name
    return sheet


def _data_rows(
    sheet: _Sheet,
    header_row: int,
    header: tuple[Any, ...],
    mapping: dict[str, int],
    issues: _Issues,
) -> list[_Row]:
    """The non-empty rows under the header, cells keyed by field, long cells refused."""
    rows: list[_Row] = []
    by_index = {index: target for target, index in mapping.items()}
    name_index = mapping["name"]
    for offset, values in enumerate(sheet.rows[header_row:]):
        number = header_row + offset + 1
        if not any(value is not None and _text(value).strip() for value in values):
            continue
        cells: dict[str, Any] = {}
        for index, value in enumerate(values[: len(header)]):
            target = by_index.get(index)
            if target is None or value is None:
                continue
            if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
                issues.add(
                    "cell_too_long",
                    "error",
                    f"The cell holds {len(value)} characters, more than {MAX_CELL_CHARS}",
                    row=number,
                    column=index,
                    field=target,
                    limit=MAX_CELL_CHARS,
                )
                continue
            cells[target] = value
        indent = None
        if sheet.indents is not None:
            row_indents = sheet.indents[header_row + offset]
            indent = row_indents[name_index] if name_index < len(row_indents) else 0.0
        rows.append(_Row(number, len(rows) + 1, cells, indent, values[: len(header)]))
        if len(rows) > MAX_ROWS:
            break
    return rows


class _Builder:
    """Turns mapped rows into the interchange document, reporting as it goes."""

    def __init__(
        self,
        result: TabularPreview,
        issues: _Issues,
        rows: list[_Row],
        header: tuple[Any, ...],
        mapping: dict[str, int],
        week: _Week,
        date_order: DateOrder | None,
    ) -> None:
        self.result = result
        self.issues = issues
        self.rows = rows
        self.header = header
        self.mapping = mapping
        self.week = week
        self.date_order = date_order
        self.refs: list[str] = []
        self.ref_rows: dict[str, int] = {}
        self._fraction: bool | None = None
        self._undated: list[int] = []

    # ── entry ──

    def build(self) -> None:
        self._assign_refs()
        dates = self._dates()
        activities = [self._activity(row, dates) for row in self.rows]
        self.result.client_visible_suggested = [
            activity["ref"]
            for row, activity in zip(self.rows, activities, strict=True)
            if self._flag(row, "client_visible")
        ]
        if self._undated:
            self.issues.add(
                "dates_missing",
                "info",
                f"{len(self._undated)} activities have no dates; they will be calculated from their links",
                row=self._undated[0],
                count=len(self._undated),
                rows=self._undated[:MAX_ISSUES_PER_CODE],
            )
        self._outline(activities)
        relationships = self._relationships()
        self._loops(relationships)
        starts = [a["start_date"] for a in activities if a["start_date"]]
        finishes = [a["end_date"] for a in activities if a["end_date"]]
        stem = re.sub(r"\.[^.]+$", "", self.result.filename.replace("\\", "/").rsplit("/", 1)[-1]).strip()
        self.result.document = {
            "format": FORMAT,
            "format_version": FORMAT_VERSION,
            "schedule": {
                "name": (stem or "Imported schedule")[:_SCHEDULE_NAME_LIMIT],
                "schedule_type": "master",
                "description": "",
                "start_date": min(starts) if starts else None,
                "end_date": max(finishes) if finishes else None,
                "status": "draft",
                "data_date": None,
                "metadata": {
                    "import": {
                        "source": "spreadsheet",
                        "file_format": self.result.file_format,
                        "filename": self.result.filename,
                        "sha256": self.result.sha256,
                        "sheet": self.result.sheet,
                    }
                },
            },
            "activities": activities,
            "relationships": relationships,
        }

    # ── ids ──

    def _assign_refs(self) -> None:
        has_id = "id" in self.mapping
        seen: dict[str, int] = {}
        for row in self.rows:
            if not has_id:
                ref = str(row.number)
            else:
                ref = _text(row.cells.get("id")).strip()
                if not ref:
                    ref = f"row-{row.number}"
                    self.issues.add(
                        "id_missing",
                        "warning",
                        "The row has no activity id; other rows cannot name it as a predecessor",
                        row=row.number,
                        column=self.mapping["id"],
                    )
                elif ref in seen:
                    self.issues.add(
                        "duplicate_id",
                        "error",
                        f"Activity id {ref!r} is also used on row {seen[ref]}",
                        row=row.number,
                        column=self.mapping["id"],
                        id=ref,
                        first_row=seen[ref],
                    )
                    ref = f"{ref}#row-{row.number}"
            seen.setdefault(ref, row.number)
            self.refs.append(ref)
            self.ref_rows[ref] = row.number
        if not has_id and "predecessors" in self.mapping:
            self.issues.add(
                "predecessors_by_row_number",
                "info",
                "No id column; predecessors are read as sheet row numbers",
                column=self.mapping["predecessors"],
                first_row=self.rows[0].number if self.rows else None,
            )

    # ── dates ──

    def _dates(self) -> dict[tuple[int, str], date | None]:
        """Every start / finish cell resolved to a date, the day/month order settled once."""
        cells: list[tuple[int, str, _DateCell]] = []
        for position, row in enumerate(self.rows):
            for target in ("start", "finish"):
                if target not in self.mapping:
                    continue
                cell = _read_date(row.cells.get(target), row.number, self.mapping[target])
                if cell is not None:
                    cells.append((position, target, cell))

        ambiguous = [cell for _, _, cell in cells if not cell.invalid and cell.value is None]
        day_first = [cell for cell in ambiguous if cell.first > 12 >= cell.second]
        month_first = [cell for cell in ambiguous if cell.second > 12 >= cell.first]
        undecided = [cell for cell in ambiguous if cell.first <= 12 and cell.second <= 12 and cell.first != cell.second]

        order: DateOrder | None = None
        if self.date_order in ("dmy", "mdy"):
            order = self.date_order
            self.result.date_order_source = "explicit"
        elif day_first and month_first:
            self.issues.add(
                "date_order_conflict",
                "error",
                "Some dates only read day first and others only month first",
                row=day_first[0].row,
                column=day_first[0].column,
                day_first_example=day_first[0].raw,
                month_first_example=month_first[0].raw,
                day_first_row=day_first[0].row,
                month_first_row=month_first[0].row,
            )
            order = "dmy"
        elif day_first or month_first:
            order = "dmy" if day_first else "mdy"
            self.result.date_order_source = "values"
            example = (day_first or month_first)[0]
            self.issues.add(
                "date_order_detected",
                "info",
                f"Dates read {'day' if order == 'dmy' else 'month'} first, as {example.raw!r} shows",
                row=example.row,
                column=example.column,
                order=order,
                example=example.raw,
            )
        elif undecided:
            suggested, suggested_by = self._suggest_order(cells)
            order = suggested
            self.result.date_order_source = "suggested"
            self.issues.add(
                "date_order_unconfirmed",
                "error",
                f"Dates such as {undecided[0].raw!r} read either day first or month first; confirm the order",
                row=undecided[0].row,
                column=undecided[0].column,
                example=undecided[0].raw,
                suggested=suggested,
                suggested_by=suggested_by,
                columns=sorted({cell.column for cell in undecided}),
            )
        if ambiguous:
            self.result.date_order = order or "dmy"

        resolved: dict[tuple[int, str], date | None] = {}
        serial_reported = False
        for position, target, cell in cells:
            value = _resolve(cell, order or "dmy")
            if value is None:
                self.issues.add(
                    "date_invalid",
                    "error",
                    f"{cell.raw!r} is not a date",
                    row=cell.row,
                    column=cell.column,
                    field=target,
                    value=cell.raw,
                )
            elif cell.serial and not serial_reported:
                serial_reported = True
                self.issues.add(
                    "date_serial_converted",
                    "warning",
                    f"Numbers such as {cell.raw} were read as spreadsheet day numbers ({value.isoformat()})",
                    row=cell.row,
                    column=cell.column,
                    example=cell.raw,
                    date=value.isoformat(),
                )
            resolved[(position, target)] = value
        return resolved

    def _suggest_order(self, cells: list[tuple[int, str, _DateCell]]) -> tuple[DateOrder, str]:
        """The order whose dates agree with the durations best; else the separator's custom."""
        if "duration" in self.mapping and "start" in self.mapping and "finish" in self.mapping:
            by_row: dict[int, dict[str, _DateCell]] = {}
            for position, target, cell in cells:
                by_row.setdefault(position, {})[target] = cell
            score = {"dmy": 0, "mdy": 0}
            week_length = len(self.week.weekdays)
            for position, pair in by_row.items():
                if "start" not in pair or "finish" not in pair:
                    continue
                duration = _read_duration(self.rows[position].cells.get("duration"), week_length, None)
                if duration is None or duration.days is None:
                    continue
                for order in ("dmy", "mdy"):
                    start, finish = _resolve(pair["start"], order), _resolve(pair["finish"], order)
                    if start and finish and self.week.count(start, finish) == max(duration.days, 1):
                        score[order] += 1
            if score["dmy"] != score["mdy"]:
                return ("dmy" if score["dmy"] > score["mdy"] else "mdy"), "durations"
        return "dmy", "convention"

    # ── one activity ──

    def _activity(self, row: _Row, dates: dict[tuple[int, str], date | None]) -> dict[str, Any]:
        position = row.ordinal - 1
        ref = self.refs[position]
        cells = row.cells
        raw_name = _text(cells.get("name"))
        name = raw_name.strip()
        if not name:
            self.issues.add(
                "name_missing",
                "error",
                "The row has no activity name",
                row=row.number,
                column=self.mapping["name"],
            )

        start = dates.get((position, "start"))
        finish = dates.get((position, "finish"))
        duration = self._duration(row)
        milestone = self._flag(row, "milestone")
        if milestone and duration not in (None, 0):
            self.issues.add(
                "milestone_duration_ignored",
                "warning",
                f"The row is flagged a milestone; its duration of {duration} days is set to 0",
                row=row.number,
                column=self.mapping.get("duration"),
                duration=duration,
            )
            duration = 0
        if duration == 0:
            milestone = True
        elif milestone and duration is None:
            duration = 0

        start, finish, duration = self._reconcile(row, start, finish, duration)

        percent = self._percent(row)
        resources = [
            {"name": part.strip(), "type": "", "allocation_pct": 100.0, "count": 1}
            for part in re.split(r"[;,\n]", _text(cells.get("resource")))
            if part.strip()
        ]
        if percent is None or percent <= 0:
            status = "not_started"
        elif percent >= 100:
            status = "completed"
        else:
            status = "in_progress"
        activity: dict[str, Any] = {
            "ref": ref,
            "activity_code": _text(cells.get("id")).strip() or None,
            "name": name,
            "description": _text(cells.get("notes")).strip(),
            "wbs_code": _text(cells.get("wbs")).strip(),
            "parent_ref": None,
            "start_date": start.isoformat() if start else "",
            "end_date": finish.isoformat() if finish else "",
            "duration_days": duration or 0,
            "progress_pct": _percent_text(percent or 0.0),
            "status": status,
            "activity_type": "milestone" if milestone else "task",
            "sort_order": row.ordinal,
            "resources": resources,
            # Hidden until a person confirms it: the sheet's "yes" only lands in
            # ``client_visible_suggested``, and the commit applies the confirmed list.
            "client_visible": False,
            "metadata": {"import_row": row.number},
        }
        if not start and not finish:
            self._undated.append(row.number)
        self._check_lengths(row, activity)
        return activity

    def _check_lengths(self, row: _Row, activity: dict[str, Any]) -> None:
        """Refuse a value longer than the column that stores it, before the database does."""
        values = {
            "name": activity["name"],
            "id": activity["activity_code"] or "",
            "wbs": activity["wbs_code"],
            "resource": max((r["name"] for r in activity["resources"]), key=len, default=""),
        }
        for target, value in values.items():
            limit = _FIELD_LIMITS[target]
            if len(value) > limit:
                self.issues.add(
                    "value_too_long",
                    "error",
                    f"The {target} {value[:40]!r}... has {len(value)} characters, more than {limit}",
                    row=row.number,
                    column=self.mapping.get(target),
                    field=target,
                    limit=limit,
                    length=len(value),
                )

    def _duration(self, row: _Row) -> int | None:
        if "duration" not in self.mapping:
            return None
        column = self.mapping["duration"]
        hint = _header_unit(_text(self.header[column]))
        amount = _read_duration(row.cells.get("duration"), len(self.week.weekdays), hint)
        if amount is None:
            return None
        raw = _text(row.cells.get("duration"))
        if amount.error is not None:
            code = {
                "elapsed_unsupported": "duration_elapsed_unsupported",
                "unit_unsupported": "duration_unit_unsupported",
                "negative": "duration_negative",
            }.get(amount.error, "duration_invalid")
            self.issues.add(
                code,
                "error",
                f"{raw!r} is not a duration in working days, weeks or hours",
                row=row.number,
                column=column,
                value=raw,
            )
            return None
        for code, params in amount.warnings:
            self.issues.add(
                f"duration_{code}",
                "warning",
                f"Duration {raw!r} was read as {amount.days} working days",
                row=row.number,
                column=column,
                value=raw,
                **params,
            )
        return amount.days

    def _reconcile(
        self, row: _Row, start: date | None, finish: date | None, duration: int | None
    ) -> tuple[date | None, date | None, int | None]:
        """Fill the missing one of start / finish / duration; the duration wins a disagreement."""
        if start and finish and finish < start:
            self.issues.add(
                "finish_before_start",
                "error",
                f"The finish {finish.isoformat()} is before the start {start.isoformat()}",
                row=row.number,
                column=self.mapping.get("finish"),
                start=start.isoformat(),
                finish=finish.isoformat(),
            )
            return start, finish, duration
        if start and finish:
            counted = self.week.count(start, finish)
            if duration is None:
                return start, finish, counted
            consistent = counted == duration or (duration == 0 and start == finish)
            if not consistent:
                expected = self.week.finish_from(start, duration)
                self.issues.add(
                    "duration_mismatch",
                    "warning",
                    f"Start {start.isoformat()} and finish {finish.isoformat()} span {counted} working days, "
                    f"the duration says {duration}; the finish is moved to {expected.isoformat()}",
                    row=row.number,
                    column=self.mapping.get("duration"),
                    start=start.isoformat(),
                    finish=finish.isoformat(),
                    span=counted,
                    duration=duration,
                    new_finish=expected.isoformat(),
                )
                return start, expected, duration
            return start, finish, duration
        if start and duration is not None:
            return start, self.week.finish_from(start, duration), duration
        if finish and duration is not None:
            return self.week.start_from(finish, duration), finish, duration
        if start or finish:
            only = start or finish
            return only, only, 1 if duration is None else duration
        return start, finish, duration

    def _percent(self, row: _Row) -> float | None:
        if "percent_complete" not in self.mapping:
            return None
        column = self.mapping["percent_complete"]
        value = row.cells.get("percent_complete")
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
        raw = _text(value).strip()
        has_sign = "%" in raw
        number, problem = parse_numeric_cell(raw.replace("%", "").strip() if isinstance(value, str) else value)
        if problem is not None or number is None:
            self.issues.add(
                "percent_invalid",
                "error",
                f"{raw!r} is not a percentage",
                row=row.number,
                column=column,
                value=raw,
            )
            return None
        if not has_sign and self._percent_is_fraction():
            number *= 100
        if number < 0 or number > 100:
            self.issues.add(
                "percent_out_of_range",
                "error",
                f"{raw!r} is outside 0 to 100 percent",
                row=row.number,
                column=column,
                value=raw,
            )
            return None
        return number

    def _percent_is_fraction(self) -> bool:
        """True when the whole column is written as 0..1 fractions (0.45 for 45 %); reported once."""
        if self._fraction is not None:
            return self._fraction
        numbers: list[float] = []
        for row in self.rows:
            value = row.cells.get("percent_complete")
            if value is None or (isinstance(value, str) and ("%" in value or not value.strip())):
                continue
            number, problem = parse_numeric_cell(value)
            if problem is None and number is not None:
                numbers.append(number)
        fraction = bool(numbers) and all(0 <= n <= 1 for n in numbers) and any(0 < n < 1 for n in numbers)
        self._fraction = fraction
        if fraction:
            self.issues.add(
                "percent_fraction_assumed",
                "warning",
                "Percent complete values are all between 0 and 1 and were read as fractions (0.45 = 45 %)",
                column=self.mapping["percent_complete"],
            )
        return fraction

    def _flag(self, row: _Row, target: str) -> bool | None:
        if target not in self.mapping:
            return None
        value = row.cells.get(target)
        if value is None:
            return False
        read = flag_value(value)
        if read is None:
            self.issues.add(
                "flag_invalid",
                "warning",
                f"{_text(value)!r} is not a yes or a no; read as no",
                row=row.number,
                column=self.mapping[target],
                field=target,
                value=_text(value),
            )
            return False
        return read

    # ── outline ──

    def _outline(self, activities: list[dict[str, Any]]) -> None:
        """Set ``parent_ref`` from the first outline source the sheet has."""
        levels: list[int | None] | None = None
        source: str | None = None
        if "outline_level" in self.mapping:
            levels = self._level_column()
            source = "outline_level"
        elif any("." in a["wbs_code"].strip(".") for a in activities):
            self._wbs_parents(activities)
            self.result.outline_source = "wbs"
            return
        elif any(row.indent for row in self.rows):
            levels = _levels_from_offsets([row.indent or 0.0 for row in self.rows])
            source = "indent"
        else:
            spaces = [len(n) - len(n.lstrip(" \t　")) for n in (_text(r.cells.get("name")) for r in self.rows)]
            if len(set(spaces)) > 1:
                levels = _levels_from_offsets([float(s) for s in spaces])
                source = "leading_spaces"
        if levels is None:
            return
        self.result.outline_source = source
        self._parents_from_levels(activities, levels)

    def _level_column(self) -> list[int | None]:
        column = self.mapping["outline_level"]
        raw_levels: list[float | None] = []
        for row in self.rows:
            value = row.cells.get("outline_level")
            if value is None or (isinstance(value, str) and not value.strip()):
                raw_levels.append(None)
                continue
            number, problem = parse_numeric_cell(value)
            if problem is not None or number is None or number < 0 or not float(number).is_integer():
                self.issues.add(
                    "outline_level_invalid",
                    "error",
                    f"{_text(value)!r} is not an outline level",
                    row=row.number,
                    column=column,
                    value=_text(value),
                )
                raw_levels.append(None)
                continue
            raw_levels.append(number)
        present = [value for value in raw_levels if value is not None]
        base = min(present) if present else 1
        return [None if value is None else int(value - base) + 1 for value in raw_levels]

    def _parents_from_levels(self, activities: list[dict[str, Any]], levels: list[int | None]) -> None:
        stack: list[tuple[int, str]] = []
        previous = 0
        for row, activity, level in zip(self.rows, activities, levels, strict=True):
            if level is None:
                level = previous or 1
            if level > previous + 1:
                self.issues.add(
                    "outline_level_jump",
                    "warning",
                    f"The row goes from level {previous} to {level}; it is placed one level deeper",
                    row=row.number,
                    level=level,
                    previous=previous,
                )
                level = previous + 1
            while stack and stack[-1][0] >= level:
                stack.pop()
            activity["parent_ref"] = stack[-1][1] if stack else None
            stack.append((level, activity["ref"]))
            previous = level

    def _wbs_parents(self, activities: list[dict[str, Any]]) -> None:
        by_code: dict[str, str] = {}
        for activity in activities:
            code = activity["wbs_code"].strip().strip(".")
            if code and code not in by_code:
                by_code[code] = activity["ref"]
        for activity in activities:
            code = activity["wbs_code"].strip().strip(".")
            parts = code.split(".")
            for cut in range(len(parts) - 1, 0, -1):
                parent = by_code.get(".".join(parts[:cut]))
                if parent is not None and parent != activity["ref"]:
                    activity["parent_ref"] = parent
                    break

    # ── network ──

    def _relationships(self) -> list[dict[str, Any]]:
        if "predecessors" not in self.mapping:
            return []
        column = self.mapping["predecessors"]
        ids = _IdIndex(self.refs)
        week_length = len(self.week.weekdays)
        relationships: list[dict[str, Any]] = []
        seen: set[tuple[str, str]] = set()
        for position, row in enumerate(self.rows):
            successor = self.refs[position]
            text = _text(row.cells.get("predecessors")).strip()
            if not text:
                continue
            for token in (part.strip() for part in _TOKEN_SPLIT.split(text)):
                if not token:
                    continue
                parsed = _parse_link(token, ids, week_length)
                if isinstance(parsed, _LinkError):
                    messages = {
                        "predecessor_unknown": f"Predecessor {token!r} names no activity in the file",
                        "predecessor_invalid": f"Predecessor {token!r} is not of the form ID, ID FS or ID FS+2d",
                        "lag_elapsed_unsupported": f"Predecessor {token!r} has an elapsed lag; use working days",
                    }
                    self.issues.add(
                        parsed.code,
                        "error",
                        messages[parsed.code],
                        row=row.number,
                        column=column,
                        token=token,
                        **parsed.params,
                    )
                    continue
                for code, params in parsed.warnings:
                    self.issues.add(
                        f"lag_{code}",
                        "warning",
                        f"Lag in {token!r} was read as {parsed.lag_days} working days",
                        row=row.number,
                        column=column,
                        token=token,
                        **params,
                    )
                if parsed.predecessor == successor:
                    self.issues.add(
                        "self_dependency",
                        "error",
                        f"The activity names itself as a predecessor ({token!r})",
                        row=row.number,
                        column=column,
                        token=token,
                    )
                    continue
                pair = (parsed.predecessor, successor)
                if pair in seen:
                    self.issues.add(
                        "duplicate_predecessor",
                        "warning",
                        f"Predecessor {parsed.predecessor!r} is listed twice; the first link is kept",
                        row=row.number,
                        column=column,
                        token=token,
                    )
                    continue
                seen.add(pair)
                relationships.append(
                    {
                        "predecessor_ref": parsed.predecessor,
                        "successor_ref": successor,
                        "relationship_type": parsed.link_type,
                        "lag_days": parsed.lag_days,
                        "metadata": {},
                    }
                )
        return relationships

    def _loops(self, relationships: list[dict[str, Any]]) -> None:
        """Report every dependency loop: Kahn's algorithm, then a walk back through what is left."""
        edges = [(rel["predecessor_ref"], rel["successor_ref"]) for rel in relationships]
        for _ in range(_MAX_CYCLES):
            residual = _kahn_residual(edges)
            if not residual:
                return
            cycle = _find_cycle(edges, residual)
            first = cycle[0]
            self.issues.add(
                "dependency_loop",
                "error",
                "Activities depend on each other in a loop: " + " -> ".join(cycle),
                row=self.ref_rows.get(first),
                column=self.mapping.get("predecessors"),
                cycle=cycle,
            )
            closing = (cycle[-2], cycle[-1])
            edges = [edge for edge in edges if edge != closing]


def _kahn_residual(edges: list[tuple[str, str]]) -> set[str]:
    """Nodes Kahn's algorithm cannot order: those on a loop or downstream of one."""
    indegree: dict[str, int] = {}
    successors: dict[str, list[str]] = {}
    for pred, succ in edges:
        indegree.setdefault(pred, 0)
        indegree[succ] = indegree.get(succ, 0) + 1
        successors.setdefault(pred, []).append(succ)
    queue = deque(node for node, degree in indegree.items() if degree == 0)
    while queue:
        node = queue.popleft()
        for succ in successors.get(node, ()):
            indegree[succ] -= 1
            if indegree[succ] == 0:
                queue.append(succ)
    return {node for node, degree in indegree.items() if degree > 0}


def _find_cycle(edges: list[tuple[str, str]], residual: set[str]) -> list[str]:
    """One loop inside ``residual``, in link order and closed (first id repeated last).

    Walks predecessor edges: every residual node has a residual predecessor, so
    the walk must come back to a node it has seen, and that stretch is a loop.
    """
    predecessors: dict[str, str] = {}
    for pred, succ in edges:
        if pred in residual and succ in residual and succ not in predecessors:
            predecessors[succ] = pred
    node = min(residual)
    order: dict[str, int] = {}
    path: list[str] = []
    while node not in order:
        order[node] = len(path)
        path.append(node)
        node = predecessors[node]
    loop = path[order[node] :]
    loop.reverse()
    return [*loop, loop[0]]


def _percent_text(value: float) -> str:
    """Percent as the interchange cleaner writes it: ``"45"``, ``"45.5"``."""
    value = round(min(100.0, max(0.0, value)), 2)
    return str(int(value)) if value == int(value) else str(value)


__all__ = [
    "MAX_CELL_CHARS",
    "MAX_COLUMNS",
    "MAX_FILE_BYTES",
    "MAX_ROWS",
    "ColumnMatch",
    "Issue",
    "TabularPreview",
    "preview",
]
