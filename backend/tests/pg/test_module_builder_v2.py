# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Module builder v2 against PostgreSQL: links, status, export, comments, deadlines.

A module with links and every feature, installed for real, read and written
through its own HTTP API and the platform routes that serve it (the link
picker, comment threads, the deadline register). Each test is built to fail on
the version of the code that would get it wrong:

- SET NULL is proven by deleting the linked rows, not by reading the DDL;
- the status default is proven by an insert that does not name the column,
  so the ORM default cannot stand in for the missing server default;
- a link into another project is refused with its own code, and a link to a
  record the caller may not see with another, so neither leaks the other;
- the picker, the labels, the export and the comment thread each answer for
  someone else's project exactly as for a record that does not exist.
"""

from __future__ import annotations

import csv
import importlib
import io
import sys
import uuid
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.modules.module_builder import links, service
from app.modules.module_builder.spec import ModuleSpec

KEY = "pg_pour_log"
BASE = "/api/v1/pg-pour-log"
TABLE = f"oe_{KEY}_pour"
CLASS_NAME = "Pour"
MB = "/api/v1/module-builder"
COLLAB = "/api/v1/collaboration"

# Close to the limit on purpose: the table name is 46 bytes, the two link
# indexes 62 and 63, and the naming convention's foreign key names run far
# past 63 and have to be shortened by SQLAlchemy without colliding.
LONG_KEY = "pg_long_register_for_checks"
LONG_TABLE = f"oe_{LONG_KEY}_inspection_item"


def _spec() -> ModuleSpec:
    return ModuleSpec.model_validate(
        {
            "key": KEY,
            "display_name": "Pour Log",
            "description": "Every concrete pour, its contract and who signs it off.",
            "schema_version": 2,
            "entity": {
                "name": "pour",
                "display_name": "Pour",
                "plural_name": "Pours",
                "project_scoped": True,
                "fields": [
                    {"name": "reference", "label": "Reference", "type": "text", "required": True},
                    {"name": "due_on", "label": "Due", "type": "date"},
                    {"name": "contract", "label": "Contract", "type": "link", "target": "contract"},
                    {"name": "responsible", "label": "Responsible", "type": "link", "target": "user"},
                ],
            },
            "rules": [
                {"code": "REFERENCE_REQUIRED", "message": "A pour needs a reference.", "kind": "required",
                 "field": "reference"},
            ],
            "features": {
                "status": {
                    "states": [
                        {"code": "open", "label": "Open"},
                        {"code": "in_progress", "label": "In progress"},
                        {"code": "done", "label": "Done", "done": True},
                    ]
                },
                "due": {"field": "due_on", "remind_days_before": 2},
                "export": True,
                "comments": True,
            },
        }
    )  # fmt: skip


def _long_spec() -> ModuleSpec:
    return ModuleSpec.model_validate(
        {
            "key": LONG_KEY,
            "display_name": "Long Register",
            "schema_version": 2,
            "entity": {
                "name": "inspection_item",
                "display_name": "Inspection item",
                "project_scoped": True,
                "fields": [
                    {"name": "title", "label": "Title", "type": "text", "required": True},
                    {"name": "contract_ref", "label": "Contract", "type": "link", "target": "contract"},
                    {"name": "responsible_x", "label": "Responsible", "type": "link", "target": "user"},
                ],
            },
            "rules": [{"code": "TITLE_REQUIRED", "message": "Title.", "kind": "required", "field": "title"}],
        }
    )


def _forget(key: str, class_name: str, table: str) -> None:
    from app.database import Base

    for name in [n for n in list(sys.modules) if n.startswith(f"app.modules.{key}")]:
        del sys.modules[name]
    existing = Base.metadata.tables.get(table)
    if existing is not None:
        Base.metadata.remove(existing)
    Base.registry._class_registry.pop(class_name, None)


@pytest.fixture
def runtime_root(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from app.core import module_runtime_root as rr

    root = tmp_path / "runtime-modules"
    monkeypatch.setenv(rr.ENV_VAR, str(root))
    before = list(rr._package_path())
    try:
        yield root
    finally:
        rr._package_path()[:] = before
        _forget(KEY, CLASS_NAME, TABLE)
        _forget(LONG_KEY, "InspectionItem", LONG_TABLE)
        importlib.invalidate_caches()


@pytest.fixture(autouse=True)
def instance_state(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.core import module_state

    monkeypatch.setattr(module_state, "_resolve_data_dir", lambda data_dir=None: tmp_path / "instance-state")


@pytest.fixture(autouse=True)
def platform(pg_engine, monkeypatch: pytest.MonkeyPatch):
    """The engine the startup hook uses, the targets' modules, and their permissions.

    The link targets' modules are not loaded by the module loader in this lane,
    so availability is answered yes here; the permissions the routes check are
    registered as each module's own startup would.
    """
    import app.database
    from app.modules.collaboration.permissions import register_collaboration_permissions
    from app.modules.contracts.permissions import register_contracts_permissions
    from app.modules.module_builder.permissions import register_module_builder_permissions
    from app.modules.projects.permissions import register_project_permissions
    from app.modules.users.permissions import register_user_permissions

    monkeypatch.setattr(app.database, "engine", pg_engine)
    monkeypatch.setattr(links, "available", lambda target: target in links.LINK_TARGETS)
    for register in (
        register_collaboration_permissions,
        register_contracts_permissions,
        register_module_builder_permissions,
        register_project_permissions,
        register_user_permissions,
    ):
        register()


@pytest.fixture
def app() -> FastAPI:
    from app.modules.collaboration.router import router as collaboration_router
    from app.modules.module_builder.router import router as builder_router

    application = FastAPI()
    application.include_router(builder_router, prefix=MB)
    application.include_router(collaboration_router, prefix=COLLAB)
    return application


@pytest_asyncio.fixture
async def installed(runtime_root: Path, app: FastAPI, pg_engine):
    await service.install(_spec(), app)
    try:
        yield
    finally:
        try:
            await service.uninstall(KEY, app, drop_data=True)
        except service.InstallRefused:
            pass


@pytest_asyncio.fixture
async def factory(installed, pg_engine) -> AsyncIterator[async_sessionmaker[AsyncSession]]:
    connection = await pg_engine.connect()
    outer = await connection.begin()
    try:
        yield async_sessionmaker(
            bind=connection, class_=AsyncSession, join_transaction_mode="create_savepoint", expire_on_commit=False
        )
    finally:
        if outer.is_active:
            await outer.rollback()
        await connection.close()


async def _user(session: AsyncSession, role: str, name: str | None = None) -> str:
    from app.modules.users.models import User

    user_id = uuid.uuid4()
    session.add(
        User(
            id=user_id,
            email=f"{role}-{user_id.hex[:8]}@example.test",
            hashed_password="x",
            full_name=name or f"Test {role}",
            role=role,
            locale="en",
            is_active=True,
            metadata_={},
        )
    )
    await session.flush()
    return str(user_id)


async def _project(session: AsyncSession, owner: str, name: str) -> uuid.UUID:
    from app.modules.projects.models import Project

    project_id = uuid.uuid4()
    session.add(
        Project(
            id=project_id,
            name=name,
            description="Module builder v2 fixture",
            currency="EUR",
            status="active",
            owner_id=uuid.UUID(owner),
            metadata_={},
        )
    )
    await session.flush()
    return project_id


async def _member(session: AsyncSession, project_id: uuid.UUID, user_id: str) -> str:
    """Put a user on a project's team. Returns the membership id."""
    from app.modules.teams.models import Team, TeamMembership

    team_id = uuid.uuid4()
    session.add(Team(id=team_id, project_id=project_id, name=f"Team {team_id.hex[:6]}", metadata_={}))
    await session.flush()
    membership_id = uuid.uuid4()
    session.add(TeamMembership(id=membership_id, team_id=team_id, user_id=uuid.UUID(user_id)))
    await session.flush()
    return str(membership_id)


async def _contract(session: AsyncSession, project_id: uuid.UUID, title: str) -> str:
    from app.modules.contracts.models import Contract

    contract_id = uuid.uuid4()
    session.add(Contract(id=contract_id, code=f"C-{contract_id.hex[:6]}", title=title, project_id=project_id))
    await session.flush()
    return str(contract_id)


def _as(app: FastAPI, factory: async_sessionmaker[AsyncSession], user_id: str, role: str) -> None:
    from app.dependencies import get_current_user_id, get_current_user_payload, get_session

    async def _session():
        async with factory() as session:
            yield session
            await session.commit()

    app.dependency_overrides[get_session] = _session
    app.dependency_overrides[get_current_user_id] = lambda: user_id
    app.dependency_overrides[get_current_user_payload] = lambda: {"sub": user_id, "role": role, "permissions": []}


@asynccontextmanager
async def _client(app: FastAPI) -> AsyncIterator[AsyncClient]:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://v2-test") as client:
        yield client


@dataclass
class World:
    admin: str
    alice: str
    bob: str
    alice_site: uuid.UUID
    alice_yard: uuid.UUID
    bob_site: uuid.UUID
    alice_contract: str
    alice_yard_contract: str
    bob_contract: str
    alice_record: str
    bob_record: str


def _body(project_id: uuid.UUID, reference: str, **extra: object) -> dict:
    return {"project_id": str(project_id), "reference": reference, **extra}


@pytest_asyncio.fixture
async def world(app: FastAPI, factory) -> World:
    async with factory() as session:
        admin = await _user(session, "admin")
        alice = await _user(session, "manager", "Alice Builder")
        bob = await _user(session, "manager", "Bob Builder")
        alice_site = await _project(session, alice, "Alice's site")
        alice_yard = await _project(session, alice, "Alice's yard")
        bob_site = await _project(session, bob, "Bob's site")
        alice_contract = await _contract(session, alice_site, "Frame works")
        alice_yard_contract = await _contract(session, alice_yard, "Yard works")
        bob_contract = await _contract(session, bob_site, "Bob's secret works")
        await session.commit()

    yesterday = (datetime.now(UTC).date() - timedelta(days=1)).isoformat()
    async with _client(app) as client:
        _as(app, factory, alice, "manager")
        a = await client.post(
            BASE, json=_body(alice_site, "A-1", contract=alice_contract, responsible=alice, due_on=yesterday)
        )
        _as(app, factory, bob, "manager")
        b = await client.post(BASE, json=_body(bob_site, "B-1", contract=bob_contract, due_on=yesterday))
    assert a.status_code == 201, a.text
    assert b.status_code == 201, b.text
    return World(
        admin=admin,
        alice=alice,
        bob=bob,
        alice_site=alice_site,
        alice_yard=alice_yard,
        bob_site=bob_site,
        alice_contract=alice_contract,
        alice_yard_contract=alice_yard_contract,
        bob_contract=bob_contract,
        alice_record=a.json()["id"],
        bob_record=b.json()["id"],
    )


# ── the table ───────────────────────────────────────────────────────────────


class TestTheTable:
    async def test_links_set_null_when_what_they_point_at_is_deleted(self, app: FastAPI, factory, world: World) -> None:
        """Deleted for real, not read off the DDL: a wrong ondelete blocks the delete or takes the row."""
        async with factory() as session:
            throwaway = await _user(session, "editor", "Leaves soon")
            await _member(session, world.alice_site, throwaway)
            gone = await _contract(session, world.alice_site, "Cancelled works")
            await session.commit()
        _as(app, factory, world.admin, "admin")
        async with _client(app) as client:
            created = await client.post(BASE, json=_body(world.alice_site, "A-2", contract=gone, responsible=throwaway))
        assert created.status_code == 201, created.text
        record = created.json()["id"]

        async with factory() as session:
            await session.execute(text("DELETE FROM oe_contracts_contract WHERE id = :id"), {"id": gone})
            await session.execute(text("DELETE FROM oe_teams_membership WHERE user_id = :id"), {"id": throwaway})
            await session.execute(text("DELETE FROM oe_users_user WHERE id = :id"), {"id": throwaway})
            await session.commit()

        async with _client(app) as client:
            after = await client.get(f"{BASE}/{record}")
        assert after.status_code == 200, after.text
        assert after.json()["contract"] is None
        assert after.json()["responsible"] is None
        assert after.json()["reference"] == "A-2"

    async def test_an_insert_that_names_no_status_starts_in_the_first_state(self, factory, world: World) -> None:
        """Raw SQL, so only the server default can answer."""
        record = uuid.uuid4()
        async with factory() as session:
            await session.execute(
                text(f"INSERT INTO {TABLE} (id, project_id, reference) VALUES (:id, :project, 'RAW')"),
                {"id": str(record), "project": str(world.alice_site)},
            )
            status = (
                await session.execute(text(f"SELECT status FROM {TABLE} WHERE id = :id"), {"id": str(record)})
            ).scalar()
        assert status == "open"

    async def test_every_declared_index_and_foreign_key_exists(self, factory, world: World) -> None:
        async with factory() as session:
            indexes = set(
                (
                    await session.execute(text("SELECT indexname FROM pg_indexes WHERE tablename = :t"), {"t": TABLE})
                ).scalars()
            )
            deletes = dict(
                (
                    await session.execute(
                        text(
                            "SELECT a.attname, CAST(c.confdeltype AS text) FROM pg_constraint c "
                            "JOIN pg_attribute a ON a.attrelid = c.conrelid AND a.attnum = ANY(c.conkey) "
                            "WHERE c.conrelid = CAST(:t AS regclass) AND c.contype = 'f'"
                        ),
                        {"t": TABLE},
                    )
                ).all()
            )
        assert {f"ix_{TABLE}_contract", f"ix_{TABLE}_responsible", f"ix_{TABLE}_status"} <= indexes
        assert {f"ix_{TABLE}_project", f"ix_{TABLE}_created"} <= indexes
        # n = SET NULL, c = CASCADE
        assert deletes == {"contract": "n", "responsible": "n", "project_id": "c"}

    async def test_names_close_to_the_limit_install_without_colliding(
        self, runtime_root: Path, app: FastAPI, pg_engine
    ) -> None:
        spec = _long_spec()
        assert spec.table_name == LONG_TABLE
        assert max(len(n) for n in spec.new_index_names) == 63
        await service.install(spec, app)
        try:
            async with pg_engine.connect() as connection:
                foreign_keys = (
                    await connection.execute(
                        text(
                            "SELECT conname FROM pg_constraint WHERE conrelid = CAST(:t AS regclass) AND contype = 'f'"
                        ),
                        {"t": LONG_TABLE},
                    )
                ).scalars()
                names = sorted(foreign_keys)
                indexes = set(
                    (
                        await connection.execute(
                            text("SELECT indexname FROM pg_indexes WHERE tablename = :t"), {"t": LONG_TABLE}
                        )
                    ).scalars()
                )
            assert len(names) == 3 and len(set(names)) == 3, names
            assert all(len(n) <= 63 for n in names)
            assert set(spec.new_index_names) <= indexes
        finally:
            await service.uninstall(LONG_KEY, app, drop_data=True)


# ── links ───────────────────────────────────────────────────────────────────


class TestLinksStayInsideTheProject:
    async def test_a_contract_of_another_of_my_projects_is_outside_this_one(
        self, app: FastAPI, factory, world: World
    ) -> None:
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            created = await client.post(BASE, json=_body(world.alice_site, "A-3", contract=world.alice_yard_contract))
            moved = await client.patch(f"{BASE}/{world.alice_record}", json={"contract": world.alice_yard_contract})
            after = await client.get(f"{BASE}/{world.alice_record}")

        assert created.status_code == 422, created.text
        assert [d["code"] for d in created.json()["detail"]] == ["link_outside_project"]
        assert moved.status_code == 422, moved.text
        assert [d["code"] for d in moved.json()["detail"]] == ["link_outside_project"]
        assert after.json()["contract"] == world.alice_contract

    async def test_someone_elses_contract_reads_as_one_that_does_not_exist(
        self, app: FastAPI, factory, world: World
    ) -> None:
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            foreign = await client.post(BASE, json=_body(world.alice_site, "A-4", contract=world.bob_contract))
            missing = await client.post(BASE, json=_body(world.alice_site, "A-4", contract=str(uuid.uuid4())))

        assert foreign.status_code == 422, foreign.text
        assert foreign.json() == missing.json()
        assert [d["code"] for d in foreign.json()["detail"]] == ["link_not_found"]

    async def test_an_unchanged_link_is_not_checked_again(self, app: FastAPI, factory, world: World) -> None:
        """A person who has left the project stays on the record until someone changes it.

        The link check covers what the write changes. Checking the stored link
        again would lock the editor out of their own record the day a team
        member leaves, and the control below shows the same link is refused
        when it is written again.
        """
        async with factory() as session:
            erin = await _user(session, "editor", "Erin Editor")
            dave = await _user(session, "editor", "Dave Departing")
            erin_site = await _project(session, erin, "Erin's site")
            membership = await _member(session, erin_site, dave)
            await session.commit()
        _as(app, factory, erin, "editor")
        async with _client(app) as client:
            assigned = await client.post(BASE, json=_body(erin_site, "E-1", responsible=dave))
        assert assigned.status_code == 201, assigned.text
        record = assigned.json()["id"]

        async with factory() as session:
            await session.execute(text("DELETE FROM oe_teams_membership WHERE id = :id"), {"id": membership})
            await session.commit()

        async with _client(app) as client:
            untouched = await client.patch(f"{BASE}/{record}", json={"reference": "E-1a"})
            reassigned = await client.patch(f"{BASE}/{record}", json={"responsible": dave})
        assert untouched.status_code == 200, untouched.text
        assert untouched.json()["responsible"] == dave
        assert reassigned.status_code == 422, reassigned.text
        assert [d["code"] for d in reassigned.json()["detail"]] == ["link_not_found"]

    async def test_a_person_must_be_on_the_records_project(self, app: FastAPI, factory, world: World) -> None:
        """Alice shares her yard with Yuri, not her site: Yuri may be picked for the yard only."""
        async with factory() as session:
            yuri = await _user(session, "editor", "Yuri Yard")
            await _member(session, world.alice_yard, yuri)
            await session.commit()
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            on_site = await client.post(BASE, json=_body(world.alice_site, "A-5", responsible=yuri))
            in_yard = await client.post(BASE, json=_body(world.alice_yard, "Y-1", responsible=yuri))
            stranger = await client.post(BASE, json=_body(world.alice_site, "A-6", responsible=world.bob))
        assert on_site.status_code == 422, on_site.text
        assert [d["code"] for d in on_site.json()["detail"]] == ["link_outside_project"]
        assert in_yard.status_code == 201, in_yard.text
        assert stranger.status_code == 422, stranger.text
        assert [d["code"] for d in stranger.json()["detail"]] == ["link_not_found"]

    async def test_an_unknown_status_is_refused(self, app: FastAPI, factory, world: World) -> None:
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            refused = await client.patch(f"{BASE}/{world.alice_record}", json={"status": "archived"})
            moved = await client.patch(f"{BASE}/{world.alice_record}", json={"status": "in_progress"})
        assert refused.status_code == 422, refused.text
        assert [d["code"] for d in refused.json()["detail"]] == ["STATUS_KNOWN"]
        assert moved.json()["status"] == "in_progress"


class TestThePicker:
    async def test_it_lists_only_contracts_of_reachable_projects(self, app: FastAPI, factory, world: World) -> None:
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            everything = await client.get(f"{MB}/lookup/contract")
            one = await client.get(f"{MB}/lookup/contract", params={"project_id": str(world.alice_site)})
            searched = await client.get(f"{MB}/lookup/contract", params={"q": "yard"})
            foreign = await client.get(f"{MB}/lookup/contract", params={"project_id": str(world.bob_site)})

        assert {i["id"] for i in everything.json()["items"]} == {world.alice_contract, world.alice_yard_contract}
        assert [i["id"] for i in one.json()["items"]] == [world.alice_contract]
        assert [i["label"] for i in searched.json()["items"]] == ["Yard works"]
        assert foreign.status_code == 404, foreign.text

    async def test_labels_leave_out_what_the_caller_may_not_see(self, app: FastAPI, factory, world: World) -> None:
        _as(app, factory, world.alice, "manager")
        ids = ",".join([world.alice_contract, world.bob_contract, str(uuid.uuid4()), "not-an-id"])
        async with _client(app) as client:
            labels = await client.get(f"{MB}/lookup/contract/labels", params={"ids": ids})
        assert labels.status_code == 200, labels.text
        assert labels.json() == {"labels": {world.alice_contract: "Frame works"}}

    async def test_an_unknown_target_is_404(self, app: FastAPI, factory, world: World) -> None:
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            unknown = await client.get(f"{MB}/lookup/invoice")
        assert unknown.status_code == 404

    async def test_people_are_those_who_share_a_project(self, app: FastAPI, factory, world: World) -> None:
        """The decided rule: an editor on project A never sees someone who is only on project B.

        Ed is on Alice's site. Bob owns only his own site, so Ed may not see,
        pick or read the name of Bob; Alice he shares a project with.
        """
        async with factory() as session:
            ed = await _user(session, "editor", "Ed Editor")
            await _member(session, world.alice_site, ed)
            await session.commit()
        _as(app, factory, ed, "editor")
        ids = ",".join([world.alice, world.bob, ed])
        async with _client(app) as client:
            everyone = await client.get(f"{MB}/lookup/user")
            site = await client.get(f"{MB}/lookup/user", params={"project_id": str(world.alice_site)})
            yard = await client.get(f"{MB}/lookup/user", params={"project_id": str(world.alice_yard)})
            labels = await client.get(f"{MB}/lookup/user/labels", params={"ids": ids})

        assert everyone.status_code == 200, everyone.text
        assert {i["id"] for i in everyone.json()["items"]} == {world.alice, ed}
        assert {i["id"] for i in site.json()["items"]} == {world.alice, ed}
        # Ed is not on Alice's yard, so naming it is the same 404 as any project he cannot reach.
        assert yard.status_code == 404, yard.text
        assert labels.json() == {"labels": {world.alice: "Alice Builder", ed: "Ed Editor"}}

    async def test_with_a_project_only_its_people_are_listed(self, app: FastAPI, factory, world: World) -> None:
        async with factory() as session:
            yuri = await _user(session, "editor", "Yuri Yard")
            await _member(session, world.alice_yard, yuri)
            await session.commit()
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            site = await client.get(f"{MB}/lookup/user", params={"project_id": str(world.alice_site)})
            yard = await client.get(f"{MB}/lookup/user", params={"project_id": str(world.alice_yard)})
            everyone = await client.get(f"{MB}/lookup/user")
        assert {i["id"] for i in site.json()["items"]} == {world.alice}
        assert {i["id"] for i in yard.json()["items"]} == {world.alice, yuri}
        assert {i["id"] for i in everyone.json()["items"]} == {world.alice, yuri}

    async def test_an_admin_sees_everyone(self, app: FastAPI, factory, world: World) -> None:
        _as(app, factory, world.admin, "admin")
        async with _client(app) as client:
            listed = await client.get(f"{MB}/lookup/user", params={"q": "Builder"})
        assert {world.alice, world.bob} <= {i["id"] for i in listed.json()["items"]}


# ── export ──────────────────────────────────────────────────────────────────


class TestExport:
    async def test_the_file_holds_what_the_list_shows_and_no_more(self, app: FastAPI, factory, world: World) -> None:
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            await client.post(BASE, json=_body(world.alice_site, "=HYPERLINK(1)"))
            listed = await client.get(BASE)
            exported = await client.get(f"{BASE}/export", params={"format": "csv"})
            foreign = await client.get(f"{BASE}/export", params={"project_id": str(world.bob_site)})

        assert exported.status_code == 200, exported.text
        assert exported.headers["content-disposition"] == f'attachment; filename="{KEY}.csv"'
        rows = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
        header, body = rows[0], rows[1:]
        assert header[:5] == ["Project", "Reference", "Due", "Contract", "Responsible"]
        assert len(body) == listed.json()["total"] == 2
        by_ref = {r[1]: r for r in body}
        assert "B-1" not in by_ref
        assert by_ref["A-1"][3] == "Frame works"
        assert by_ref["A-1"][4] == "Alice Builder"
        assert by_ref["A-1"][5] == "Open"
        assert "'=HYPERLINK(1)" in by_ref
        assert foreign.status_code == 404, foreign.text

    async def test_the_workbook_matches_the_csv(self, app: FastAPI, factory, world: World) -> None:
        from openpyxl import load_workbook

        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            workbook = await client.get(f"{BASE}/export", params={"format": "xlsx"})
            unknown = await client.get(f"{BASE}/export", params={"format": "pdf"})
        assert workbook.status_code == 200, workbook.text
        sheet = load_workbook(io.BytesIO(workbook.content)).active
        values = list(sheet.iter_rows(values_only=True))
        assert values[0][1] == "Reference"
        assert [row[1] for row in values[1:]] == ["A-1"]
        assert unknown.status_code == 422


# ── comments ────────────────────────────────────────────────────────────────


class TestComments:
    async def test_a_record_takes_comments_from_whoever_can_reach_it(self, app: FastAPI, factory, world: World) -> None:
        entity_type = f"built.{KEY}"
        _as(app, factory, world.alice, "manager")
        async with _client(app) as client:
            posted = await client.post(
                f"{COLLAB}/comments/",
                json={"entity_type": entity_type, "entity_id": world.alice_record, "text": "Pour checked."},
            )
            listed = await client.get(
                f"{COLLAB}/comments/", params={"entity_type": entity_type, "entity_id": world.alice_record}
            )
        assert posted.status_code == 201, posted.text
        assert listed.status_code == 200, listed.text

        _as(app, factory, world.bob, "manager")
        async with _client(app) as client:
            foreign_read = await client.get(
                f"{COLLAB}/comments/", params={"entity_type": entity_type, "entity_id": world.alice_record}
            )
            foreign_write = await client.post(
                f"{COLLAB}/comments/",
                json={"entity_type": entity_type, "entity_id": world.alice_record, "text": "Hello"},
            )
            missing = await client.get(
                f"{COLLAB}/comments/", params={"entity_type": entity_type, "entity_id": str(uuid.uuid4())}
            )
        assert foreign_read.status_code == 404
        assert foreign_write.status_code == 404
        assert foreign_read.json() == missing.json()

    async def test_a_key_that_is_not_an_installed_module_is_closed(self, app: FastAPI, factory, world: World) -> None:
        _as(app, factory, world.admin, "admin")
        async with _client(app) as client:
            unknown = await client.get(
                f"{COLLAB}/comments/", params={"entity_type": "built.not_installed", "entity_id": world.alice_record}
            )
            shipped = await client.get(
                f"{COLLAB}/comments/", params={"entity_type": "built.projects", "entity_id": str(world.alice_site)}
            )
        assert unknown.status_code == 404
        assert shipped.status_code == 404


# ── deadlines ───────────────────────────────────────────────────────────────


class TestDeadlines:
    async def test_the_register_holds_overdue_records_of_reachable_projects(
        self, app: FastAPI, factory, world: World
    ) -> None:
        from app.modules.deadlines.service import compute_deadlines

        _as(app, factory, world.alice, "manager")
        yesterday = (datetime.now(UTC).date() - timedelta(days=1)).isoformat()
        async with _client(app) as client:
            finished = await client.post(BASE, json=_body(world.alice_site, "A-done", due_on=yesterday, status="done"))
        assert finished.status_code == 201, finished.text

        async with factory() as session:
            mine = await compute_deadlines(session, [world.alice_site], module="built_modules")
            everyone = await compute_deadlines(session, None, module="built_modules", now=datetime.now(UTC))

        assert [i.entity_id for i in mine.items] == [world.alice_record]
        item = mine.items[0]
        assert item.entity_type == f"built.{KEY}"
        assert item.classification == "overdue"
        assert item.source_label == "Pour Log"
        assert item.owner_user_id == world.alice
        assert item.remind_days == 2
        assert item.action_url == f"/projects/{world.alice_site}/modules/{KEY}"
        assert {i.entity_id for i in everyone.items} == {world.alice_record, world.bob_record}
