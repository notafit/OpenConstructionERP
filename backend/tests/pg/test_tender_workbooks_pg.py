# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""PG: the tender workbooks are the stored package, read through the routes the contractor uses.

The comparison workbook must say what the comparison endpoint says: the same
totals per bidder, a line a bidder left out shown as missing rather than as 0,
and the project's own currency rather than a default. The bill for bidders
must carry the package's lines and quantities and none of the contractor's
own rates or totals, anywhere in the file.

Both routes are called as handlers on a real session, so the package, its
bids and the bill are the stored rows and the comparison is the service's own
answer, not a hand-built one.

Gated by ``OE_TEST_DB=pg`` (see conftest).
"""

from __future__ import annotations

import io
import uuid
import zipfile
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.modules.boq.models import BOQ, Position
from app.modules.projects.models import Project
from app.modules.tendering.models import TenderBid, TenderPackage
from app.modules.tendering.router import (
    export_bid_comparison_xlsx,
    export_bill_for_bidders_xlsx,
    export_package_gaeb_x83,
)
from app.modules.tendering.service import TenderingService
from app.modules.tendering.workbooks import LOWEST_FILL, MISSING_FILL
from app.modules.users.models import User

# The contractor's own pricing. Distinctive figures, so a leak into the bill
# for bidders is found by a plain text search of the file.
LINES = (
    ("01.010", "Excavation", "m3", "120", "18.37", "2204.40"),
    ("01.020", "Blinding concrete", "m2", "80", "41.93", "3354.40"),
    ("01.030", "Formwork", "m2", "60", "57.61", "3456.60"),
)


async def _stored_tender(session):
    owner = User(email=f"wb-{uuid.uuid4().hex[:8]}@example.test", hashed_password="x", full_name="Estimator")
    session.add(owner)
    await session.flush()
    project = Project(name="Logistics hub", owner_id=owner.id, currency="CHF")
    session.add(project)
    await session.flush()
    boq = BOQ(project_id=project.id, name="Main bill")
    session.add(boq)
    await session.flush()
    header = Position(
        boq_id=boq.id, ordinal="01", description="Groundworks", unit="", quantity="0", unit_rate="0", total="9015.40"
    )
    session.add(header)
    await session.flush()
    lines = [
        Position(boq_id=boq.id, parent_id=header.id, ordinal=o, description=d, unit=u, quantity=q, unit_rate=r, total=t)
        for o, d, u, q, r, t in LINES
    ]
    session.add_all(lines)
    await session.flush()
    return owner, project, boq, header, lines


def _bid(package, name: str, rates: dict, total: str) -> TenderBid:
    return TenderBid(
        package_id=package.id,
        company_name=name,
        contact_email="",
        total_amount=total,
        currency="CHF",
        status="submitted",
        line_items=[
            {
                "position_id": str(pos.id),
                "description": pos.description,
                "unit": pos.unit,
                "quantity": float(pos.quantity),
                "unit_rate": rate,
                "total": float(Decimal(rate) * Decimal(pos.quantity)),
            }
            for pos, rate in rates.items()
        ],
        metadata_={},
    )


async def _body(response) -> bytes:
    chunks = [chunk async for chunk in response.body_iterator]
    return b"".join(c if isinstance(c, bytes) else c.encode() for c in chunks)


def _row_of(ws, description: str) -> int:
    for row in ws.iter_rows():
        if len(row) > 1 and row[1].value == description:
            return row[1].row
    raise AssertionError(f"no row for {description!r}")


def _column_of(ws, header: str) -> int:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(header):
                return cell.column
    raise AssertionError(f"no column for {header!r}")


@pytest.mark.asyncio
async def test_the_comparison_workbook_says_what_the_comparison_endpoint_says(pg_session) -> None:
    owner, project, boq, _header, lines = await _stored_tender(pg_session)
    package = TenderPackage(
        project_id=project.id, boq_id=boq.id, name="Groundworks lot", status="evaluating", metadata_={}
    )
    pg_session.add(package)
    await pg_session.flush()
    excavation, blinding, formwork = lines
    # Beta leaves the blinding out and states a total of its own.
    pg_session.add_all(
        [
            _bid(package, "Alpha Bau", {excavation: "17.00", blinding: "40.00", formwork: "60.00"}, "8840.00"),
            _bid(package, "Beta Tiefbau", {excavation: "19.50", formwork: "52.00"}, "5500.00"),
        ]
    )
    await pg_session.flush()

    service = TenderingService(pg_session)
    comparison = await service.compare_bids(package.id)
    # Looked up by firm, not by column: bids from one transaction share their
    # creation time on PostgreSQL, so position in the list says nothing here.
    blinding_row = next(r for r in comparison.rows if r.description == "Blinding concrete")
    beta_blinding = next(b for b in blinding_row.bids if b["company_name"] == "Beta Tiefbau")
    assert beta_blinding["priced"] is False, "the endpoint must tell a missing price from a zero"
    assert beta_blinding["unit_rate"] is None, "a missing price carries no figure, not 0"
    totals_by_firm = {t["company_name"]: t for t in comparison.bid_totals}
    alpha_total, beta_total = totals_by_firm["Alpha Bau"], totals_by_firm["Beta Tiefbau"]
    # The header row is not a line; Beta priced two of the three.
    assert (alpha_total["matched_lines"], alpha_total["total_lines"]) == (3, 3)
    assert (beta_total["matched_lines"], beta_total["total_lines"]) == (2, 3)

    response = await export_bid_comparison_xlsx(
        package.id, str(owner.id), {}, pg_session, service, None, locale="de", accept_language="en"
    )
    assert response.headers["content-language"] == "de"
    ws = load_workbook(io.BytesIO(await _body(response))).active
    alpha, beta = _column_of(ws, "Alpha Bau"), _column_of(ws, "Beta Tiefbau")

    totals_row = _row_of(ws, "Angebotssumme laut Angebot")
    for column, bid_total in ((alpha, alpha_total), (beta, beta_total)):
        assert Decimal(str(ws.cell(totals_row, column + 1).value)) == Decimal(str(bid_total["total"]))

    missing = ws.cell(_row_of(ws, "Blinding concrete"), beta)
    assert missing.value == "nicht angeboten"
    assert missing.fill.fgColor.rgb.endswith(MISSING_FILL)
    excavation_row = _row_of(ws, "Excavation")
    assert ws.cell(excavation_row, alpha).fill.fgColor.rgb.endswith(LOWEST_FILL)
    assert not ws.cell(excavation_row, beta).fill.fgColor.rgb.endswith(LOWEST_FILL)
    assert ws.cell(excavation_row, 1).value == "01.010"
    assert ws.cell(excavation_row, 3).value == "m³"

    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "CHF" in text
    assert "EUR" not in text, "a CHF package must not be labelled in a default currency"


@pytest.mark.asyncio
async def test_the_bill_for_bidders_holds_the_package_lines_and_no_contractor_price(pg_session) -> None:
    owner, project, boq, _header, lines = await _stored_tender(pg_session)
    # A package over the first two lines only.
    package = TenderPackage(
        project_id=project.id,
        boq_id=boq.id,
        name="Earth and blinding",
        status="issued",
        deadline="2026-11-30",
        metadata_={"scope_position_ids": [str(lines[0].id), str(lines[1].id)]},
    )
    pg_session.add(package)
    await pg_session.flush()

    response = await export_bill_for_bidders_xlsx(
        package.id,
        str(owner.id),
        {},
        pg_session,
        TenderingService(pg_session),
        None,
        locale=None,
        accept_language="xx-YY",
    )
    assert response.headers["content-language"] == "en", "an unknown language falls back to English and says so"
    data = await _body(response)

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        everything = "".join(archive.read(name).decode("utf-8", "replace") for name in archive.namelist())
    for _o, _d, _u, _q, rate, total in LINES:
        assert rate not in everything, f"contractor rate {rate} leaked"
        assert total not in everything, f"contractor total {total} leaked"
    assert "9015.4" not in everything
    assert "Formwork" not in everything, "a line outside the package is not sent"

    ws = load_workbook(io.BytesIO(data)).active
    excavation = _row_of(ws, "Excavation")
    assert ws.cell(excavation, 4).value == 120
    assert ws.cell(excavation, 5).value is None
    assert ws.cell(excavation, 5).protection.locked is False
    # The letterhead block moved the table down; the formula moved with it.
    assert ws.cell(excavation, 6).value == f'=IF(E{excavation}="","",D{excavation}*E{excavation})'
    _row_of(ws, "Groundworks")
    total_row = _row_of(ws, "Total")
    assert ws.cell(total_row, 6).value == f"=SUM(F{excavation - 1}:F{_row_of(ws, 'Blinding concrete')})"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("route", "language"),
    [(export_bid_comparison_xlsx, True), (export_bill_for_bidders_xlsx, True), (export_package_gaeb_x83, False)],
)
async def test_another_users_package_reads_as_missing(pg_session, route, language) -> None:
    """Each export carries prices or the bill, so a stranger's request is a 404 like the screens'."""
    _owner, project, boq, _header, _lines = await _stored_tender(pg_session)
    package = TenderPackage(project_id=project.id, boq_id=boq.id, name="Private lot", status="evaluating", metadata_={})
    stranger = User(email=f"other-{uuid.uuid4().hex[:8]}@example.test", hashed_password="x", full_name="Other")
    pg_session.add_all([package, stranger])
    await pg_session.flush()

    with pytest.raises(HTTPException) as caught:
        await route(
            package.id,
            str(stranger.id),
            {},
            pg_session,
            TenderingService(pg_session),
            None,
            **({"locale": None, "accept_language": None} if language else {}),
        )
    assert caught.value.status_code == 404


@pytest.mark.asyncio
async def test_the_package_x83_holds_the_lines_the_excel_bill_holds(pg_session) -> None:
    """Both files a bidder receives list the package's lines and nothing else, unpriced."""
    from lxml import etree

    from tests.unit.test_gaeb_export_xsd import _load_schema

    owner, project, boq, _header, lines = await _stored_tender(pg_session)
    package = TenderPackage(
        project_id=project.id,
        boq_id=boq.id,
        name="Earth and blinding",
        status="issued",
        metadata_={"scope_position_ids": [str(lines[0].id), str(lines[1].id)]},
    )
    pg_session.add(package)
    await pg_session.flush()

    response = await export_package_gaeb_x83(
        package.id, str(owner.id), {}, pg_session, TenderingService(pg_session), None
    )
    assert ".X83" in response.headers["content-disposition"]
    xml = (await _body(response)).decode("utf-8")

    doc = etree.fromstring(xml.encode("utf-8"))
    schema = _load_schema("83")
    assert schema.validate(doc), "; ".join(f"{e.line}:{e.message}" for e in schema.error_log[:8])
    names = [etree.QName(el).localname for el in doc.iter() if isinstance(el.tag, str)]
    assert names.count("Item") == 2
    assert "UP" not in names
    assert "IT" not in names
    assert "Excavation" in xml
    assert "Blinding concrete" in xml
    assert "Formwork" not in xml, "a line outside the package is not sent"
    for _o, _d, _u, _q, rate, total in LINES:
        assert rate not in xml, f"contractor rate {rate} leaked"
        assert total not in xml, f"contractor total {total} leaked"
