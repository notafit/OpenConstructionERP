# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Write the Excel 97-2003 copies of the Hungarian test workbooks.

The .xls files under ``tests/fixtures/xls/`` are committed rather than built
at test time, because the only library that writes the format, xlwt, has had
no release since 2017 and is not a dependency of this project. Each one is the
cell-for-cell copy of a workbook the tests build as .xlsx, so the parity tests
can read both and require the same bill out of each.

Change a builder and the parity test for it fails until the copy is written
again::

    pip install xlwt==1.3.0
    cd backend && python -m tests.fixtures.make_hu_xls
"""

from __future__ import annotations

import datetime as dt
import io
from collections.abc import Callable
from pathlib import Path

from openpyxl import load_workbook

from tests.fixtures.hu_boq import flat_xlsx, trade_workbook_xlsx
from tests.unit.test_hungary_workbook_import import (
    _building_workbook,
    _infra_rows,
    _infra_workbook,
    _rows_like_the_template,
)

XLS_DIR = Path(__file__).with_name("xls")

#: File name of each copy, and the builder of the .xlsx it copies.
SOURCES: dict[str, Callable[[], bytes]] = {
    "hu_building.xls": lambda: _building_workbook(_rows_like_the_template()),
    "hu_infrastructure.xls": lambda: _infra_workbook(_infra_rows()),
    "hu_flat.xls": flat_xlsx,
    "hu_trades_hidden.xls": lambda: trade_workbook_xlsx(hidden=True),
}


def to_xls(xlsx: bytes) -> bytes:
    """The same sheets, cells, sheet states and active sheet, as BIFF8."""
    import xlwt

    source = load_workbook(io.BytesIO(xlsx))
    target = xlwt.Workbook(encoding="utf-8")
    date_style = xlwt.easyxf(num_format_str="yyyy-mm-dd")
    for worksheet in source.worksheets:
        sheet = target.add_sheet(worksheet.title, cell_overwrite_ok=True)
        if worksheet.sheet_state != "visible":
            sheet.visibility = 1 if worksheet.sheet_state == "hidden" else 2
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, dt.datetime | dt.date):
                    sheet.write(cell.row - 1, cell.column - 1, cell.value, date_style)
                else:
                    sheet.write(cell.row - 1, cell.column - 1, cell.value)
    target.active_sheet = source.worksheets.index(source.active)
    buffer = io.BytesIO()
    target.save(buffer)
    return buffer.getvalue()


def main() -> None:
    XLS_DIR.mkdir(exist_ok=True)
    for name, build in SOURCES.items():
        (XLS_DIR / name).write_bytes(to_xls(build()))
        print(f"wrote {XLS_DIR / name}")


if __name__ == "__main__":
    main()
