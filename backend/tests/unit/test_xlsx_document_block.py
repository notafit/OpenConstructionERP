# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The document block an exporter asks for, with or without a letterhead.

``apply_company_header(ws, details=[...])`` prints the title and the detail
lines (which project, which bill, which currency, when) above the table
whether or not the company has a profile. The profile only adds the firm's
logo and its lines on top. Without ``details`` the old contract holds: no
profile, no change, which ``test_xlsx_branding`` pins for every other
exporter.

The readers are stubbed exactly as ``test_xlsx_branding`` stubs them, so a
profile saved on this machine cannot colour a result.
"""

from __future__ import annotations

import base64
import io
import zipfile
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PIL import Image as PILImage

from app.core import pdf_branding, xlsx_branding
from app.core.company_profile import DEFAULT_COMPANY_PROFILE
from app.core.sheet_header import HEADER_SEARCH_ROWS
from app.core.xlsx_branding import apply_company_header

_PROFILE = {
    "legal_name": "Acme & Sons Construction GmbH",
    "address": "Hauptstrasse 1\n10115 Berlin",
    "registration_line": "HRB 12345",
    "phone": "+49 30 1234567",
    "email": "office@acme.example",
    "website": "acme.example",
}

_DETAILS = [
    "Project: Riverside HQ  |  Standard: din276  |  Region: DACH",
    "Currency: EUR  |  Exported: 2026-10-04",
]


@pytest.fixture
def use_profile(monkeypatch: pytest.MonkeyPatch):
    def _set(
        profile: dict[str, str] | None = None,
        *,
        branding: dict[str, Any] | None = None,
        appearance: dict[str, Any] | None = None,
    ) -> None:
        full = dict(DEFAULT_COMPANY_PROFILE)
        full.update(profile or {})
        monkeypatch.setattr(pdf_branding, "_read_company_profile", lambda *_a, **_kw: dict(full))
        monkeypatch.setattr(pdf_branding, "_read_branding", lambda *_a, **_kw: dict(branding or {}))
        monkeypatch.setattr(pdf_branding, "_read_appearance", lambda *_a, **_kw: dict(appearance or {}))

    _set()
    return _set


def _png_data_url(width: int = 400, height: int = 100) -> str:
    out = io.BytesIO()
    PILImage.new("RGB", (width, height), (20, 60, 140)).save(out, format="PNG")
    return "data:image/png;base64," + base64.b64encode(out.getvalue()).decode("ascii")


def _sample() -> tuple[Workbook, Any]:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    for col, header in enumerate(("Pos.", "Description", "Unit", "Quantity", "Unit Rate", "Total"), 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    ws.append(["01.001", "Excavation", "m3", 120, 18.5, 2220])
    ws.append(["01.002", "Backfill", "m3", 80, 9.0, 720])
    ws["F4"] = "=SUM(F2:F3)"
    ws.freeze_panes = "A2"
    return wb, ws


def _save(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _parts(blob: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name != "docProps/core.xml"}


def _above(ws: Any, start: int) -> list[Any]:
    """Every value above the table, top to bottom."""
    return [cell.value for row in ws.iter_rows(min_row=1, max_row=start - 1) for cell in row if cell.value]


# -- Without a company profile -------------------------------------------------


def test_no_profile_still_prints_the_title_and_the_details(use_profile) -> None:
    wb, ws = _sample()

    start = apply_company_header(ws, title="Tender BOQ", details=_DETAILS)
    reloaded = load_workbook(io.BytesIO(_save(wb)))["BOQ"]

    assert start > 1
    assert _above(reloaded, start) == ["Tender BOQ", *_DETAILS]
    # The table is whole, one block lower, and the pane and the formula moved with it.
    assert [reloaded.cell(row=start, column=c).value for c in (1, 2)] == ["Pos.", "Description"]
    assert reloaded.cell(row=start + 1, column=2).value == "Excavation"
    assert reloaded.freeze_panes == f"A{start + 1}"
    assert reloaded.cell(row=start + 3, column=6).value == f"=SUM(F{start + 1}:F{start + 2})"


def test_no_profile_block_carries_no_company_at_all(use_profile) -> None:
    wb, ws = _sample()

    start = apply_company_header(ws, title="Tender BOQ", details=_DETAILS)

    assert ws._images == []
    # No rule: the thin line separates a company block, and there is none.
    for row in range(1, start):
        for col in range(1, 7):
            assert ws.cell(row=row, column=col).border.bottom.style is None, (row, col)
    assert ws.oddHeader.left.text in (None, "")
    assert ws.oddFooter.left.text in (None, "")
    assert ws.oddHeader.right.text == "Tender BOQ"
    assert ws.oddFooter.right.text == "&P / &N"


def test_no_profile_and_no_details_is_still_identical(use_profile) -> None:
    # Every exporter that does not ask for a block keeps the old promise.
    untouched, _ = _sample()
    branded, ws = _sample()

    assert apply_company_header(ws, title="Tender BOQ", details=[]) == 1
    assert apply_company_header(ws, title="Tender BOQ", details=["", "   "]) == 1
    assert _parts(_save(branded)) == _parts(_save(untouched))


def test_letterhead_switched_off_still_prints_the_block(use_profile) -> None:
    use_profile(_PROFILE, appearance={"show_letterhead": False})
    wb, ws = _sample()

    start = apply_company_header(ws, title="Tender BOQ", details=_DETAILS)

    values = _above(ws, start)
    assert values == ["Tender BOQ", *_DETAILS]
    assert _PROFILE["legal_name"] not in values


# -- With a company profile ----------------------------------------------------


def test_profile_puts_the_company_on_top_of_the_block(use_profile) -> None:
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url()})
    wb, ws = _sample()

    start = apply_company_header(ws, title="Tender BOQ", details=_DETAILS)
    reloaded = load_workbook(io.BytesIO(_save(wb)))["BOQ"]

    values = _above(reloaded, start)
    assert values[0] == _PROFILE["legal_name"]
    assert values[-3:] == ["Tender BOQ", *_DETAILS]
    assert len(reloaded._images) == 1
    assert reloaded.freeze_panes == f"A{start + 1}"
    assert reloaded.oddFooter.left.text == "Acme && Sons Construction GmbH"


def test_profile_without_details_is_what_it_always_was(use_profile) -> None:
    use_profile(_PROFILE)
    wb_old, ws_old = _sample()
    wb_new, ws_new = _sample()

    assert apply_company_header(ws_old, title="Tender BOQ") == apply_company_header(
        ws_new, title="Tender BOQ", details=None
    )
    assert _parts(_save(wb_old)) == _parts(_save(wb_new))


def test_a_failed_letterhead_still_leaves_the_block(use_profile, monkeypatch: pytest.MonkeyPatch) -> None:
    # A logo that breaks the build costs the company part, never the line that
    # says which project and which bill the sheet is.
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url()})

    def _boom(*_args: Any, **_kwargs: Any) -> Any:
        raise RuntimeError("logo failed")

    monkeypatch.setattr(xlsx_branding, "_logo_picture", _boom)
    wb, ws = _sample()

    start = apply_company_header(ws, title="Tender BOQ", details=_DETAILS)

    assert _above(ws, start) == ["Tender BOQ", *_DETAILS]
    assert ws._images == []
    assert ws.cell(row=start, column=1).value == "Pos."


def test_a_block_that_cannot_be_written_leaves_the_sheet_as_it_was(
    use_profile, monkeypatch: pytest.MonkeyPatch
) -> None:
    use_profile(_PROFILE)
    untouched, _ = _sample()
    branded, ws = _sample()

    def _boom(*_args: Any, **_kwargs: Any) -> Any:
        raise RuntimeError("write failed")

    monkeypatch.setattr(xlsx_branding, "_write_letterhead", _boom)

    assert apply_company_header(ws, title="Tender BOQ", details=_DETAILS) == 1
    assert _parts(_save(branded)) == _parts(_save(untouched))


# -- The lines themselves ------------------------------------------------------


def test_detail_lines_are_text_never_formulas(use_profile) -> None:
    wb, ws = _sample()

    start = apply_company_header(ws, title="=HYPERLINK(1)", details=["=1+1", "+49 30 1"])
    blob = _save(wb)
    reloaded = load_workbook(io.BytesIO(blob))["BOQ"]

    assert _above(reloaded, start) == ["=HYPERLINK(1)", "=1+1", "+49 30 1"]
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        # The only formula on the sheet is the table's own SUM.
        assert archive.read("xl/worksheets/sheet1.xml").count(b"<f>") == 1


def test_blank_detail_lines_are_dropped_and_whitespace_folded(use_profile) -> None:
    wb, ws = _sample()

    start = apply_company_header(ws, title=None, details=["", "  Project:\n  Riverside  ", None])  # type: ignore[list-item]

    assert _above(ws, start) == ["Project: Riverside"]


def test_the_block_never_pushes_the_header_out_of_the_importers_reach(use_profile) -> None:
    # The importers look for the header in the first rows only; the tallest
    # letterhead plus the most detail lines a caller can ask for stays inside.
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url()})
    wb, ws = _sample()
    ws.column_dimensions["A"].width = 120  # forces the logo onto a row of its own

    start = apply_company_header(ws, title="T", subtitle="S", details=[f"line {i}" for i in range(10)])

    assert start <= HEADER_SEARCH_ROWS
    shown = [v for v in _above(ws, start) if isinstance(v, str) and v.startswith("line ")]
    assert shown == [f"line {i}" for i in range(xlsx_branding.MAX_DETAIL_LINES)]


def test_the_bill_reimports_under_a_block_without_a_profile(use_profile) -> None:
    from app.modules.boq import router as boq_router
    from app.modules.boq.importers import excel as boq_excel

    wb, ws = _sample()
    apply_company_header(ws, title="Description", details=_DETAILS)
    blob = _save(wb)

    for parse in (boq_excel._parse_rows_from_excel, boq_router._parse_rows_from_excel):
        rows, meta = parse(blob)
        assert [row.get("description") for row in rows][:2] == ["Excavation", "Backfill"]
        assert meta["original_columns"][:2] == ["Pos.", "Description"]
