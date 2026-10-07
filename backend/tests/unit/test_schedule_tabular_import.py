# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Unit tests for the spreadsheet schedule import parser (pure, no database).

Covers header recognition in ten languages, the day/month order of text dates,
duration and predecessor grammar, outline sources, limits, CSV dialects, .xlsx
workbooks built with openpyxl in ``tmp_path``, and the contract that matters for
wiring: the parsed document goes through the interchange parser, validator and
cleaner untouched.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest

from app.modules.schedule import tabular_headers
from app.modules.schedule import tabular_import as ti
from app.modules.schedule.schedule_clean import clean_document
from app.modules.schedule.schedule_interchange import parse_document, validate_document
from app.modules.schedule.tabular_headers import (
    FIELDS,
    HEADER_SYNONYMS,
    LANGUAGES,
    TEMPLATE_HEADERS,
    match_header,
)
from app.modules.schedule.tabular_import import preview

BACKEND = Path(ti.__file__).resolve().parents[3]
FIXTURES = BACKEND / "tests" / "fixtures" / "schedule_tabular"
ROUTER = BACKEND / "app" / "modules" / "schedule" / "router.py"


# ── helpers ──────────────────────────────────────────────────────────────────


def _csv(rows: list[list[Any]], *, delimiter: str = ",") -> bytes:
    out = io.StringIO()
    csv.writer(out, delimiter=delimiter, lineterminator="\n").writerows(rows)
    return out.getvalue().encode("utf-8")


def _codes(result: ti.TabularPreview, severity: str | None = None) -> list[str]:
    return [i.code for i in result.issues if severity is None or i.severity == severity]


def _issue(result: ti.TabularPreview, code: str) -> ti.Issue:
    found = [i for i in result.issues if i.code == code]
    assert found, f"no {code} in {_codes(result)}"
    return found[0]


def _acts(result: ti.TabularPreview) -> dict[str, dict[str, Any]]:
    assert result.document is not None
    return {a["ref"]: a for a in result.document["activities"]}


def _ordered(result: ti.TabularPreview) -> list[dict[str, Any]]:
    assert result.document is not None
    return list(result.document["activities"])


def _nth(result: ti.TabularPreview, index: int) -> dict[str, Any]:
    return _ordered(result)[index]


def _links(result: ti.TabularPreview) -> set[tuple[str, str, str, int]]:
    assert result.document is not None
    return {
        (r["predecessor_ref"], r["successor_ref"], r["relationship_type"], r["lag_days"])
        for r in result.document["relationships"]
    }


def _assert_interchange_clean(result: ti.TabularPreview) -> None:
    """The document imports through the interchange path with nothing to repair."""
    parsed = parse_document(result.document)
    assert validate_document(parsed) == []
    cleaned = clean_document(parsed)
    assert cleaned.actions == []


def _simple(header: list[str], rows: list[list[Any]], **kwargs: Any) -> ti.TabularPreview:
    return preview(_csv([header, *rows]), "plan.csv", **kwargs)


def _xlsx(tmp_path: Path, rows: list[list[Any]], *, indents: dict[int, int] | None = None, **extra: Any) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = extra.get("title", "Schedule")
    for row in rows:
        sheet.append(row)
    for row_number, indent in (indents or {}).items():
        sheet.cell(row=row_number, column=extra.get("name_column", 2)).alignment = Alignment(indent=indent)
    for title, other_rows, state in extra.get("sheets", ()):
        other = workbook.create_sheet(title, 0 if extra.get("others_first") else None)
        for row in other_rows:
            other.append(row)
        other.sheet_state = state
    path = tmp_path / "plan.xlsx"
    workbook.save(path)
    return path.read_bytes()


# ── header vocabulary ────────────────────────────────────────────────────────


@pytest.mark.parametrize("lang", LANGUAGES)
def test_template_headers_map_back_exactly_in_every_language(lang: str) -> None:
    for header, field in zip(TEMPLATE_HEADERS[lang], FIELDS, strict=True):
        found = match_header(header)
        assert (found.field, found.confidence) == (field, 1.0), (lang, header, found)


@pytest.mark.parametrize("lang", LANGUAGES)
def test_every_listed_synonym_maps_to_its_field(lang: str) -> None:
    for field in FIELDS:
        for header in HEADER_SYNONYMS[field][lang]:
            found = match_header(header)
            assert found.field == field, (lang, header, found)
            assert found.confidence in (1.0, 0.8)


@pytest.mark.parametrize(
    ("header", "field"),
    [
        ("Task Name", "name"),
        ("Vorgaenger", "predecessors"),
        ("Fecha de inicio", "start"),
        ("Date de fin", "finish"),
        ("Продолжительность", "duration"),
        ("Antecessoras", "predecessors"),
        ("Data inizio", "start"),
        ("Einde", "finish"),
        ("Dni", "duration"),
        ("Bitiş Tarihi", "finish"),
    ],
)
def test_synonyms_read_at_point_eight(header: str, field: str) -> None:
    found = match_header(header)
    assert (found.field, found.confidence, found.tier) == (field, 0.8, "synonym")


@pytest.mark.parametrize(
    ("header", "field"),
    [("Start date planned", "start"), ("Duraton", "duration"), ("Cost code", None), ("Planned start date", "start")],
)
def test_fuzzy_tier(header: str, field: str | None) -> None:
    found = match_header(header)
    assert found.field == field
    if field:
        assert found.confidence == 0.5


@pytest.mark.parametrize(
    "header",
    ["Actual Start", "Baseline Finish", "Remaining Duration", "Unique ID", "Total Float", "Critical", "Activity Type"],
)
def test_other_measurements_of_the_same_thing_stay_unmapped(header: str) -> None:
    assert match_header(header).field is None


def test_vocabulary_refuses_a_header_that_reads_two_ways(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(HEADER_SYNONYMS["notes"], "en", (*HEADER_SYNONYMS["notes"]["en"], "Start"))
    with pytest.raises(ValueError, match="reads as both"):
        tabular_headers._build_header_index()


def test_link_type_aliases_are_unambiguous() -> None:
    aliases = tabular_headers.RELATIONSHIP_TYPE_ALIASES
    assert set(aliases.values()) == {"FS", "SS", "FF", "SF"}
    assert all(len(alias) == 2 for alias in aliases)


def test_header_collision_keeps_the_exact_column_and_reports_the_other() -> None:
    result = _simple(["Task Name", "Name", "Start"], [["x", "Pour slab", "2026-05-04"]])
    assert result.mapping["name"] == 1
    collision = _issue(result, "header_collision")
    assert collision.severity == "warning"
    assert collision.params["field"] == "name"
    assert collision.params["ignored"] == [0]
    assert _nth(result, 0)["name"] == "Pour slab"


def test_german_header_row_maps_every_column() -> None:
    result = preview((FIXTURES / "ambiguous_de.csv").read_bytes(), "plan.csv", date_order="dmy")
    assert result.encoding == "cp1252"
    assert result.delimiter == ";"
    assert result.mapping == {
        "id": 0,
        "name": 1,
        "duration": 2,
        "start": 3,
        "finish": 4,
        "predecessors": 5,
        "milestone": 6,
        "client_visible": 7,
        "notes": 8,
    }


def test_explicit_column_mapping_overrides_and_unmaps() -> None:
    header = ["Code", "Text", "Begin", "Comments"]
    rows = [["A", "Excavate", "2026-05-04", "deep"]]
    detected = _simple(header, rows)
    assert detected.mapping.get("name") is None or detected.mapping["name"] == 3
    result = _simple(header, rows, column_mapping={"1": "name", "3": ""})
    assert result.mapping["name"] == 1
    assert "notes" not in result.mapping
    assert result.columns[1].tier == "override"
    assert _acts(result)["A"]["name"] == "Excavate"
    assert _acts(result)["A"]["description"] == ""


@pytest.mark.parametrize("mapping", [{"9": "name"}, {"0": "colour"}, {"0": "name", "1": "name"}])
def test_invalid_column_mapping_is_an_error(mapping: dict[str, str]) -> None:
    result = _simple(["ID", "Name"], [["1", "A"]], column_mapping=mapping)
    assert "column_mapping_invalid" in _codes(result, "error")


def test_description_column_stands_in_for_a_missing_name() -> None:
    result = _simple(["ID", "Description", "Start"], [["1", "Survey", "2026-05-04"]])
    assert "name_from_description" in _codes(result, "info")
    assert _nth(result, 0)["name"] == "Survey"


def test_no_name_column_is_an_error() -> None:
    result = _simple(["ID", "Start", "Finish"], [["1", "2026-05-04", "2026-05-05"]])
    assert "name_column_missing" in _codes(result, "error")


# ── dates ────────────────────────────────────────────────────────────────────


def test_day_first_is_detected_from_a_day_above_twelve() -> None:
    result = _simple(["Name", "Start"], [["A", "03/04/2026"], ["B", "25/04/2026"]])
    assert result.date_order == "dmy"
    assert result.date_order_source == "values"
    assert not result.has_errors
    assert _nth(result, 0)["start_date"] == "2026-04-03"


def test_month_first_is_detected_from_a_day_above_twelve() -> None:
    result = _simple(["Name", "Start"], [["A", "03/04/2026"], ["B", "04/25/2026"]])
    assert result.date_order == "mdy"
    assert _nth(result, 0)["start_date"] == "2026-03-04"


def test_ambiguous_dates_need_confirmation_with_a_suggestion() -> None:
    result = preview((FIXTURES / "ambiguous_de.csv").read_bytes(), "plan.csv")
    issue = _issue(result, "date_order_unconfirmed")
    assert issue.severity == "error"
    assert issue.params["suggested"] == "dmy"
    assert issue.params["suggested_by"] == "durations"
    assert result.date_order_source == "suggested"


def test_suggestion_follows_the_durations_for_month_first_files() -> None:
    rows = [["A", "05/04/2026", "05/08/2026", "5"], ["B", "06/01/2026", "06/02/2026", "2"]]
    result = _simple(["Name", "Start", "Finish", "Duration"], rows)
    assert _issue(result, "date_order_unconfirmed").params["suggested"] == "mdy"


def test_confirmed_order_resolves_ambiguous_dates() -> None:
    result = preview((FIXTURES / "ambiguous_de.csv").read_bytes(), "plan.csv", date_order="dmy")
    assert not result.has_errors, result.issues
    acts = _acts(result)
    assert (acts["10"]["start_date"], acts["10"]["end_date"]) == ("2026-05-04", "2026-05-08")
    assert result.date_order_source == "explicit"
    _assert_interchange_clean(result)


def test_conflicting_orders_are_an_error() -> None:
    result = _simple(["Name", "Start"], [["A", "25/04/2026"], ["B", "04/25/2026"]])
    assert "date_order_conflict" in _codes(result, "error")


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        ("2026-05-04", "2026-05-04"),
        ("2026-05-04 08:00", "2026-05-04"),
        ("2026-05-04T08:00:00", "2026-05-04"),
        ("Mon 5/4/26 8:00 AM", "2026-05-04"),
        ("May 4, 2026", "2026-05-04"),
        ("4 May 2026", "2026-05-04"),
        ("04-May-26", "2026-05-04"),
    ],
)
def test_text_date_forms(cell: str, expected: str) -> None:
    result = _simple(["Name", "Start"], [["A", cell]], date_order="mdy")
    assert _nth(result, 0)["start_date"] == expected, result.issues


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        ("1. März 2026", "2026-03-01"),
        ("1 Mrz 2026", "2026-03-01"),
        ("12. Dez. 2026", "2026-12-12"),
        ("3. Jänner 2027", "2027-01-03"),
        ("Montag, 4. Mai 2026", "2026-05-04"),
        ("1er mars 2026", "2026-03-01"),
        ("15 févr. 2026", "2026-02-15"),
        ("4 août 2026", "2026-08-04"),
        ("lundi 7 décembre 2026", "2026-12-07"),
        ("12 de marzo de 2026", "2026-03-12"),
        ("1 ene 2026", "2026-01-01"),
        ("3 dic. 2026", "2026-12-03"),
        ("4 maggio 2026", "2026-05-04"),
        ("10 giu 2026", "2026-06-10"),
        ("4 mei 2026", "2026-05-04"),
        ("1 mrt 2026", "2026-03-01"),
        ("15 okt. 2026", "2026-10-15"),
        ("4 maja 2026 r.", "2026-05-04"),
        ("1 marca 2026", "2026-03-01"),
        ("22 października 2026", "2026-10-22"),
        ("5 paź 2026", "2026-10-05"),
        ("1 styczeń 2026", "2026-01-01"),
        ("1º de março de 2026", "2026-03-01"),
        ("4 de maio de 2026", "2026-05-04"),
        ("15 set 2026", "2026-09-15"),
        ("2 out. 2026", "2026-10-02"),
        ("1 марта 2026", "2026-03-01"),
        ("1 марта 2026 г.", "2026-03-01"),
        ("1 марта 2026г.", "2026-03-01"),
        ("14 мая 2026", "2026-05-14"),
        ("3 сент. 2026", "2026-09-03"),
        ("5 янв 2026", "2026-01-05"),
        ("31 декабря 2026", "2026-12-31"),
        ("1 май 2026", "2026-05-01"),
        ("15-фев-26", "2026-02-15"),
        ("пн, 4 мая 2026", "2026-05-04"),
        ("March 1st, 2026", "2026-03-01"),
        ("MÄRZ 1, 2026", "2026-03-01"),
    ],
)
def test_month_names_in_eight_more_languages(cell: str, expected: str) -> None:
    result = _simple(["Name", "Start"], [["A", cell]])
    assert _nth(result, 0)["start_date"] == expected, result.issues
    assert "date_order_ambiguous" not in _codes(result)


@pytest.mark.parametrize("lang", sorted(tabular_headers.MONTH_NAMES))
def test_every_listed_month_word_reads_as_its_month(lang: str) -> None:
    months = tabular_headers.MONTH_NAMES[lang]
    assert len(months) == 12
    for number, words in enumerate(months, start=1):
        for word in words:
            assert tabular_headers.month_number(word) == number, (lang, word)
            assert tabular_headers.month_number(word.upper() + ".") == number, (lang, word)


def test_a_word_that_is_no_month_still_fails_the_date() -> None:
    assert tabular_headers.month_number("days") is None
    result = _simple(["Name", "Start"], [["A", "12 Tage 2026"]])
    assert "date_invalid" in _codes(result, "error")


def test_unreadable_date_is_an_error_on_its_cell() -> None:
    result = _simple(["Name", "Start"], [["A", "next week"]])
    issue = _issue(result, "date_invalid")
    assert (issue.row, issue.column) == (2, 1)


def test_xlsx_date_cells_need_no_order(tmp_path: Path) -> None:
    data = _xlsx(
        tmp_path,
        [["ID", "Name", "Start", "Finish"], [1, "Slab", datetime(2026, 5, 4), date(2026, 5, 8)]],
    )
    result = preview(data, "plan.xlsx")
    assert result.file_format == "xlsx"
    assert not result.has_errors, result.issues
    assert result.date_order is None
    act = _nth(result, 0)
    assert (act["start_date"], act["end_date"], act["duration_days"]) == ("2026-05-04", "2026-05-08", 5)


def test_serial_day_numbers_are_converted_with_a_warning(tmp_path: Path) -> None:
    data = _xlsx(tmp_path, [["Name", "Start"], ["Slab", 46146]])
    result = preview(data, "plan.xlsx")
    assert _nth(result, 0)["start_date"] == "2026-05-04"
    assert "date_serial_converted" in _codes(result, "warning")


# ── durations ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("cell", "days", "warning"),
    [
        ("5", 5, None),
        ("5d", 5, None),
        ("5 days", 5, None),
        ("5 days?", 5, None),
        ("5 Tage", 5, None),
        ("5 дн", 5, None),
        ("5 дней", 5, None),
        ("5 jours", 5, None),
        ("5 días", 5, None),
        ("2w", 10, None),
        ("2 Wochen", 10, None),
        ("40h", 5, "duration_hours_converted"),
        ("12 hrs", 2, "duration_rounded"),
        ("2,5", 3, "duration_rounded"),
    ],
)
def test_duration_forms(cell: str, days: int, warning: str | None) -> None:
    result = _simple(["Name", "Duration"], [["A", cell]])
    assert _nth(result, 0)["duration_days"] == days
    assert not result.has_errors, result.issues
    if warning:
        assert warning in _codes(result, "warning")


@pytest.mark.parametrize(
    ("cell", "code"),
    [
        ("5ed", "duration_elapsed_unsupported"),
        ("5 edays", "duration_elapsed_unsupported"),
        ("3 KT", "duration_elapsed_unsupported"),
        ("2mo", "duration_unit_unsupported"),
        ("-2", "duration_negative"),
        ("soon", "duration_invalid"),
        ("99999", "duration_invalid"),
    ],
)
def test_duration_errors(cell: str, code: str) -> None:
    result = _simple(["Name", "Duration"], [["A", cell]])
    issue = _issue(result, code)
    assert (issue.severity, issue.row, issue.column) == ("error", 2, 1)


def test_header_unit_applies_to_bare_numbers() -> None:
    result = _simple(["Name", "Dauer [Std]"], [["A", "16"]])
    assert _nth(result, 0)["duration_days"] == 2
    assert "duration_hours_converted" in _codes(result)


def test_duration_wins_when_dates_disagree() -> None:
    result = _simple(["Name", "Start", "Finish", "Duration"], [["A", "2026-05-04", "2026-05-06", "5"]])
    issue = _issue(result, "duration_mismatch")
    assert issue.severity == "warning"
    assert issue.params["span"] == 3
    act = _nth(result, 0)
    assert (act["start_date"], act["end_date"], act["duration_days"]) == ("2026-05-04", "2026-05-08", 5)


def test_missing_date_is_derived_from_the_duration() -> None:
    result = _simple(
        ["Name", "Start", "Finish", "Duration"],
        [["A", "2026-05-08", "", "3"], ["B", "", "2026-05-08", "3"], ["C", "2026-05-04", "2026-05-15", ""]],
    )
    acts = _ordered(result)
    assert acts[0]["end_date"] == "2026-05-12"  # Fri + weekend
    assert acts[1]["start_date"] == "2026-05-06"
    assert acts[2]["duration_days"] == 10
    assert not result.has_errors


def test_finish_before_start_is_an_error() -> None:
    result = _simple(["Name", "Start", "Finish"], [["A", "2026-05-08", "2026-05-04"]])
    assert "finish_before_start" in _codes(result, "error")


def test_the_project_working_week_is_used() -> None:
    # Sunday to Thursday: 2026-05-03 is a Sunday, five working days end on Thursday.
    rows = [["A", "2026-05-03", "", "5"], ["B", "2026-05-03", "2026-05-07", ""]]
    result = _simple(["Name", "Start", "Finish", "Duration"], rows, work_weekdays={6, 0, 1, 2, 3})
    acts = _ordered(result)
    assert acts[0]["end_date"] == "2026-05-07"
    assert acts[1]["duration_days"] == 5


@pytest.mark.parametrize(
    "weekdays", [frozenset({0, 1, 2, 3, 4}), frozenset({6, 0, 1, 2, 3}), frozenset({0, 1, 2, 3, 4, 5})]
)
def test_finish_from_and_count_are_inverse(weekdays: frozenset[int]) -> None:
    week = ti._Week(weekdays, frozenset({date(2026, 5, 14)}))
    start = date(2026, 5, 1)
    for offset in range(10):
        for days in range(1, 25):
            begin = start + timedelta(days=offset)
            assert week.count(begin, week.finish_from(begin, days)) == days
            finish = begin + timedelta(days=40)
            assert week.count(week.start_from(finish, days), finish) == days


# ── predecessors ─────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("token", "link_type", "lag"),
    [
        ("A10", "FS", 0),
        ("A10FS", "FS", 0),
        ("A10SS", "SS", 0),
        ("A10FF", "FF", 0),
        ("A10SF", "SF", 0),
        ("A10FS+2d", "FS", 2),
        ("A10 FS + 2 days", "FS", 2),
        ("A10FS-3d", "FS", -3),
        ("A10 SS -1", "SS", -1),
        ("A10+1w", "FS", 5),
        ("A10EA", "FS", 0),
        ("A10AA+2 Tage", "SS", 2),
        ("A10EE", "FF", 0),
        ("A10AE", "SF", 0),
        ("A10FD", "FS", 0),
        ("A10DD", "SS", 0),
        ("A10FC", "FS", 0),
        ("A10CC", "SS", 0),
        ("A10ОН", "FS", 0),
        ("A10НН+2д", "SS", 2),
        ("A10ОО", "FF", 0),
        ("A10НО", "SF", 0),
        ("A10ON", "FS", 0),
        ("A10NN", "SS", 0),
        ("a10ss", "SS", 0),
    ],
)
def test_predecessor_grammar(token: str, link_type: str, lag: int) -> None:
    result = _simple(["ID", "Name", "Predecessors"], [["A10", "First", ""], ["B20", "Second", token]])
    assert not result.has_errors, result.issues
    assert _links(result) == {("A10", "B20", link_type, lag)}


def test_whole_token_is_tried_as_an_id_before_a_type_suffix() -> None:
    rows = [["A10", "a", ""], ["A10FS", "b", ""], ["C", "c", "A10FS; A10FS+1d"]]
    result = _simple(["ID", "Name", "Predecessors"], rows)
    assert _links(result) == {("A10FS", "C", "FS", 0)}
    assert "duplicate_predecessor" in _codes(result, "warning")


def test_hyphenated_ids_with_leads() -> None:
    rows = [["A-10", "a", ""], ["B", "b", "A-10-2d"], ["C", "c", "A-10FS"]]
    result = _simple(["ID", "Name", "Predecessors"], rows)
    assert _links(result) == {("A-10", "B", "FS", -2), ("A-10", "C", "FS", 0)}


def test_leads_survive_into_the_document() -> None:
    result = _simple(["ID", "Name", "Predecessors"], [["1", "a", ""], ["2", "b", "1FS-2d"]])
    assert result.document["relationships"][0]["lag_days"] == -2
    _assert_interchange_clean(result)


def test_without_an_id_column_predecessors_are_sheet_row_numbers() -> None:
    rows = [["Survey", ""], ["Excavate", "2FS+2d"], ["Pour", "3SS, 2"]]
    result = _simple(["Name", "Predecessors"], rows)
    assert _links(result) == {("2", "3", "FS", 2), ("3", "4", "SS", 0), ("2", "4", "FS", 0)}
    assert "predecessors_by_row_number" in _codes(result, "info")


@pytest.mark.parametrize(
    ("token", "code"),
    [
        ("ZZ", "predecessor_unknown"),
        ("A10FS+2ed", "lag_elapsed_unsupported"),
        ("A10FS+2mo", "predecessor_invalid"),
        ("B20", "self_dependency"),
    ],
)
def test_predecessor_errors(token: str, code: str) -> None:
    result = _simple(["ID", "Name", "Predecessors"], [["A10", "First", ""], ["B20", "Second", token]])
    issue = _issue(result, code)
    assert (issue.severity, issue.row, issue.column, issue.params["token"]) == ("error", 3, 2, token)


def test_lag_in_hours_is_converted_with_a_warning() -> None:
    result = _simple(["ID", "Name", "Predecessors"], [["A", "a", ""], ["B", "b", "A+16h"]])
    assert _links(result) == {("A", "B", "FS", 2)}
    assert "lag_hours_converted" in _codes(result, "warning")


def test_dependency_loop_is_reported_with_its_cycle() -> None:
    rows = [["A", "a", "C"], ["B", "b", "A"], ["C", "c", "B"], ["D", "d", "C"], ["E", "e", ""]]
    result = _simple(["ID", "Name", "Predecessors"], rows)
    loops = [i for i in result.issues if i.code == "dependency_loop"]
    assert len(loops) == 1
    cycle = loops[0].params["cycle"]
    assert cycle[0] == cycle[-1]
    assert set(cycle) == {"A", "B", "C"}
    assert loops[0].severity == "error"


def test_two_loops_are_both_reported() -> None:
    rows = [["A", "a", "B"], ["B", "b", "A"], ["C", "c", "D"], ["D", "d", "C"]]
    result = _simple(["ID", "Name", "Predecessors"], rows)
    cycles = [frozenset(i.params["cycle"]) for i in result.issues if i.code == "dependency_loop"]
    assert sorted(cycles, key=sorted) == [frozenset({"A", "B"}), frozenset({"C", "D"})]


def test_a_chain_without_a_loop_passes() -> None:
    rows = [[str(n), f"t{n}", str(n - 1) if n > 1 else ""] for n in range(1, 200)]
    result = _simple(["ID", "Name", "Predecessors"], rows)
    assert "dependency_loop" not in _codes(result)
    assert result.relationship_count == 198


# ── outline ──────────────────────────────────────────────────────────────────


def _parents(result: ti.TabularPreview) -> dict[str, str | None]:
    return {ref: act["parent_ref"] for ref, act in _acts(result).items()}


def test_outline_from_a_level_column() -> None:
    rows = [["1", "Project", "1"], ["2", "Phase", "2"], ["3", "Task", "3"], ["4", "Phase 2", "2"], ["5", "Jump", "4"]]
    result = _simple(["ID", "Name", "Outline Level"], rows)
    assert result.outline_source == "outline_level"
    assert _parents(result) == {"1": None, "2": "1", "3": "2", "4": "1", "5": "4"}
    assert "outline_level_jump" in _codes(result, "warning")


def test_outline_from_dotted_wbs() -> None:
    rows = [["Site", "1"], ["Clear", "1.1"], ["Grade", "1.2"], ["Detail", "1.2.1"], ["Build", "2"], ["Frame", "2.1"]]
    result = _simple(["Name", "WBS"], rows)
    assert result.outline_source == "wbs"
    assert _parents(result) == {"2": None, "3": "2", "4": "2", "5": "4", "6": None, "7": "6"}


def test_level_column_beats_wbs() -> None:
    rows = [["Site", "1.1", "1"], ["Clear", "1.2", "1"]]
    result = _simple(["Name", "WBS", "Level"], rows)
    assert result.outline_source == "outline_level"
    assert _parents(result) == {"2": None, "3": None}


def test_outline_from_xlsx_indent(tmp_path: Path) -> None:
    rows = [["ID", "Name"], [1, "Project"], [2, "Phase"], [3, "Task"], [4, "Phase 2"]]
    data = _xlsx(tmp_path, rows, indents={3: 1, 4: 2, 5: 1})
    result = preview(data, "plan.xlsx")
    assert result.outline_source == "indent"
    assert _parents(result) == {"1": None, "2": "1", "3": "2", "4": "1"}


def test_outline_from_leading_spaces() -> None:
    rows = [["Project"], ["  Phase"], ["    Task"], ["  Phase 2"]]
    result = _simple(["Name"], rows)
    assert result.outline_source == "leading_spaces"
    assert _parents(result) == {"2": None, "3": "2", "4": "3", "5": "2"}
    assert _nth(result, 2)["name"] == "Task"


def test_a_space_after_every_delimiter_is_not_an_outline() -> None:
    data = b"Name; Start\nA; 2026-05-04\nB; 2026-05-05\n"
    result = preview(data, "plan.csv")
    assert result.outline_source is None
    assert set(_parents(result).values()) == {None}


# ── milestones, percent, resources, flags ────────────────────────────────────


def test_zero_duration_is_a_milestone() -> None:
    result = _simple(["Name", "Start", "Duration"], [["Handover", "2026-05-04", "0"]])
    act = _nth(result, 0)
    assert (act["activity_type"], act["duration_days"], act["end_date"]) == ("milestone", 0, "2026-05-04")


def test_milestone_flag_drops_a_duration_with_a_warning() -> None:
    rows = [["Topping out", "2026-05-04", "5", "ja"], ["Pour", "2026-05-04", "2", "nein"]]
    result = _simple(["Name", "Start", "Duration", "Meilenstein"], rows)
    acts = _ordered(result)
    assert acts[0]["activity_type"] == "milestone"
    assert acts[0]["duration_days"] == 0
    assert acts[1]["activity_type"] == "task"
    assert "milestone_duration_ignored" in _codes(result, "warning")


def test_milestone_flag_without_duration() -> None:
    result = _simple(["Name", "Start", "Milestone"], [["Permit", "2026-05-04", "yes"]])
    act = _nth(result, 0)
    assert (act["activity_type"], act["duration_days"], act["end_date"]) == ("milestone", 0, "2026-05-04")
    _assert_interchange_clean(result)


@pytest.mark.parametrize(
    ("cells", "expected", "warning"),
    [
        (["45", "100", "0"], ["45", "100", "0"], False),
        (["45%", "12.5 %", ""], ["45", "12.5", "0"], False),
        (["0.45", "1", "0"], ["45", "100", "0"], True),
    ],
)
def test_percent_complete(cells: list[str], expected: list[str], warning: bool) -> None:
    rows = [[f"t{n}", cell] for n, cell in enumerate(cells)]
    result = _simple(["Name", "% Complete"], rows)
    assert [a["progress_pct"] for a in result.document["activities"]] == expected
    assert ("percent_fraction_assumed" in _codes(result, "warning")) is warning
    assert [a["status"] for a in result.document["activities"]][:2] == (
        ["in_progress", "completed"] if expected[1] == "100" else ["in_progress", "in_progress"]
    )


@pytest.mark.parametrize(("cell", "code"), [("150", "percent_out_of_range"), ("half", "percent_invalid")])
def test_percent_errors(cell: str, code: str) -> None:
    result = _simple(["Name", "% Complete"], [["A", cell]])
    assert code in _codes(result, "error")


def test_resources_notes_and_client_visibility() -> None:
    header = ["Name", "Resource Names", "Notes", "Client Visible"]
    rows = [
        ["Pour", "Concrete crew; Pump truck", "Night shift", "yes"],
        ["Cure", "", "", "no"],
        ["Strip", "", "", "maybe"],
    ]
    result = _simple(header, rows)
    acts = _ordered(result)
    assert acts[0]["resources"] == [
        {"name": "Concrete crew", "type": "", "allocation_pct": 100.0, "count": 1},
        {"name": "Pump truck", "type": "", "allocation_pct": 100.0, "count": 1},
    ]
    assert acts[0]["description"] == "Night shift"
    # The sheet's "yes" is a suggestion; nothing in the document is visible until confirmed.
    assert [act["client_visible"] for act in acts] == [False, False, False]
    assert result.client_visible_suggested == [acts[0]["ref"]]
    assert result.to_dict()["client_visible_suggested"] == [acts[0]["ref"]]
    issue = _issue(result, "flag_invalid")
    assert (issue.severity, issue.row, issue.params["field"]) == ("warning", 4, "client_visible")


# ── limits and file checks ───────────────────────────────────────────────────


def test_five_thousand_rows_are_accepted_and_one_more_is_refused() -> None:
    rows = [[f"Task {n}", "1"] for n in range(5000)]
    ok = _simple(["Name", "Duration"], rows)
    assert ok.activity_count == 5000
    assert "too_many_rows" not in _codes(ok)
    over = _simple(["Name", "Duration"], [*rows, ["one more", "1"]])
    assert "too_many_rows" in _codes(over, "error")
    assert over.document is None


def test_far_too_many_rows_stop_reading_early() -> None:
    rows = [[f"Task {n}"] for n in range(6000)]
    result = _simple(["Name"], rows)
    assert _issue(result, "too_many_rows").params["limit"] == 5000


def test_too_many_columns() -> None:
    header = ["Name", *[f"Extra {n}" for n in range(60)]]
    result = _simple(header, [["A", *["x"] * 60]])
    assert "too_many_columns" in _codes(result, "error")


def test_overlong_cell_is_an_error_on_its_cell() -> None:
    result = _simple(["Name", "Notes"], [["A", "x" * 2001], ["B", "y" * 2000]])
    issue = _issue(result, "cell_too_long")
    assert (issue.row, issue.column) == (2, 1)
    assert len([i for i in result.issues if i.code == "cell_too_long"]) == 1


def test_file_size_limit() -> None:
    result = preview(b"Name\n" + b"a\n" * (ti.MAX_FILE_BYTES // 2 + 1), "big.csv")
    assert _codes(result) == ["file_too_large"]
    assert len(result.sha256) == 64


def test_empty_file() -> None:
    assert _codes(preview(b"", "plan.csv")) == ["file_empty"]


def test_binary_files_are_refused() -> None:
    assert _codes(preview(b"%PDF-1.7\n...", "plan.csv")) == ["unsupported_file_type"]
    assert _codes(preview(b"Name\nA\n", "plan.xlsx")) == ["file_type_mismatch"]


def test_broken_workbook_is_an_issue_not_an_exception() -> None:
    result = preview(b"PK\x03\x04" + b"\x00" * 200, "plan.xlsx")
    assert _codes(result) == ["file_unreadable"]


def test_decompression_bomb_is_refused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    data = _xlsx(tmp_path, [["Name"], ["A"]])
    monkeypatch.setattr(ti, "MAX_UNCOMPRESSED_BYTES", 1024)
    assert _codes(preview(data, "plan.xlsx")) == ["file_too_large_uncompressed"]


def test_issue_cap_per_code() -> None:
    rows = [["", "2026-05-04"] for _ in range(120)]
    result = _simple(["Name", "Start"], rows)
    assert len([i for i in result.issues if i.code == "name_missing"]) == ti.MAX_ISSUES_PER_CODE
    assert _issue(result, "issues_truncated").params["by_code"] == {"name_missing": 70}


def test_issue_shape() -> None:
    result = preview((FIXTURES / "broken.csv").read_bytes(), "broken.csv")
    for issue in result.to_dict()["issues"]:
        assert set(issue) == {"code", "severity", "row", "column", "message", "params"}
        assert issue["severity"] in ("error", "warning", "info")


def test_broken_fixture_reports_each_problem_where_it_is() -> None:
    result = preview((FIXTURES / "broken.csv").read_bytes(), "broken.csv")
    found = {(i.code, i.row) for i in result.issues}
    assert ("name_missing", 3) in found
    assert ("date_invalid", 4) in found
    assert ("duration_elapsed_unsupported", 5) in found
    assert ("predecessor_unknown", 5) in found
    assert ("duplicate_id", 6) in found
    loop = _issue(result, "dependency_loop")
    assert set(loop.params["cycle"]) == {"A", "B", "C"}
    assert result.has_errors
    assert result.document is not None


# ── CSV dialects ─────────────────────────────────────────────────────────────


def test_utf8_with_bom() -> None:
    data = "\ufeffName;Start\nBaugrube aushüben;2026-05-04\n".encode()
    result = preview(data, "plan.csv")
    assert result.encoding == "utf-8-sig"
    assert result.mapping["name"] == 0
    assert _nth(result, 0)["name"] == "Baugrube aushüben"


def test_cp1252_semicolon() -> None:
    data = "Vorgangsname;Dauer\nBühne aufbauen;3\n".encode("cp1252")
    result = preview(data, "plan.csv")
    assert (result.encoding, result.delimiter) == ("cp1252", ";")
    assert _nth(result, 0)["name"] == "Bühne aufbauen"


def test_cp1251_russian_header_row_is_read_as_cyrillic() -> None:
    text = (
        "Код;Наименование;Начало;Окончание;Длительность;Предшественники\n"
        "A10;Устройство котлована;04.05.2026;08.05.2026;5;\n"
        "A20;Бетонирование фундамента;11.05.2026;15.05.2026;5;A10\n"
    )
    result = preview(text.encode("cp1251"), "plan.csv", date_order="dmy")
    assert (result.encoding, result.delimiter) == ("cp1251", ";")
    assert result.mapping == {"id": 0, "name": 1, "start": 2, "finish": 3, "duration": 4, "predecessors": 5}
    assert _acts(result)["A10"]["name"] == "Устройство котлована"
    assert _acts(result)["A20"]["start_date"] == "2026-05-11"
    assert not result.has_errors, result.issues


def test_cp1251_short_header_row_with_the_number_sign() -> None:
    result = preview("№;Наименование\n1;Кладка\n".encode("cp1251"), "plan.csv")
    assert result.encoding == "cp1251"
    assert result.mapping == {"id": 0, "name": 1}
    assert _acts(result)["1"]["name"] == "Кладка"


@pytest.mark.parametrize(
    "name",
    ["Кладка стен", "Начало работ", "Изкоп на котлован", "Монтаж металлоконструкций каркаса здания"],
)
def test_cp1251_names_under_english_headers(name: str) -> None:
    result = preview(f"Name,Duration\n{name},3\n".encode("cp1251"), "plan.csv")
    assert result.encoding == "cp1251"
    assert _nth(result, 0)["name"] == name


def test_cp1250_polish_names() -> None:
    text = "Nazwa zadania;Czas trwania\nZbrojenie płyty;3\nWykop pod ławy;2\n"
    result = preview(text.encode("cp1250"), "plan.csv")
    assert result.encoding == "cp1250"
    assert [a["name"] for a in _ordered(result)] == ["Zbrojenie płyty", "Wykop pod ławy"]


def test_cp1254_turkish_headers() -> None:
    text = "Görev Adı;Başlangıç;Süre\nKazı işleri;2026-05-04;3\n"
    result = preview(text.encode("cp1254"), "plan.csv")
    assert result.encoding == "cp1254"
    assert result.mapping == {"name": 0, "start": 1, "duration": 2}
    assert _nth(result, 0)["name"] == "Kazı işleri"


@pytest.mark.parametrize(
    "text",
    [
        "Vorgangsname;Dauer;Notizen\nBeton 25 m³ einbauen;3;1.200 €\nÄußere Wände;2;Größe prüfen\n",
        "Nom de la tâche;Durée\nTerrassement général;3\nPose des fenêtres;2\n",
        "Nombre de tarea;Duración\nCimentación;3\nAlbañilería y señalización;2\n",
    ],
)
def test_western_files_stay_cp1252(text: str) -> None:
    result = preview(text.encode("cp1252"), "plan.csv")
    assert result.encoding == "cp1252"
    assert _nth(result, 0)["name"] == text.split("\n")[1].split(";")[0]


def test_utf8_cyrillic_is_not_second_guessed() -> None:
    result = preview("Наименование;Длительность\nКладка;3\n".encode(), "plan.csv")
    assert result.encoding == "utf-8-sig"  # the first probe, which also reads UTF-8 without a mark
    assert _nth(result, 0)["name"] == "Кладка"


def test_sep_directive_and_tabs() -> None:
    result = preview(b"sep=|\nName|Duration\nA|2\n", "plan.csv")
    assert result.delimiter == "|"
    assert result.header_row == 2
    assert _nth(result, 0)["duration_days"] == 2
    tabbed = preview(b"Name\tDuration\nA\t2\n", "plan.txt")
    assert tabbed.delimiter == "\t"


def test_utf16_csv() -> None:
    data = "Name,Duration\nA,2\n".encode("utf-16")
    result = preview(data, "plan.csv")
    assert _nth(result, 0)["duration_days"] == 2


def test_letterhead_rows_above_the_header() -> None:
    rows = [["Acme Builders Ltd"], ["Master schedule"], [], ["ID", "Name", "Start"], ["1", "Survey", "2026-05-04"]]
    result = preview(_csv(rows), "plan.csv")
    assert result.header_row == 4
    assert _nth(result, 0)["metadata"]["import_row"] == 5


# ── xlsx ─────────────────────────────────────────────────────────────────────


def test_xlsx_picks_the_sheet_with_the_schedule(tmp_path: Path) -> None:
    data = _xlsx(
        tmp_path,
        [["ID", "Name", "Duration"], [10, "Pour", 2]],
        title="Programme",
        sheets=[("Cover", [["Project 12 Oak Lane"]], "visible"), ("Secret", [["ID", "Name"]], "hidden")],
        others_first=True,
    )
    result = preview(data, "plan.xlsx")
    assert result.sheet == "Programme"
    assert "sheet_selected" in _codes(result, "info")
    assert _acts(result)["10"]["activity_code"] == "10"


def test_xlsx_round_number_ids_read_as_integers(tmp_path: Path) -> None:
    data = _xlsx(tmp_path, [["ID", "Name", "Predecessors"], [1.0, "A", None], [2.0, "B", 1]])
    result = preview(data, "plan.xlsx")
    assert _links(result) == {("1", "2", "FS", 0)}


# ── fixtures and round trip ──────────────────────────────────────────────────


def test_residential_us_export() -> None:
    result = preview((FIXTURES / "residential_us.csv").read_bytes(), "residential_us.csv")
    assert not result.has_errors, result.issues
    assert result.date_order == "mdy"
    assert result.outline_source == "outline_level"
    acts = _acts(result)
    assert acts["4"]["start_date"] == "2026-05-04"
    assert acts["7"]["resources"][1]["name"] == "Pump truck"
    assert acts["3"]["activity_type"] == "milestone"
    assert acts["4"]["status"] == "completed"
    assert acts["4"]["parent_ref"] == "2"
    assert ("5", "7", "FS", 1) in _links(result)
    assert ("8", "9", "SS", 2) in _links(result)
    assert "duration_mismatch" not in _codes(result)
    _assert_interchange_clean(result)


def test_roundtrip_fixture_uses_the_export_header_row() -> None:
    source = ROUTER.read_text(encoding="utf-8")
    body = source[source.index("def _render_schedule_csv") :]
    header_block = body[body.index("writer.writerow(") : body.index("]", body.index("writer.writerow("))]
    exported = re.findall(r'"([^"]+)"', header_block)
    with (FIXTURES / "roundtrip_export.csv").open(encoding="utf-8", newline="") as handle:
        assert next(csv.reader(handle)) == exported


def test_roundtrip_with_the_schedule_csv_export() -> None:
    result = preview((FIXTURES / "roundtrip_export.csv").read_bytes(), "roundtrip_export.csv")
    assert not result.has_errors, result.issues
    assert "duration_mismatch" not in _codes(result)
    with (FIXTURES / "roundtrip_export.csv").open(encoding="utf-8", newline="") as handle:
        exported = list(csv.DictReader(handle))
    acts = _acts(result)
    for row in exported:
        act = acts[row["Activity Code"]]
        assert act["activity_code"] == row["Activity Code"]
        assert act["wbs_code"] == row["WBS"]
        assert (act["start_date"], act["end_date"]) == (row["Start"], row["End"])
        assert act["duration_days"] == int(row["Duration (days)"])
        assert float(act["progress_pct"]) == float(row["Progress (%)"])
    assert acts["A40"]["name"] == "=Pour slab (night)"
    assert acts["A50"]["activity_type"] == "milestone"
    assert _links(result) == {
        ("A10", "A20", "FS", 0),
        ("A20", "A30", "SS", 2),
        ("A30", "A40", "FS", 0),
        ("A20", "A40", "FF", 3),
        ("A40", "A50", "FS", 0),
    }
    _assert_interchange_clean(result)


def test_preview_is_stable_and_fingerprinted() -> None:
    data = (FIXTURES / "roundtrip_export.csv").read_bytes()
    first, second = preview(data, "a.csv"), preview(data, "a.csv")
    assert first.to_dict() == second.to_dict()
    assert first.document["schedule"]["metadata"]["import"]["sha256"] == first.sha256


@pytest.mark.parametrize("lang", LANGUAGES)
def test_template_row_in_every_language_parses(lang: str) -> None:
    values = {
        "id": "A1",
        "name": "Task",
        "wbs": "1",
        "outline_level": "1",
        "start": "2026-05-04",
        "finish": "2026-05-08",
        "duration": "5",
        "predecessors": "",
        "percent_complete": "10",
        "milestone": "",
        "resource": "Crew",
        "notes": "n",
        "client_visible": "",
    }
    rows = [list(TEMPLATE_HEADERS[lang]), [values[f] for f in FIELDS], ["A2", "Next", "2", "1", "", "", "1", "A1"]]
    result = preview(_csv(rows, delimiter=";"), f"template_{lang}.csv")
    assert not result.has_errors, (lang, result.issues)
    assert set(result.mapping) == set(FIELDS)
    assert _links(result) == {("A1", "A2", "FS", 0)}
    _assert_interchange_clean(result)


def test_whole_document_is_json_safe() -> None:
    import json

    result = preview((FIXTURES / "residential_us.csv").read_bytes(), "residential_us.csv")
    json.dumps(result.to_dict())


def test_zip_that_is_not_a_workbook() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "hello")
    assert _codes(preview(buffer.getvalue(), "plan.xlsx")) == ["file_unreadable"]


def test_values_wider_than_their_columns_are_refused_in_the_preview() -> None:
    header = ["ID", "Name", "WBS", "Resource"]
    rows = [["A", "Pour", "1." * 25 + "1", ""], ["B", "Cure", "2", "R" * 256], ["C", "Strip", "3", "R" * 255]]
    result = _simple(header, rows)
    found = {(i.params["field"], i.row, i.column) for i in result.issues if i.code == "value_too_long"}
    assert found == {("wbs", 2, 2), ("resource", 3, 3)}


def test_schedule_name_from_a_long_filename_is_cut_to_fit() -> None:
    result = preview(b"Name\nA\n", "x" * 300 + ".csv")
    assert len(result.document["schedule"]["name"]) == 255


@pytest.mark.parametrize(
    "text",
    [
        "sep=;\nName;Start\nA;1\n",
        "Name;Rate\nA;125,5\nB;3,25\n",
        "Name\tStart\nA\t1\n",
        "Acme Ltd\nMaster schedule\nName,Start,Finish\nA,1,2\nB,3,4\n",
        'Name,Notes\nA,"x; y"\n',
    ],
)
def test_promoted_delimiter_sniffer_matches_the_boq_original(text: str) -> None:
    from app.core.csv_dialect import sniff_delimiter
    from app.modules.boq.importers.excel import _sniff_delimiter

    assert sniff_delimiter(text) == _sniff_delimiter(text)
