# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Pure tests for the NRM 1 elemental rollup.

The engine regroups a bill; it never invents or loses money. Every test that
builds a plan therefore checks the one invariant a cost plan has to keep:
facilitating works estimate (group 0) + building works estimate (groups 1-8)
+ groups 9-14 priced in the bill + not allocated == direct cost, exactly, as
Decimals. The fixtures are chosen so that a plausible wrong
implementation fails: a leaf with its own bad code under a well-coded section
(inheriting would hide the typo), ``2.10`` next to ``2.1`` (a float parse merges
them), an element number the table does not list (dropping it loses money).
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.modules.cost_plan.engine import (
    MAX_LISTED_UNALLOCATED,
    LeafInput,
    MarkupInput,
    build_cost_plan,
    load_nrm1_catalogue,
    normalise_code,
    place_code,
)
from app.modules.cost_plan.export import build_cost_plan_workbook
from app.modules.cost_plan.schemas import CostPlanResponse

D = Decimal


def _leaf(amount: str, code: object = None, inherited: object = None, *, placeholder: bool = False) -> LeafInput:
    return LeafInput(
        id=uuid.uuid4(),
        ordinal="x",
        description="item",
        amount=D(amount),
        code=code,
        inherited_code=inherited,
        placeholder=placeholder,
    )


def _plan(leaves: list[LeafInput], markups: list[MarkupInput] | None = None, gifa: str | None = None):
    return build_cost_plan(
        boq_id=uuid.uuid4(),
        boq_name="Bill",
        project_id=uuid.uuid4(),
        currency="GBP",
        leaves=leaves,
        markups=markups or [],
        gifa=D(gifa) if gifa is not None else None,
        gifa_source="entered" if gifa is not None else "none",
    )


def _group(plan: CostPlanResponse, code: str):
    for group in [*plan.groups, *plan.addon_groups]:
        if group.code == code:
            return group
    raise AssertionError(f"group {code} missing")


def _element(plan: CostPlanResponse, code: str):
    for group in [*plan.groups, *plan.addon_groups]:
        for element in group.elements:
            if element.code == code:
                return element
    raise AssertionError(f"element {code} missing")


def _assert_conserved(plan: CostPlanResponse, leaves: list[LeafInput]) -> None:
    """The cost plan invariant, exact."""
    expected = sum((leaf.amount for leaf in leaves), D("0"))
    addons = sum((g.total for g in plan.addon_groups), D("0"))
    assert plan.direct_cost.total == expected
    facilitating = plan.facilitating_works_estimate.total
    building = plan.building_works_estimate.total
    assert facilitating + building + addons + plan.unallocated.total == expected
    assert facilitating == sum((g.total for g in plan.groups if g.code == "0"), D("0"))
    assert building == sum((g.total for g in plan.groups if g.code != "0"), D("0"))
    for group in [*plan.groups, *plan.addon_groups]:
        level = group.group_level.total if group.group_level else D("0")
        assert group.total == sum((e.total for e in group.elements), D("0")) + level
    assert plan.grand_total.total == plan.direct_cost.total + plan.markups_total.total


# ── The element table ────────────────────────────────────────────────────────


def test_table_carries_groups_0_to_14_in_order_with_works_then_addons() -> None:
    table = load_nrm1_catalogue()
    assert [g.code for g in table.groups] == [str(n) for n in range(15)]
    assert [g.kind for g in table.groups] == ["works"] * 9 + ["addon"] * 6
    superstructure = table.group("2")
    assert superstructure is not None
    assert [e.code for e in superstructure.elements] == [f"2.{n}" for n in range(1, 9)]
    assert superstructure.elements[0].name == "Frame"
    assert superstructure.elements[-1].name == "Internal doors"


def test_table_codes_are_unique_and_belong_to_their_group() -> None:
    table = load_nrm1_catalogue()
    seen: set[str] = set()
    for group in table.groups:
        for element in group.elements:
            assert element.code not in seen
            seen.add(element.code)
            assert element.code.split(".")[0] == group.code


# ── Codes ────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("2.5", "2.5"),
        ("02.05", "2.5"),
        ("2.6.1.", "2.6.1"),
        (" 5.10 ", "5.10"),
        ("2.10", "2.10"),
        (7, "7"),
        (2.1, "2.1"),
        ("abc", None),
        ("2,5", None),
        ("", None),
        (None, None),
        (True, None),
    ],
)
def test_normalise_code(raw: object, expected: str | None) -> None:
    assert normalise_code(raw) == expected


@pytest.mark.parametrize(
    ("raw", "reason", "group", "element"),
    [
        ("2.5", "matched", "2", "2.5"),
        ("02.05", "matched", "2", "2.5"),
        ("2.6.1", "matched", "2", "2.6"),
        ("5.10", "matched", "5", "5.10"),
        ("2", "group_level", "2", None),
        (1, "group_level", "1", None),
        ("2.10", "unknown_element", "2", None),
        ("5.99", "unknown_element", "5", None),
        ("9", "group_level", "9", None),
        ("13.1", "unknown_element", "13", None),
        ("15.1", "unknown_group", None, None),
        ("abc", "invalid_code", None, None),
        ("   ", "no_code", None, None),
        (None, "no_code", None, None),
    ],
)
def test_place_code(raw: object, reason: str, group: str | None, element: str | None) -> None:
    placement = place_code(raw, load_nrm1_catalogue())
    assert placement.reason == reason
    assert placement.group_code == group
    assert placement.element_code == element


def test_ten_and_one_are_different_elements() -> None:
    """A float parse would put 5.10 (lifts) on 5.1 (sanitary)."""
    plan = _plan([_leaf("100", "5.1"), _leaf("7", "5.10")])
    assert _element(plan, "5.1").total == D("100")
    assert _element(plan, "5.10").total == D("7")


# ── Rollup ───────────────────────────────────────────────────────────────────


def test_every_kind_of_code_is_conserved_and_placed() -> None:
    leaves = [
        _leaf("1000.10", "1.1"),
        _leaf("250.05", "02.05"),
        _leaf("80", "2.6.1"),
        _leaf("40", None, inherited="2.7"),  # uncoded item under a coded section
        _leaf("30", "2"),
        _leaf("20", "5.99"),
        _leaf("15", "9.1"),  # preliminaries priced in the bill
        _leaf("12.34", "15.1"),
        _leaf("9", "abc", inherited="2.5"),  # own bad code wins over the section
        _leaf("3.33", None),
    ]
    plan = _plan(leaves)
    _assert_conserved(plan, leaves)

    assert _element(plan, "1.1").total == D("1000.10")
    assert _element(plan, "2.5").total == D("250.05")
    assert _element(plan, "2.6").total == D("80")
    assert _element(plan, "2.7").total == D("40")
    superstructure = _group(plan, "2")
    assert superstructure.group_level is not None
    assert superstructure.group_level.total == D("30")
    assert superstructure.total == D("400.05")
    services = _group(plan, "5")
    assert services.group_level is not None
    assert services.group_level.codes == ["5.99"]
    assert services.total == D("20")
    assert _group(plan, "9").total == D("15")
    assert plan.building_works_estimate.total == D("1420.15")
    assert plan.facilitating_works_estimate.total == D("0")

    assert plan.unallocated.total == D("24.67")
    assert plan.unallocated.position_count == 3
    reasons = sorted(p.reason for p in plan.unallocated.positions)
    assert reasons == ["invalid_code", "no_code", "unknown_group"]
    bad = next(p for p in plan.unallocated.positions if p.reason == "invalid_code")
    assert bad.code == "abc"
    assert plan.inherited_count == 1
    assert plan.allocated_count == 7
    assert plan.position_count == 10


def test_an_empty_bill_is_all_zeros_without_dividing_by_zero() -> None:
    plan = _plan([], gifa="100")
    assert plan.direct_cost.total == 0
    assert plan.grand_total.total == 0
    assert plan.grand_total.share_pct is None
    assert plan.grand_total.cost_per_m2 == D("0.00")


# ── Cost per m2 and share of total ──────────────────────────────────────────


def test_cost_per_m2_with_gifa() -> None:
    leaves = [_leaf("1000", "1.1"), _leaf("500", "2.1")]
    plan = _plan(leaves, gifa="250")
    assert plan.gifa == D("250")
    assert plan.gifa_source == "entered"
    assert _element(plan, "1.1").cost_per_m2 == D("4.00")
    assert _element(plan, "2.1").cost_per_m2 == D("2.00")
    assert plan.building_works_estimate.cost_per_m2 == D("6.00")
    assert plan.grand_total.cost_per_m2 == D("6.00")
    assert "no_gifa" not in plan.warnings


def test_cost_per_m2_without_gifa_is_null_everywhere_not_zero() -> None:
    plan = _plan([_leaf("1000", "1.1")])
    assert plan.gifa is None
    assert plan.gifa_source == "none"
    assert _element(plan, "1.1").cost_per_m2 is None
    assert _group(plan, "1").cost_per_m2 is None
    assert plan.grand_total.cost_per_m2 is None
    assert "no_gifa" in plan.warnings


def test_zero_gifa_is_treated_as_absent() -> None:
    plan = _plan([_leaf("1000", "1.1")], gifa="0")
    assert plan.gifa is None
    assert plan.gifa_source == "none"
    assert plan.grand_total.cost_per_m2 is None


def test_share_is_of_the_cost_plan_total_including_markups() -> None:
    markups = [
        MarkupInput(
            id=None,
            name="OH&P",
            category="overhead",
            markup_type="fixed",
            apply_to="direct_cost",
            amount=D("1000"),
            fixed_amount=D("1000"),
        )
    ]
    plan = _plan([_leaf("3000", "1.1")], markups)
    assert plan.grand_total.total == D("4000")
    assert _element(plan, "1.1").share_pct == D("75.00")
    assert plan.markups[0].share_pct == D("25.00")
    assert plan.grand_total.share_pct == D("100.00")


# ── Markup cascade (read, not recomputed) ───────────────────────────────────


def _pct_line(name: str, pct: str, apply_to: str, amount: str) -> MarkupInput:
    return MarkupInput(
        id=uuid.uuid4(),
        name=name,
        category="overhead",
        markup_type="percentage",
        apply_to=apply_to,
        amount=D(amount),
        percentage=D(pct),
    )


def test_cascade_keeps_order_amounts_and_shows_each_lines_real_base() -> None:
    leaves = [_leaf("1000", "1.1")]
    markups = [
        _pct_line("Prelims", "10", "direct_cost", "100"),
        _pct_line("OH&P", "10", "cumulative", "110"),
        MarkupInput(
            id=uuid.uuid4(),
            name="Fees",
            category="other",
            markup_type="fixed",
            apply_to="direct_cost",
            amount=D("500"),
            fixed_amount=D("500"),
        ),
        _pct_line("VAT", "20", "subtotal", "342"),
    ]
    plan = _plan(leaves, markups)
    assert [m.name for m in plan.markups] == ["Prelims", "OH&P", "Fees", "VAT"]
    # The amounts are the bill's, passed through untouched.
    assert [m.total for m in plan.markups] == [D("100"), D("110"), D("500"), D("342")]
    # Base: direct cost for a direct line, running subtotal for a compounding one.
    assert plan.markups[0].base == D("1000")
    assert plan.markups[1].base == D("1100")
    assert plan.markups[2].base is None
    assert plan.markups[3].base == D("1710")
    assert [m.running_total for m in plan.markups] == [D("1100"), D("1210"), D("1710"), D("2052")]
    assert plan.grand_total.total == D("2052")
    assert plan.markups_total.total == D("1052")
    _assert_conserved(plan, leaves)


def test_scoped_cascade_does_not_claim_a_single_base() -> None:
    line = MarkupInput(
        id=uuid.uuid4(),
        name="Section OH",
        category="overhead",
        markup_type="percentage",
        apply_to="direct_cost",
        amount=D("12"),
        percentage=D("12"),
        scoped=True,
    )
    plan = _plan([_leaf("1000", "1.1")], [_pct_line("Prelims", "10", "direct_cost", "100"), line])
    assert all(m.base is None for m in plan.markups)
    assert "scoped_markups" in plan.warnings


def test_addon_groups_in_the_bill_next_to_a_markup_stack_are_flagged() -> None:
    plan = _plan([_leaf("100", "9")], [_pct_line("Prelims", "10", "direct_cost", "10")])
    assert "addons_in_bill_and_markups" in plan.warnings
    plain = _plan([_leaf("100", "9")])
    assert "addons_in_bill_and_markups" not in plain.warnings


# ── Unallocated listing ─────────────────────────────────────────────────────


def test_unallocated_list_is_capped_but_total_and_count_are_not() -> None:
    leaves = [_leaf("1", None) for _ in range(MAX_LISTED_UNALLOCATED + 5)]
    plan = _plan(leaves)
    assert plan.unallocated.position_count == MAX_LISTED_UNALLOCATED + 5
    assert plan.unallocated.total == D(MAX_LISTED_UNALLOCATED + 5)
    assert len(plan.unallocated.positions) == MAX_LISTED_UNALLOCATED
    assert plan.unallocated.positions_truncated is True
    _assert_conserved(plan, leaves)


# ── Empty placeholder rows ──────────────────────────────────────────────────


def test_a_blank_row_in_a_fully_coded_bill_raises_no_warning_and_counts_nowhere() -> None:
    """The estimator just clicked "Add Position": the plan must look the same."""
    coded = [_leaf("1000", "1.1"), _leaf("500", "2.5")]
    leaves = [*coded, _leaf("0", None, placeholder=True)]
    plan = _plan(leaves)
    _assert_conserved(plan, leaves)

    assert "unallocated_positions" not in plan.warnings
    assert plan.unallocated.position_count == 0
    assert plan.unallocated.positions == []
    assert plan.unallocated.positions_truncated is False
    assert plan.position_count == 2
    assert plan.allocated_count == plan.position_count
    without = _plan(coded)
    assert plan.warnings == without.warnings
    assert plan.direct_cost == without.direct_cost
    assert [g.position_count for g in plan.groups] == [g.position_count for g in without.groups]


def test_a_blank_row_under_a_coded_section_does_not_bump_the_element_count() -> None:
    leaves = [_leaf("40", None, inherited="2.7"), _leaf("0", None, inherited="2.7", placeholder=True)]
    plan = _plan(leaves)
    assert _element(plan, "2.7").position_count == 1
    assert _group(plan, "2").position_count == 1
    assert plan.position_count == 1
    assert plan.inherited_count == 1


def test_an_unallocated_real_row_is_still_listed_next_to_a_blank_one() -> None:
    leaves = [_leaf("0", None, placeholder=True), _leaf("5", None)]
    plan = _plan(leaves)
    assert plan.unallocated.position_count == 1
    assert len(plan.unallocated.positions) == 1
    assert plan.unallocated.positions_truncated is False
    assert "unallocated_positions" in plan.warnings


def test_a_row_flagged_as_placeholder_that_carries_money_is_still_counted() -> None:
    """The flag never outranks the money: a wrong flag must not lose an amount."""
    leaves = [_leaf("12.50", None, placeholder=True)]
    plan = _plan(leaves)
    _assert_conserved(plan, leaves)
    assert plan.unallocated.total == D("12.50")
    assert plan.unallocated.position_count == 1
    assert plan.position_count == 1


# ── Facilitating works vs building works ────────────────────────────────────


def test_facilitating_works_stay_out_of_the_building_works_estimate() -> None:
    """NRM 1: the building works estimate is groups 1-8 only.

    The wrong answer is 4,400,000, the sum of groups 0-8, and the cost per m2
    a benchmark is read against would be overstated by the demolition.
    """
    leaves = [
        _leaf("400000", "0.2"),  # major demolition works
        _leaf("1500000", "1.1"),
        _leaf("2000000", "2.1"),
        _leaf("500000", "8.1"),
        _leaf("100000", "9"),
    ]
    plan = _plan(leaves, gifa="2000")
    _assert_conserved(plan, leaves)

    assert plan.facilitating_works_estimate.total == D("400000")
    assert plan.building_works_estimate.total == D("4000000")
    assert plan.building_works_estimate.total != D("4400000")
    assert plan.building_works_estimate.cost_per_m2 == D("2000.00")
    assert plan.facilitating_works_estimate.cost_per_m2 == D("200.00")
    # Shares are of the cost plan total, 4,500,000 here.
    assert plan.facilitating_works_estimate.share_pct == D("8.89")
    assert plan.building_works_estimate.share_pct == D("88.89")
    assert _group(plan, "0").total == D("400000")


def test_the_workbook_prints_both_estimates_with_the_plans_figures() -> None:
    leaves = [_leaf("400000", "0.2"), _leaf("4000000", "2.1"), _leaf("7", None)]
    plan = _plan(leaves, gifa="2000")
    book = load_workbook(io.BytesIO(build_cost_plan_workbook(plan)))
    rows = list(book.worksheets[0].iter_rows(values_only=True))

    def row(text: str) -> tuple:
        matches = [r for r in rows if len(r) > 1 and r[1] == text]
        assert len(matches) == 1, text
        return matches[0]

    facilitating = row("Facilitating works estimate")
    building = row("Building works estimate")
    assert facilitating[4] == pytest.approx(400000.0)
    assert building[4] == pytest.approx(4000000.0)
    assert building[5] == pytest.approx(2000.0)
    assert building[0] == "1-8"

    # Group 0, its subtotal, groups 1-8, their subtotal: in that order.
    def at(predicate) -> int:
        return next(i for i, r in enumerate(rows) if predicate(r))

    group_0 = at(lambda r: r[0] == "0")
    group_1 = at(lambda r: r[0] == "1")
    group_8 = at(lambda r: r[0] == "8")
    assert group_0 < rows.index(facilitating) < group_1 < group_8 < rows.index(building)


def test_json_money_is_positional_text() -> None:
    plan = _plan([_leaf("0.00000001", "1.1"), _leaf("1E+3", "1.1")])
    payload = plan.model_dump(mode="json")
    assert payload["direct_cost"]["total"] == "1000.00000001"
    assert "E" not in payload["groups"][1]["total"]
