# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""A Hungarian workbook imports whole, with its numbers, and says what it left out.

A user in Hungary installed the Hungarian country pack, imported their bills
and "did not get a good result". The flat single-sheet bill already imported
correctly; what still went wrong, each checked against the fixtures in
:mod:`tests.fixtures.hu_boq`:

* a workbook that opens on its summary sheet ("Főösszesítő", the trade names
  beside "Anyag összesen | Díj összesen") had that sheet taken for the bill:
  the import was three empty headings, and every line on the trade sheets
  behind it was never read;
* a bill priced by trade, one sheet per trade, imported the first trade only;
* headings the other Hungarian exports write ("Tételszöveg", "Egységár
  anyag", "Díj egységre", "Anyag összege") were not recognised, so the bill
  imported nothing or every line at a rate of zero, with no error;
* "12.500", which a Hungarian bill means as twelve thousand five hundred,
  was read as twelve and a half;
* the project page's Import button posts to the smart route, which read the
  spreadsheet on its own and bypassed the Hungarian workbook profile.

Every one of those failed quietly. The tests below hold the numbers and the
report: what was read, what was not, and why.
"""

from __future__ import annotations

import asyncio
import io

import pytest
from openpyxl import Workbook

from app.modules.boq.importers._base import ImportedBOQ, ImporterParseError
from app.modules.boq.importers._encoding import dot_groups_thousands, parse_numeric_cell
from app.modules.boq.importers.excel import ExcelImporter
from app.modules.boq.importers.hungary_workbook import _number
from tests.fixtures import hu_boq


def _parse(content: bytes) -> ImportedBOQ:
    return asyncio.run(ExcelImporter.parse(content))


def _lines(result: ImportedBOQ) -> list:
    return [position for position in result.positions if not position.is_section]


def _coded(result: ImportedBOQ, code: str) -> list[dict]:
    return [issue for issue in (*result.warnings, *result.errors) if issue.get("code") == code]


# ── A workbook with a sheet per trade ──────────────────────────────────────


def test_every_trade_sheet_is_read_and_the_summary_sheet_is_not() -> None:
    result = _parse(hu_boq.trade_workbook_xlsx())

    assert result.errors == []
    got = [(p.metadata.get("import_sheet"), p.ordinal, p.unit, p.quantity, p.unit_rate) for p in _lines(result)]
    want = [
        (sheet, ordinal, unit, qty, pytest.approx(rate))
        for sheet, ordinal, unit, qty, rate in hu_boq.EXPECTED_TRADE_LINES
    ]
    assert got == want
    assert result.metadata["item_sheets"] == ["Építészet", "Épületgépészet", "Villamos"]


def test_the_bill_total_is_the_sum_of_every_trade() -> None:
    result = _parse(hu_boq.trade_workbook_xlsx())

    total = sum(p.quantity * p.unit_rate for p in _lines(result))
    architectural = sum(qty * (material + fee) for _, _, qty, material, fee in hu_boq.EXPECTED_LINES)
    assert total == pytest.approx(architectural + 36 * 3900 + 120 * 1050)


def test_each_trade_opens_with_a_section_named_after_its_sheet() -> None:
    result = _parse(hu_boq.trade_workbook_xlsx())

    sections = [(p.ordinal, p.description) for p in result.positions if p.is_section]
    assert ("1", "Építészet") in sections
    assert ("2", "Épületgépészet") in sections
    assert ("3", "Villamos") in sections
    ordinals = [p.ordinal for p in result.positions]
    assert len(ordinals) == len(set(ordinals)), "line numbers repeat across trades and must not collide"


def test_the_sheets_that_were_not_read_are_named_with_the_reason() -> None:
    result = _parse(hu_boq.trade_workbook_xlsx(hidden=True, copy=True))

    notes = {issue["sheet"]: issue["reason"] for issue in _coded(result, "sheet_not_read")}
    assert notes == {
        "Záradék": "no_item_header",
        "Főösszesítő": "no_item_header",
        "Építészet (2)": "duplicate",
        "Segéd": "hidden",
    }
    assert "Munkaanyag" not in [p.description for p in result.positions]
    assert len(_lines(result)) == len(hu_boq.EXPECTED_TRADE_LINES)


def test_a_total_line_on_a_trade_sheet_names_that_sheet() -> None:
    result = _parse(hu_boq.trade_workbook_xlsx())

    summary = {(w.get("sheet"), w["label"]) for w in _coded(result, "summary_row_skipped")}
    assert ("Épületgépészet", "Épületgépészeti csővezeték szerelése összesen:") in summary
    assert ("Építészet", "Mindösszesen") in summary


def test_a_single_sheet_bill_keeps_its_own_numbers_and_gets_no_sheet_section() -> None:
    result = _parse(hu_boq.flat_xlsx())

    assert [p.ordinal for p in _lines(result)] == [line[0] for line in hu_boq.EXPECTED_LINES]
    assert "Költségvetés" not in [p.description for p in result.positions if p.is_section]
    assert _coded(result, "sheet_not_read") == []


def test_trades_that_number_their_lines_apart_keep_their_numbers() -> None:
    """Prefixing is for collisions only: a workbook numbered 1-8, 9, 10 keeps those numbers."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    first = workbook.create_sheet("Építészet")
    first.append(hu_boq.HEADER)
    first.append(["1", "", "Földkiemelés", "10", "m3", "0", "1 850", "", ""])
    second = workbook.create_sheet("Villamos")
    second.append(hu_boq.HEADER)
    second.append(["2", "", "Védőcső", "120", "m", "450", "600", "", ""])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = _parse(buffer.getvalue())

    assert [p.ordinal for p in _lines(result)] == ["1", "2"]


# ── Header variants ────────────────────────────────────────────────────────


@pytest.mark.parametrize("variant", sorted(hu_boq.HEADER_VARIANTS))
def test_every_hungarian_way_of_heading_the_bill_reads_the_same_lines(variant: str) -> None:
    result = _parse(hu_boq.header_variant_xlsx(variant))

    assert result.errors == []
    got = [(p.ordinal, p.unit, p.quantity, p.unit_rate) for p in _lines(result)]
    want = [(ssz, unit, qty, pytest.approx(material + fee)) for ssz, unit, qty, material, fee in hu_boq.EXPECTED_LINES]
    assert got == want
    assert result.metadata["header_report"]["unrecognised"] == []


def test_a_header_with_no_known_columns_is_an_error_naming_the_headings() -> None:
    result = _parse(hu_boq.unrecognised_header_xlsx())

    assert result.positions == []
    [error] = _coded(result, "header_not_recognised")
    # The headings reported are the table's own, not the title above it.
    assert error["unrecognised"] == ["Sor", "Kód", "Munka", "Darab", "Ár"]
    assert error["recognised"] == {"ME": "unit"}
    assert error["missing"] == ["description"]


def test_a_header_naming_nothing_to_price_by_is_an_error_not_an_empty_bill() -> None:
    result = _parse(hu_boq.description_only_xlsx())

    assert _lines(result) == []
    [error] = _coded(result, "header_not_recognised")
    assert error["missing"] == ["quantity_or_rate"]
    assert "Darab" in error["unrecognised"]
    assert error["recognised"]["Tétel szövege"] == "description"


def test_the_report_says_which_heading_fed_which_column() -> None:
    result = _parse(hu_boq.flat_xlsx())

    recognised = result.metadata["header_report"]["recognised"]
    assert recognised["Menny."] == "quantity"
    assert recognised["Anyag egységár"] == "unit_rate"
    assert recognised["Tételszám"] == "classification"


# ── Numbers ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("builder", [hu_boq.dot_thousands_csv, hu_boq.dot_thousands_xlsx])
def test_dots_between_thousands_are_thousands_in_a_hungarian_bill(builder) -> None:
    result = _parse(builder())

    assert result.errors == []
    got = [(p.ordinal, p.quantity, p.unit_rate) for p in _lines(result)]
    assert got == [(ordinal, qty, pytest.approx(rate)) for ordinal, qty, rate in hu_boq.EXPECTED_DOT_THOUSANDS]


def test_every_cell_read_with_dot_thousands_is_reported_once() -> None:
    result = _parse(hu_boq.dot_thousands_csv())

    notes = sorted((w["row"], w["column"], w["text"], w["value"]) for w in _coded(result, "dot_read_as_thousands"))
    assert notes == [
        (2, "labour_rate", "1.200", 1200.0),
        (2, "quantity", "12.500", 12500.0),
        (3, "labour_rate", "12.500 Ft", 12500.0),
        (4, "material_rate", "3.200", 3200.0),
    ]


def test_an_english_bill_keeps_reading_a_lone_dot_as_the_decimal_point() -> None:
    content = b"Description;Quantity;Unit;Rate\r\nConcrete;12.500;m3;1.250\r\n"
    result = _parse(content)

    [line] = _lines(result)
    assert (line.quantity, line.unit_rate) == (12.5, 1.25)
    assert _coded(result, "dot_read_as_thousands") == []


@pytest.mark.parametrize(
    ("text", "value"),
    [
        ("12.500", 12500.0),
        ("1.250.000", 1250000.0),
        ("12.500 Ft", 12500.0),
        ("-1.500", -1500.0),
        ("1.234,5", 1234.5),
        ("1,234.5", 1234.5),
        ("12.5", 12.5),
        ("1.2345", 1.2345),
        ("1,234", 1.234),
        ("12 500", 12500.0),
    ],
)
def test_mixed_decimal_conventions_read_by_their_last_separator(text: str, value: float) -> None:
    """Only dots between groups of three with no comma are the ambiguous case the flag decides."""
    assert parse_numeric_cell(text, dot_thousands=True) == (value, None)


def test_a_cell_that_is_already_a_number_is_never_reinterpreted() -> None:
    assert not dot_groups_thousands(12.5)
    assert parse_numeric_cell(12.5, dot_thousands=True) == (12.5, None)


@pytest.mark.parametrize(
    ("text", "value"),
    [("12.500", 12500.0), ("12 500 Ft", 12500.0), ("12.500 Ft", 12500.0), ("1 234,5", 1234.5), ("420,50", 420.5)],
)
def test_the_workbook_profile_reads_hungarian_numbers_the_same_way(text: str, value: float) -> None:
    assert _number(text) == value


def test_a_cell_that_is_not_a_number_is_an_error_on_its_row() -> None:
    rows = [hu_boq.HEADER, ["1", "", "Földkiemelés", "sok", "m3", "0", "1 850", "", ""]]
    result = _parse("".join(";".join(row) + "\r\n" for row in rows).encode())

    assert _lines(result) == []
    [error] = result.errors
    assert error["row"] == 2
    assert "sok" in error["error"]


# ── Files that cannot be read ──────────────────────────────────────────────


def test_a_truncated_workbook_is_refused_with_a_reason() -> None:
    with pytest.raises(ImporterParseError):
        _parse(hu_boq.truncated_xlsx())


def test_an_empty_upload_is_refused() -> None:
    with pytest.raises(ImporterParseError):
        _parse(b"")


# ── The smart route reads a spreadsheet the way /import/auto/ does ─────────


def test_the_smart_route_reads_a_trade_workbook_through_the_same_importer() -> None:
    from app.modules.boq.router import _native_spreadsheet_import

    native = asyncio.run(_native_spreadsheet_import(hu_boq.trade_workbook_xlsx()))

    assert native is not None
    assert len(_lines(native)) == len(hu_boq.EXPECTED_TRADE_LINES)


def test_the_smart_route_keeps_the_building_workbook_profile() -> None:
    """The chapter workbook's item code is composed down its heading tree; the plain reader loses it."""
    from app.modules.boq.router import _native_spreadsheet_import

    workbook = Workbook()
    workbook.active.title = "Címlap"
    chapter = workbook.create_sheet("MA-01_ÁLT")
    chapter["B5"] = "EGYSÉGES MAGASÉPÍTÉSI ÁGAZATI TÉTELREND"
    for column, text in {
        "T": "Tétel szövege",
        "U": "Mennyiség",
        "V": "Egység",
        "X": "Anyag egységár",
        "Y": "Díj egységár",
    }.items():
        chapter[f"{column}5"] = text
    chapter["B6"], chapter["D6"], chapter["F6"], chapter["H6"] = "MA", "01", "001", "Felvonulás"
    chapter["J7"], chapter["T7"], chapter["U7"], chapter["V7"] = "001", "Felvonulási épület", 1, "klt"
    chapter["X7"], chapter["Y7"] = 100000, 50000
    buffer = io.BytesIO()
    workbook.save(buffer)

    native = asyncio.run(_native_spreadsheet_import(buffer.getvalue()))

    assert native is not None
    [line] = _lines(native)
    assert line.classification == {"tetelrend": "MA-01-001-001"}
    assert line.unit_rate == 150000.0


def test_the_smart_route_hands_back_bytes_that_are_no_workbook() -> None:
    from app.modules.boq.router import _native_spreadsheet_import

    assert asyncio.run(_native_spreadsheet_import(hu_boq.truncated_xlsx())) is None
