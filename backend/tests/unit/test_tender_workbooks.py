# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The tender workbooks read the comparison the way the screen does.

Pure renderer tests over a hand-built comparison answer. The round trip through
the stored package, the comparison endpoint and the export route is in
``tests/pg/test_tender_workbooks_pg.py``; these pin the readings the sheet adds
on top of the numbers: which price is lowest, which is an outlier, what a
missing price looks like, and which language the text is in.
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.modules.tendering.schemas import BidComparisonResponse, BidComparisonRow
from app.modules.tendering.workbooks import (
    LOWEST_FILL,
    MISSING_FILL,
    bidder_lines,
    classify_rate,
    render_bidder_xlsx,
    render_comparison_xlsx,
    unit_code,
)
from app.modules.tendering.xlsx_translations import _STRINGS, SUPPORTED_XLSX_LOCALES, resolve_xlsx_locale

A, B, C = "a" * 8, "b" * 8, "c" * 8


def _entry(bid_id: str, rate: str, qty: str = "10", *, priced: bool = True) -> dict:
    r = Decimal(rate)
    return {
        "bid_id": bid_id,
        "company_name": bid_id,
        # The comparison sends no figure for an unpriced line, as the service does.
        "unit_rate": float(r) if priced else None,
        "total": float(r * Decimal(qty)) if priced else None,
        "deviation_pct": 0.0,
        "priced": priced,
    }


def _comparison() -> BidComparisonResponse:
    return BidComparisonResponse(
        package_id=uuid.uuid4(),
        package_name="Shell and core",
        bid_count=3,
        bid_companies=["Alpha", "Beta", "Gamma"],
        budget_total=Decimal("3000"),
        rows=[
            BidComparisonRow(position_id="s", ordinal="01", description="Earthworks", unit="", bids=[]),
            BidComparisonRow(
                position_id="p1",
                ordinal="01.010",
                description="Excavation",
                unit="m3",
                budget_quantity=10,
                budget_rate=Decimal("100"),
                budget_total=Decimal("1000"),
                bids=[_entry(A, "90"), _entry(B, "100"), _entry(C, "50")],
            ),
            BidComparisonRow(
                position_id="p2",
                ordinal="01.020",
                description="Backfill",
                unit="m3",
                budget_quantity=10,
                budget_rate=Decimal("200"),
                budget_total=Decimal("2000"),
                bids=[_entry(A, "200"), _entry(B, "0", priced=False), _entry(C, "260")],
            ),
        ],
        bid_totals=[
            {
                "bid_id": A,
                "company_name": "Alpha",
                "total": 2900.0,
                "currency": "CHF",
                "deviation_pct": -3.3,
                "deviation_known": True,
                "status": "submitted",
            },
            {
                "bid_id": B,
                "company_name": "Beta",
                "total": 1000.0,
                "currency": "CHF",
                "deviation_pct": -66.7,
                "deviation_known": True,
                "status": "submitted",
            },
            # Quoted in another currency: shown, never ranked.
            {
                "bid_id": C,
                "company_name": "Gamma",
                "total": 3100.0,
                "currency": "USD",
                "deviation_pct": 0.0,
                "deviation_known": False,
                "status": "submitted",
            },
        ],
    )


def _sheet(data: bytes):
    return load_workbook(io.BytesIO(data)).active


def _row_of(ws, description: str) -> int:
    for row in ws.iter_rows():
        if len(row) > 1 and row[1].value == description:
            return row[1].row
    raise AssertionError(f"no row for {description!r}")


def _columns_of(ws, header: str) -> list[int]:
    """Columns whose top header cell names ``header`` (merged over price + total)."""
    return [c.column for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith(header)]


@pytest.mark.parametrize(
    ("rate", "rates", "expected"),
    [
        (Decimal("120"), [Decimal("100"), Decimal("100"), Decimal("120")], "high"),
        (Decimal("80"), [Decimal("100"), Decimal("100"), Decimal("80")], "low"),
        (Decimal("110"), [Decimal("100"), Decimal("110")], None),
        # One price is no median, a zero is not a price.
        (Decimal("100"), [Decimal("100"), Decimal("0")], None),
        (Decimal("0"), [Decimal("100"), Decimal("200")], None),
    ],
)
def test_outliers_follow_the_screens_rule(rate, rates, expected) -> None:
    assert classify_rate(rate, rates) == expected


def test_lowest_missing_and_foreign_currency_are_read_like_the_screen() -> None:
    ws = _sheet(render_comparison_xlsx(_comparison(), project_name="Depot", reporting_currency="CHF", locale="en"))
    alpha, beta, gamma = _columns_of(ws, "Alpha")[0], _columns_of(ws, "Beta")[0], _columns_of(ws, "Gamma")[0]

    excavation = _row_of(ws, "Excavation")
    # Gamma is cheapest on paper but quotes in USD, so Alpha is the lowest.
    assert ws.cell(excavation, alpha).fill.fgColor.rgb.endswith(LOWEST_FILL)
    assert not ws.cell(excavation, gamma).fill.fgColor.rgb.endswith(LOWEST_FILL)

    backfill = _row_of(ws, "Backfill")
    missing = ws.cell(backfill, beta)
    assert missing.value == "not priced", "a missing price is a word, not 0"
    assert missing.fill.fgColor.rgb.endswith(MISSING_FILL)
    assert ws.cell(backfill, beta + 1).value is None

    section = _row_of(ws, "Earthworks")
    assert all(ws.cell(section, col).value is None for col in (alpha, beta, gamma)), "a header asks no price"

    totals = _row_of(ws, "Bid total as submitted")
    assert ws.cell(totals, alpha + 1).value == 2900
    assert ws.cell(totals, beta + 1).value == 1000
    # Beta is cheapest only by the line it left out: an incomplete bid is not
    # ranked lowest, and Alpha is the only complete one in CHF.
    assert not ws.cell(totals, beta + 1).fill.fgColor.rgb.endswith(LOWEST_FILL)
    assert not ws.cell(totals, alpha + 1).fill.fgColor.rgb.endswith(LOWEST_FILL)
    lines = _row_of(ws, "Lines priced")
    assert ws.cell(lines, beta + 1).value == "1 of 2"
    assert ws.cell(lines, beta + 1).fill.fgColor.rgb.endswith(MISSING_FILL)
    assert ws.cell(lines, alpha + 1).value == "2 of 2"
    assert ws.cell(excavation, 3).value == "m³", "units print as the BOQ editor shows them"
    deviation = _row_of(ws, "Deviation from own estimate")
    assert ws.cell(deviation, gamma + 1).value == "n/a"


def test_an_outlier_keeps_its_number_and_shows_a_mark() -> None:
    comparison = _comparison()
    comparison.rows[1].bids = [_entry(A, "100"), _entry(B, "100"), _entry(C, "100")]
    comparison.rows[2].bids = [_entry(A, "100"), _entry(B, "100"), _entry(C, "100")]
    comparison.rows[1].bids[0] = _entry(A, "150")
    ws = _sheet(render_comparison_xlsx(comparison, project_name="", reporting_currency="", locale="en"))
    cell = ws.cell(_row_of(ws, "Excavation"), _columns_of(ws, "Alpha")[0])
    assert cell.value == 150
    assert "▲" in cell.number_format


def test_the_sheet_speaks_the_requested_language_and_falls_back_to_english() -> None:
    assert resolve_xlsx_locale("de-AT", None) == "de"
    assert resolve_xlsx_locale(None, "xx-YY, ru;q=0.8") == "ru"
    assert resolve_xlsx_locale("xx", "yy") == "en"
    ws = _sheet(render_comparison_xlsx(_comparison(), project_name="", reporting_currency="CHF", locale="de"))
    assert ws.cell(_row_of(ws, "Backfill"), _columns_of(ws, "Beta")[0]).value == "nicht angeboten"
    _row_of(ws, "Angebotssumme laut Angebot")


def test_every_language_carries_every_key() -> None:
    keys = set(_STRINGS["en"])
    for locale in SUPPORTED_XLSX_LOCALES:
        assert set(_STRINGS[locale]) == keys, locale


def test_the_bill_for_bidders_carries_no_contractor_price() -> None:
    section = SimpleNamespace(
        id=uuid.uuid4(),
        parent_id=None,
        ordinal="02",
        description="Concrete",
        unit="",
        quantity="0",
        unit_rate="0",
        total="98765.43",
    )
    other = SimpleNamespace(
        id=uuid.uuid4(),
        parent_id=None,
        ordinal="03",
        description="Roofing",
        unit="",
        quantity="0",
        unit_rate="0",
        total="0",
    )
    wall = SimpleNamespace(
        id=uuid.uuid4(),
        parent_id=section.id,
        ordinal="02.010",
        description="Wall",
        unit="m3",
        quantity="12.5",
        unit_rate="187.31",
        total="2341.38",
    )
    roof = SimpleNamespace(
        id=uuid.uuid4(),
        parent_id=other.id,
        ordinal="03.010",
        description="Roof",
        unit="m2",
        quantity="40",
        unit_rate="55.55",
        total="2222.00",
    )
    everything = [section, wall, other, roof]

    lines = bidder_lines(everything, [wall])
    assert [line.description for line in lines] == ["Concrete", "Wall"], "only the package's lines and their header"

    data = render_bidder_xlsx(
        lines, project_name="Depot", package_name="Lot 2", currency="CHF", deadline="2026-11-30", locale="en"
    )
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    wall_row = _row_of(ws, "Wall")
    assert ws.cell(wall_row, 5).value is None
    assert ws.cell(wall_row, 5).protection.locked is False
    assert ws.cell(wall_row, 4).protection.locked is not False
    assert ws.cell(wall_row, 6).value == f'=IF(E{wall_row}="","",D{wall_row}*E{wall_row})'
    assert ws.protection.sheet is True
    for row in ws.iter_rows():
        for cell in row:
            text = str(cell.value or "")
            for secret in ("187.31", "2341.38", "98765.43", "55.55", "Roof"):
                assert secret not in text, (cell.coordinate, text)


@pytest.mark.parametrize(
    ("token", "locale", "expected"),
    [
        ("m2", "en", "m²"),
        ("m3", "de", "m³"),
        ("lsum", "de", "psch"),
        ("pcs", "de", "Stk"),
        ("lsum", "en", "lsum"),
        ("Stk", "de", "Stk"),
        ("", "de", ""),
    ],
)
def test_units_read_like_the_boq_editor(token: str, locale: str, expected: str) -> None:
    assert unit_code(token, locale) == expected
