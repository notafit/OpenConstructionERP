# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The drawing set split, the project COBie and BI reports render off the loop.

Companion to ``test_exports_render_off_the_event_loop.py``, same method: a spy
records the thread the heavy call ran on, then calls the real one, and the test
fails if that thread is the loop's own, which is the thread the test coroutine
runs on.

A drawing set has no page limit and every page costs a text extraction and a
rendered thumbnail. The project COBie export unions every element of every BIM
model on the project into one workbook, then reopens it to add sheets. A BI
report lays out every KPI row and its drill-down rows with reportlab, and the
scheduler runs the same code on the same loop. All three used to do that work
inline in an ``async def``, which holds every other request on the worker.

The spies sit on the expensive library calls themselves (pdfplumber's
``extract_text`` and ``to_image``, the COBie builder, the report builder), not
on a helper around them, so moving the work back inline turns these red.
"""

from __future__ import annotations

import io
import threading
import uuid
from collections.abc import Callable
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock, MagicMock

import pytest


def _recording(seen: list[int], real: Callable[..., Any]) -> Callable[..., Any]:
    def spy(*args: Any, **kwargs: Any) -> Any:
        seen.append(threading.get_ident())
        return real(*args, **kwargs)

    return spy


# ── Drawing set split ─────────────────────────────────────────────────────


def _drawing_set(pages: list[list[str]]) -> bytes:
    canvas = pytest.importorskip("reportlab.pdfgen.canvas", reason="reportlab is not installed")
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for lines in pages:
        for offset, line in enumerate(lines):
            pdf.drawString(60, 760 - offset * 20, line)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _split_service(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    """A SheetService whose persistence is faked and whose files land in tmp."""
    from app.modules.documents import service as documents_service

    monkeypatch.setattr(documents_service, "UPLOAD_BASE", tmp_path / "uploads")
    monkeypatch.setattr(documents_service, "SHEET_THUMB_BASE", tmp_path / "sheets")

    class _Documents:
        def __init__(self, _session: Any) -> None:
            pass

        async def create(self, document: Any) -> Any:
            document.id = uuid.uuid4()
            return document

    monkeypatch.setattr(documents_service, "DocumentRepository", _Documents)
    monkeypatch.setattr(documents_service, "_register_version_safely", AsyncMock(return_value=None))
    service = documents_service.SheetService(MagicMock())
    service.repo = SimpleNamespace(create_many=AsyncMock(side_effect=lambda rows: rows))
    monkeypatch.setattr(service, "_supersede_previous_sheets", AsyncMock(return_value=0))
    return service


def _upload(content: bytes) -> Any:
    from fastapi import UploadFile

    return UploadFile(filename="drawings.pdf", file=io.BytesIO(content), headers={"content-type": "application/pdf"})


@pytest.mark.asyncio
async def test_drawing_set_pages_are_read_and_thumbnailed_off_the_loop(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import pdfplumber.page

    content = _drawing_set(
        [
            ["SHEET NO: A-101", "TITLE: Ground Floor Plan", "SCALE: 1:100", "REV B"],
            ["SHEET NO: S-201", "TITLE: Foundation Layout"],
            ["no title block on this page"],
        ]
    )
    text_threads: list[int] = []
    render_threads: list[int] = []
    monkeypatch.setattr(
        pdfplumber.page.Page, "extract_text", _recording(text_threads, pdfplumber.page.Page.extract_text)
    )
    monkeypatch.setattr(pdfplumber.page.Page, "to_image", _recording(render_threads, pdfplumber.page.Page.to_image))
    service = _split_service(monkeypatch, tmp_path)

    sheets = await service.split_pdf_to_sheets(uuid.uuid4(), _upload(content), str(uuid.uuid4()))

    loop_thread = threading.get_ident()
    assert len(text_threads) == 3 and loop_thread not in text_threads
    assert len(render_threads) == 3 and loop_thread not in render_threads
    # The rows built back on the loop carry what the worker read, page by page.
    assert [s.page_number for s in sheets] == [1, 2, 3]
    assert sheets[0].sheet_number == "A-101"
    assert sheets[0].discipline == "Architectural"
    assert sheets[1].sheet_number == "S-201"
    assert sheets[1].discipline == "Structural"
    assert sheets[2].sheet_number is None
    assert all(s.is_current for s in sheets)
    for sheet in sheets:
        assert sheet.thumbnail_path and Path(sheet.thumbnail_path).is_file()


@pytest.mark.asyncio
async def test_thumbnails_render_one_at_a_time(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    # PDFium must not be entered from two threads at once, and the split now
    # runs in a worker thread, so two uploads could reach it together. Every
    # render has to hold the module's lock.
    import pdfplumber.page

    from app.modules.documents import service as documents_service

    held: list[bool] = []
    real = pdfplumber.page.Page.to_image

    def spy(self: Any, *args: Any, **kwargs: Any) -> Any:
        held.append(documents_service._PDFIUM_RENDER_LOCK.locked())
        return real(self, *args, **kwargs)

    monkeypatch.setattr(pdfplumber.page.Page, "to_image", spy)
    service = _split_service(monkeypatch, tmp_path)

    await service.split_pdf_to_sheets(uuid.uuid4(), _upload(_drawing_set([["A-101"], ["A-102"]])), "u1")

    assert held == [True, True]


@pytest.mark.asyncio
async def test_each_page_drops_its_parsed_layout_before_the_next_is_read(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    # pdfplumber keeps every Page in ``pdf.pages`` and each Page caches its
    # parsed layout and objects until the document closes. A set has no page
    # limit and several uploads now split side by side in worker threads, so
    # holding every page's layout at once multiplies the peak by the page
    # count and again by the number of uploads. While page N is read, no
    # earlier page may still hold its parse.
    import pdfplumber.page

    real = pdfplumber.page.Page.extract_text
    still_cached: list[list[int]] = []

    def spy(self: Any, *args: Any, **kwargs: Any) -> Any:
        still_cached.append(
            [
                p.page_number
                for p in self.pdf.pages
                if p.page_number < self.page_number and (hasattr(p, "_layout") or hasattr(p, "_objects"))
            ]
        )
        return real(self, *args, **kwargs)

    monkeypatch.setattr(pdfplumber.page.Page, "extract_text", spy)
    service = _split_service(monkeypatch, tmp_path)

    sheets = await service.split_pdf_to_sheets(
        uuid.uuid4(), _upload(_drawing_set([["A-101"], ["A-102"], ["A-103"], ["A-104"]])), "u1"
    )

    assert [s.page_number for s in sheets] == [1, 2, 3, 4]
    assert still_cached == [[], [], [], []]


@pytest.mark.asyncio
async def test_a_page_whose_text_cannot_be_read_still_drops_its_parse(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    # The release runs in a ``finally``: a page that fails mid-read must not
    # leave its layout behind for the rest of a long failing request.
    import pdfplumber.page

    # Page objects, not numbers: closing the document rebuilds ``pdf.pages``
    # and closes the fresh copies, so a page number would be closed either way.
    # Holding the objects also keeps their ids from being reused.
    closed: list[Any] = []
    failed: list[Any] = []
    real_close = pdfplumber.page.Page.close

    def close_spy(self: Any) -> None:
        closed.append(self)
        real_close(self)

    def broken(self: Any, *_args: Any, **_kwargs: Any) -> Any:
        failed.append(self)
        _ = self.layout  # the parse is what the release has to drop
        raise RuntimeError("damaged content stream")

    monkeypatch.setattr(pdfplumber.page.Page, "close", close_spy)
    monkeypatch.setattr(pdfplumber.page.Page, "extract_text", broken)
    service = _split_service(monkeypatch, tmp_path)

    from fastapi import HTTPException

    with pytest.raises(HTTPException) as caught:
        await service.split_pdf_to_sheets(uuid.uuid4(), _upload(_drawing_set([["A-101"], ["A-102"]])), "u1")

    assert caught.value.status_code == 422
    assert len(failed) == 1
    assert any(page is failed[0] for page in closed)
    assert not hasattr(failed[0], "_layout")


@pytest.mark.asyncio
async def test_a_failed_thumbnail_still_leaves_the_page_its_row(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import pdfplumber.page

    from app.modules.documents import service as documents_service

    def broken(*_args: Any, **_kwargs: Any) -> Any:
        raise RuntimeError("renderer unavailable")

    monkeypatch.setattr(pdfplumber.page.Page, "to_image", broken)
    service = _split_service(monkeypatch, tmp_path)

    sheets = await service.split_pdf_to_sheets(uuid.uuid4(), _upload(_drawing_set([["A-101"], ["A-102"]])), "u1")

    assert [s.page_number for s in sheets] == [1, 2]
    assert all(s.thumbnail_path is None for s in sheets)
    # The lock is released after a failed render, or the next upload would hang.
    assert not documents_service._PDFIUM_RENDER_LOCK.locked()


@pytest.mark.asyncio
async def test_an_unreadable_drawing_set_is_still_a_422(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from fastapi import HTTPException

    service = _split_service(monkeypatch, tmp_path)

    with pytest.raises(HTTPException) as caught:
        await service.split_pdf_to_sheets(uuid.uuid4(), _upload(b"%PDF-1.4 this is not a pdf"), "u1")

    assert caught.value.status_code == 422
    service.repo.create_many.assert_not_awaited()


@pytest.mark.asyncio
async def test_drawing_index_tables_are_read_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    # The completeness check lifts the index tables from every page of the
    # index PDF when no page is named.
    from app.modules.documents import service as documents_service
    from app.modules.documents import sheet_index

    seen: list[int] = []
    expected = [
        sheet_index.ExpectedSheet(sheet_number="A-101", sheet_number_norm="A101", sheet_title="Plan", revision="B")
    ]
    monkeypatch.setattr(sheet_index, "parse_index_tables", _recording(seen, lambda _path, _page: expected))
    project_id = uuid.uuid4()
    index_doc = SimpleNamespace(id=uuid.uuid4(), project_id=project_id, file_path="index.pdf")

    class _Documents:
        def __init__(self, _session: Any) -> None:
            pass

        async def get_by_id(self, _doc_id: Any) -> Any:
            return index_doc

    monkeypatch.setattr(documents_service, "DocumentRepository", _Documents)
    service = documents_service.SheetService(MagicMock())
    # Stop right after the expected set is built: the rest is persistence.
    monkeypatch.setattr(service, "list_sheets", AsyncMock(side_effect=_Stop))

    with pytest.raises(_Stop):
        await service.check_completeness(project_id, index_document_id=index_doc.id)

    assert seen and threading.get_ident() not in seen


class _Stop(Exception):
    pass


# ── Project COBie ─────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_project_cobie_is_built_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.modules.bim_hub import repository as bim_repository
    from app.modules.reporting import exporters
    from app.modules.reporting.service import ReportingService

    project_id = uuid.uuid4()
    models = [
        SimpleNamespace(id=uuid.uuid4(), project_id=project_id, name="Architecture", discipline="Architectural"),
        SimpleNamespace(id=uuid.uuid4(), project_id=project_id, name="Services", discipline="MEP"),
    ]
    elements = {
        models[0].id: [
            SimpleNamespace(
                id=uuid.uuid4(),
                stable_id="room-101",
                element_type="IfcSpace",
                name="Office 101",
                storey="Level 1",
                discipline="Architectural",
                asset_info={},
                is_tracked_asset=False,
                quantities={"area": 42.5},
                properties={},
            )
        ],
        models[1].id: [
            SimpleNamespace(
                id=uuid.uuid4(),
                stable_id="ahu-01",
                element_type="AirHandlingUnit",
                name="AHU-01",
                storey="Level 1",
                discipline="MEP",
                asset_info={"manufacturer": "Generic", "model": "AH-100"},
                is_tracked_asset=True,
                quantities={},
                properties={},
            )
        ],
    }

    class _Models:
        def __init__(self, _session: Any) -> None:
            pass

        async def list_for_project(self, _project_id: Any, **_kwargs: Any) -> tuple[list[Any], int]:
            return models, len(models)

    class _Elements:
        def __init__(self, _session: Any) -> None:
            pass

        async def list_for_model(self, model_id: Any, **_kwargs: Any) -> tuple[list[Any], int]:
            rows = elements[model_id]
            return rows, len(rows)

    seen: list[int] = []
    monkeypatch.setattr(bim_repository, "BIMModelRepository", _Models)
    monkeypatch.setattr(bim_repository, "BIMElementRepository", _Elements)
    monkeypatch.setattr(exporters, "export_project_cobie", _recording(seen, exporters.export_project_cobie))
    service = ReportingService(MagicMock())
    monkeypatch.setattr(service, "_lookup_project_name", AsyncMock(return_value="Riverside Block C"))

    filename, media_type, blob = await service.export_project_cobie(project_id)

    assert blob[:2] == b"PK"
    assert filename.endswith(".xlsx")
    assert "spreadsheet" in media_type
    assert seen and threading.get_ident() not in seen
    # Both models reached the one workbook.
    from openpyxl import load_workbook

    names = {row[0] for row in load_workbook(io.BytesIO(blob))["Component"].iter_rows(min_row=2, values_only=True)}
    assert "AHU-01" in names


# ── BI report ─────────────────────────────────────────────────────────────


@pytest.mark.asyncio
@pytest.mark.parametrize(("output_format", "magic"), [("pdf", b"%PDF"), ("xlsx", b"PK")])
async def test_bi_report_file_is_rendered_off_the_loop(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, output_format: str, magic: bytes
) -> None:
    from app.modules.bi_dashboards import kpis
    from app.modules.bi_dashboards import service as bi_service

    monkeypatch.setenv("BI_REPORTS_DIR", str(tmp_path))
    seen: list[int] = []
    monkeypatch.setattr(bi_service, "build_report", _recording(seen, bi_service.build_report))
    monkeypatch.setattr(
        kpis,
        "compute",
        AsyncMock(
            return_value=SimpleNamespace(value="1250.00", unit="currency", source_record_count=3, breakdown={"eur": 1})
        ),
    )
    report = SimpleNamespace(
        id=uuid.uuid4(),
        code="monthly_cost",
        name="Monthly cost",
        description="Cost to date",
        output_format=output_format,
        query_spec_json={"kpis": ["eac", "cpi"]},
    )
    session = MagicMock()
    session.flush = AsyncMock()
    service = bi_service.BIDashboardsService(session)
    service.repo = SimpleNamespace(get_report=AsyncMock(return_value=report))

    response = await service.run_report(report.id)

    assert response is not None and response.row_count == 2
    assert seen and threading.get_ident() not in seen
    run = session.add.call_args.args[0]
    assert run.status == "success"
    assert Path(run.file_path).read_bytes()[: len(magic)] == magic


@pytest.mark.asyncio
async def test_bi_report_without_a_file_renders_nothing(monkeypatch: pytest.MonkeyPatch) -> None:
    # ``produce_file=False`` is the preview path: rows only, no render at all.
    from app.modules.bi_dashboards import kpis
    from app.modules.bi_dashboards import service as bi_service

    seen: list[int] = []
    monkeypatch.setattr(bi_service, "build_report", _recording(seen, bi_service.build_report))
    monkeypatch.setattr(
        kpis,
        "compute",
        AsyncMock(return_value=SimpleNamespace(value="1", unit="", source_record_count=1, breakdown={})),
    )
    report = SimpleNamespace(
        id=uuid.uuid4(),
        code="x",
        name="x",
        description=None,
        output_format="pdf",
        query_spec_json={"kpis": ["eac"]},
    )
    session = MagicMock()
    session.flush = AsyncMock()
    service = bi_service.BIDashboardsService(session)
    service.repo = SimpleNamespace(get_report=AsyncMock(return_value=report))

    response = await service.run_report(report.id, produce_file=False)

    assert response is not None and response.file_url is None
    assert seen == []


# ── Schedule MSPDI ────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_msp_xml_is_serialised_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    # A programme exports up to 5,000 tasks with their predecessor links.
    from app.modules.schedule import mspdi_export
    from app.modules.schedule import router as schedule_router

    seen: list[int] = []
    monkeypatch.setattr(mspdi_export, "build_mspdi_xml", _recording(seen, mspdi_export.build_mspdi_xml))
    monkeypatch.setattr(schedule_router, "_verify_schedule_owner", AsyncMock(return_value=None))
    first, second = uuid.uuid4(), uuid.uuid4()

    def activity(act_id: uuid.UUID, name: str, deps: list[dict[str, Any]]) -> SimpleNamespace:
        return SimpleNamespace(
            id=act_id,
            name=name,
            start_date="2026-10-05",
            end_date="2026-10-09",
            duration_days=5,
            progress_pct="0",
            activity_type="task",
            wbs_code="1.1",
            constraint_type=None,
            constraint_date=None,
            dependencies=deps,
        )

    activities = [
        activity(first, "Excavation", []),
        activity(second, "Blinding concrete", [{"activity_id": str(first), "type": "FS", "lag_days": 1}]),
    ]
    service = SimpleNamespace(
        get_schedule=AsyncMock(return_value=SimpleNamespace(name="Riverside programme")),
        list_activities_for_schedule=AsyncMock(return_value=(activities, len(activities))),
    )
    no_relationships = MagicMock()
    no_relationships.scalars.return_value.all.return_value = []
    session = MagicMock()
    session.execute = AsyncMock(return_value=no_relationships)

    response = await schedule_router.export_schedule_msp_xml(
        _user_id="u1", payload={}, schedule_id=uuid.uuid4(), service=service, session=session
    )

    body = b"".join([chunk async for chunk in response.body_iterator])
    assert b"<Name>Blinding concrete</Name>" in body
    assert b"<PredecessorUID>1</PredecessorUID>" in body
    assert seen and threading.get_ident() not in seen
