# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""DB-backed tests for the NRM 1 cost plan.

The plan is a regrouping of a bill, so the tests that matter compare it with the
bill as the rest of the platform reads it: its direct cost and grand total must
equal both ``compute_boq_totals`` (the bill list) and ``get_boq_structured``
(the editor and every export) to the cent. The markup case is built to tell a
plan that reads the bill's cascade from one that reorders or recomputes it:
three lines inserted in reverse ``sort_order``, a compounding line whose amount
differs from what it would be on direct cost, a lump sum, and an inactive line
that must contribute nothing.
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal

import pytest
import pytest_asyncio
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

import app.modules.boq.models  # noqa: F401
import app.modules.projects.models  # noqa: F401
from app.modules.boq.models import BOQ, BOQMarkup, Position
from app.modules.boq.service import BOQService
from app.modules.cost_plan.router import _verify_boq_access, export_nrm1_cost_plan
from app.modules.cost_plan.schemas import CostPlanResponse
from app.modules.cost_plan.service import build_nrm1_cost_plan, parse_gifa
from app.modules.projects.models import Project
from tests._pg import transactional_session

D = Decimal
CENT = D("0.01")


@pytest_asyncio.fixture
async def session() -> AsyncSession:
    """Isolated PostgreSQL session, FK triggers off, rolled back on teardown."""
    async with transactional_session(disable_fks=True) as sess:
        yield sess


async def _project(session: AsyncSession, *, gfa: str | None = "1000") -> Project:
    project = Project(
        name="Elemental Office",
        owner_id=uuid.uuid4(),
        currency="GBP",
        fx_rates=[{"code": "USD", "rate": "2", "label": "US Dollar"}],
        gross_floor_area=gfa,
    )
    session.add(project)
    await session.flush()
    return project


def _pos(boq: BOQ, ordinal: str, total: str, *, parent: Position | None = None, **extra: object) -> Position:
    return Position(
        boq_id=boq.id,
        parent_id=parent.id if parent is not None else None,
        ordinal=ordinal,
        description=extra.pop("description", f"Item {ordinal}"),
        unit=extra.pop("unit", "m2"),
        quantity=extra.pop("quantity", "1"),
        unit_rate=extra.pop("unit_rate", total),
        total=total,
        **extra,
    )


async def _bill(session: AsyncSession, project: Project) -> BOQ:
    """A bill exercising every placement path, with known totals.

    Elements: 1.1 = 1000, 2.5 = 250, 2.7 = 40 (inherited from its section),
    group 2 level = 30, 5.10 = 200 (a 100 USD line at 2), 9 = 15 (preliminaries
    priced in the bill). Unallocated: 12.34 (no code) + 9 (own bad code under a
    coded section). A section header with a stale total must not leak in.
    Direct cost = 1556.34.
    """
    boq = BOQ(project_id=project.id, name="Stage 2 cost plan")
    session.add(boq)
    await session.flush()

    walls = _pos(boq, "02", "999", unit="", quantity="0", unit_rate="0", classification={"nrm": "2.7"})
    session.add(walls)
    await session.flush()
    session.add_all(
        [
            _pos(boq, "01.001", "1000", classification={"nrm": "1.1"}),
            _pos(boq, "01.002", "250", classification={"nrm": "02.05"}),
            _pos(boq, "02.001", "40", parent=walls),
            _pos(boq, "02.002", "9", parent=walls, classification={"nrm": "abc"}),
            _pos(boq, "03.001", "30", classification={"nrm": "2"}),
            _pos(boq, "04.001", "100", classification={"nrm": "5.10"}, metadata_={"currency": "USD"}),
            _pos(boq, "05.001", "15", classification={"nrm": "9"}),
            _pos(boq, "06.001", "12.34"),
        ]
    )
    await session.flush()
    return boq


async def _markups(session: AsyncSession, boq: BOQ) -> None:
    """Inserted in reverse order on purpose: ``sort_order`` must decide."""
    lines = [
        BOQMarkup(
            boq_id=boq.id,
            name="Inactive risk",
            category="contingency",
            markup_type="percentage",
            percentage="50",
            apply_to="cumulative",
            sort_order=3,
            is_active=False,
        ),
        BOQMarkup(
            boq_id=boq.id,
            name="Design fees",
            category="other",
            markup_type="fixed",
            fixed_amount="500",
            apply_to="direct_cost",
            sort_order=2,
        ),
        BOQMarkup(
            boq_id=boq.id,
            name="Overheads and profit",
            category="profit",
            markup_type="percentage",
            percentage="10",
            apply_to="cumulative",
            sort_order=1,
        ),
        BOQMarkup(
            boq_id=boq.id,
            name="Preliminaries",
            category="overhead",
            markup_type="percentage",
            percentage="13",
            apply_to="direct_cost",
            sort_order=0,
        ),
    ]
    for line in lines:
        session.add(line)
        await session.flush()


def _all_groups(plan: CostPlanResponse):
    return [*plan.groups, *plan.addon_groups]


def _element_total(plan: CostPlanResponse, code: str) -> Decimal:
    for group in _all_groups(plan):
        for element in group.elements:
            if element.code == code:
                return element.total
    raise AssertionError(code)


# ── Parity with the bill ─────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_plan_conserves_the_bill_and_matches_both_bill_rollups(session: AsyncSession) -> None:
    project = await _project(session)
    boq = await _bill(session, project)
    await _markups(session, boq)

    plan = await build_nrm1_cost_plan(session, boq.id)

    direct = D("1556.34")
    assert plan.direct_cost.total == direct
    addons = sum((g.total for g in plan.addon_groups), D("0"))
    facilitating = plan.facilitating_works_estimate.total
    building = plan.building_works_estimate.total
    assert facilitating + building + addons + plan.unallocated.total == direct
    assert facilitating == D("0")
    assert building == D("1520")
    assert addons == D("15")
    assert plan.unallocated.total == D("21.34")
    assert plan.unallocated.position_count == 2

    assert _element_total(plan, "1.1") == D("1000")
    assert _element_total(plan, "2.5") == D("250")
    assert _element_total(plan, "2.7") == D("40")
    assert _element_total(plan, "5.10") == D("200")
    assert plan.inherited_count == 1

    totals = (await BOQService(session).compute_boq_totals([boq.id]))[boq.id]
    structured = await BOQService(session).get_boq_structured(boq.id)
    assert float(plan.direct_cost.total.quantize(CENT)) == totals["direct_cost"]
    assert float(plan.grand_total.total.quantize(CENT)) == totals["grand_total"]
    assert plan.direct_cost.total.quantize(CENT) == D(str(structured.direct_cost)).quantize(CENT)
    assert plan.grand_total.total.quantize(CENT) == D(str(structured.grand_total)).quantize(CENT)


@pytest.mark.asyncio
async def test_cascade_follows_sort_order_and_compounds_on_the_running_subtotal(session: AsyncSession) -> None:
    project = await _project(session)
    boq = await _bill(session, project)
    await _markups(session, boq)

    plan = await build_nrm1_cost_plan(session, boq.id)

    assert [m.name for m in plan.markups] == ["Preliminaries", "Overheads and profit", "Design fees"]
    direct = plan.direct_cost.total
    prelims, ohp, fees = plan.markups
    assert prelims.total == direct * D("13") / D("100")
    assert prelims.base == direct
    # The distinguishing case: on direct cost this line would be 155.634.
    assert ohp.base == direct + prelims.total
    assert ohp.total == (direct + prelims.total) * D("10") / D("100")
    assert ohp.total != direct * D("10") / D("100")
    assert fees.total == D("500")
    assert fees.base is None
    assert fees.running_total == direct + prelims.total + ohp.total + D("500")
    # The inactive 50% line is neither printed nor priced.
    assert plan.grand_total.total == fees.running_total
    assert plan.markups_total.total == prelims.total + ohp.total + fees.total


@pytest.mark.asyncio
async def test_blank_rows_from_add_position_change_nothing_in_the_plan(session: AsyncSession) -> None:
    """A row nobody has typed into is not a position: no count, no warning."""
    project = await _project(session)
    boq = await _bill(session, project)
    before = await build_nrm1_cost_plan(session, boq.id)

    walls = next(p for p in await BOQService(session).position_repo.list_all_for_boq(boq.id) if p.ordinal == "02")
    # A real unit, so the row is not mistaken for a section header.
    blank = {"description": "", "unit": "m2", "quantity": "0", "unit_rate": "0"}
    session.add_all([_pos(boq, "07.001", "0", **blank), _pos(boq, "02.003", "0", parent=walls, **blank)])
    await session.flush()
    after = await build_nrm1_cost_plan(session, boq.id)

    assert after.position_count == before.position_count
    assert after.allocated_count == before.allocated_count
    assert after.unallocated.position_count == before.unallocated.position_count == 2
    assert len(after.unallocated.positions) == 2
    assert after.warnings == before.warnings
    assert _element_count(after, "2.7") == _element_count(before, "2.7") == 1
    assert after.direct_cost.total == before.direct_cost.total


def _element_count(plan: CostPlanResponse, code: str) -> int:
    for group in _all_groups(plan):
        for element in group.elements:
            if element.code == code:
                return element.position_count
    raise AssertionError(code)


# ── GIFA ─────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_gifa_comes_from_the_project_and_an_entered_value_wins(session: AsyncSession) -> None:
    project = await _project(session, gfa="1000")
    boq = await _bill(session, project)

    from_project = await build_nrm1_cost_plan(session, boq.id)
    assert from_project.gifa_source == "project"
    assert from_project.gifa == D("1000")
    assert from_project.grand_total.cost_per_m2 == (from_project.grand_total.total / D("1000")).quantize(CENT)

    entered = await build_nrm1_cost_plan(session, boq.id, entered_gifa=D("500"))
    assert entered.gifa_source == "entered"
    assert entered.grand_total.cost_per_m2 == (entered.grand_total.total / D("500")).quantize(CENT)
    assert entered.grand_total.cost_per_m2 != from_project.grand_total.cost_per_m2


@pytest.mark.asyncio
@pytest.mark.parametrize("stored", [None, "", "n/a", "0", "-5"])
async def test_an_unusable_project_area_means_no_cost_per_m2(session: AsyncSession, stored: str | None) -> None:
    project = await _project(session, gfa=stored)
    boq = await _bill(session, project)
    plan = await build_nrm1_cost_plan(session, boq.id)
    assert plan.gifa is None
    assert plan.gifa_source == "none"
    assert plan.grand_total.cost_per_m2 is None
    assert "no_gifa" in plan.warnings


def test_parse_gifa() -> None:
    assert parse_gifa("1250.5") == D("1250.5")
    assert parse_gifa(" 800 ") == D("800")
    assert parse_gifa("abc") is None
    assert parse_gifa("NaN") is None
    assert parse_gifa("0") is None


# ── Access ───────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_a_stranger_and_a_missing_bill_both_get_404(session: AsyncSession) -> None:
    project = await _project(session)
    boq = await _bill(session, project)

    await _verify_boq_access(session, boq.id, str(project.owner_id))
    with pytest.raises(HTTPException) as stranger:
        await _verify_boq_access(session, boq.id, str(uuid.uuid4()))
    assert stranger.value.status_code == 404
    with pytest.raises(HTTPException) as missing:
        await _verify_boq_access(session, uuid.uuid4(), str(project.owner_id))
    assert missing.value.status_code == 404


# ── Excel ────────────────────────────────────────────────────────────────────


def _rows(sheet) -> list[tuple]:
    return [tuple(cell for cell in row) for row in sheet.iter_rows(values_only=True)]


def _row_by_label(rows: list[tuple], text: str) -> tuple:
    for row in rows:
        if len(row) > 1 and row[1] == text:
            return row
    raise AssertionError(f"no row labelled {text!r}")


@pytest.mark.asyncio
async def test_export_has_the_same_structure_and_totals_as_the_plan(session: AsyncSession) -> None:
    project = await _project(session)
    boq = await _bill(session, project)
    await _markups(session, boq)
    plan = await build_nrm1_cost_plan(session, boq.id)

    response = await export_nrm1_cost_plan(
        boq.id, str(project.owner_id), session, gifa=None, locale=None, accept_language="en-GB"
    )
    assert response.headers["Content-Language"] == "en"
    assert "NRM1" in response.headers["Content-Disposition"]
    book = load_workbook(io.BytesIO(response.body))
    rows = _rows(book.worksheets[0])

    assert _row_by_label(rows, "Building works estimate")[4] == pytest.approx(float(plan.building_works_estimate.total))
    assert _row_by_label(rows, "Facilitating works estimate")[4] == pytest.approx(0.0)
    assert _row_by_label(rows, "Not allocated to an element")[4] == pytest.approx(float(plan.unallocated.total))
    assert _row_by_label(rows, "Direct cost of the bill")[4] == pytest.approx(float(plan.direct_cost.total))
    total_row = _row_by_label(rows, "Cost plan total")
    assert total_row[4] == pytest.approx(float(plan.grand_total.total))
    assert total_row[5] == pytest.approx(float(plan.grand_total.cost_per_m2))
    ohp = _row_by_label(rows, "Overheads and profit")
    assert ohp[3] == "10% on running subtotal"
    assert ohp[4] == pytest.approx(float(plan.markups[1].total))
    assert _row_by_label(rows, "External walls")[0] == "2.5"
    assert _row_by_label(rows, "Main contractor's preliminaries")[4] == pytest.approx(15.0)

    unallocated = _rows(book.worksheets[1])
    assert len(unallocated) == 1 + plan.unallocated.position_count
    assert {row[2] or "" for row in unallocated[1:]} == {"", "abc"}


@pytest.mark.asyncio
async def test_export_language_follows_the_request_and_says_when_it_fell_back(session: AsyncSession) -> None:
    project = await _project(session)
    boq = await _bill(session, project)

    german = await export_nrm1_cost_plan(boq.id, str(project.owner_id), session, gifa=None, locale="de")
    assert german.headers["Content-Language"] == "de"
    rows = _rows(load_workbook(io.BytesIO(german.body)).worksheets[0])
    assert _row_by_label(rows, "Kosten der Bauleistungen")

    fallback = await export_nrm1_cost_plan(boq.id, str(project.owner_id), session, gifa=None, locale="fr")
    assert fallback.headers["Content-Language"] == "en"
